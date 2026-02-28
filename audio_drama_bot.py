"""
Audio Drama Bot - All-in-One Script for Google Colab

Features:
1. Telegram Bot для прийому файлів (.txt, .pdf, .epub, .fb2, .doc, .docx)
2. Gemini Integration для переробки тексту на сценарій по ролях
3. Генерація описів голосів персонажів (мін. 20 слів)
4. Тестові голоси 15+ сек через Qwen3-TTS-VoiceDesign
5. Інтерфейс перегенерації голосів через Telegram команди
6. Pexels API для фонових звуків
7. Мікшування аудіо (репліки + фонові звуки)
8. Збірка фінального аудіофайлу
9. Контрольні точки (1 хв, 30 хв) з інлайн кнопками
10. Консистентність голосів між сесіями

API Keys (from Colab secrets):
- TELEGRAM_BOT_TOKEN
- TELEGRAM_CHAT_ID
- GEMINI_API_KEY
- PEXELS_API_KEY

Version: 1.1.20
Author: VIBEMODLY Team
Created: 2026-02-22
"""

# =============================================================================
# SECTION 1: IMPORTS та DEPENDENCIES
# =============================================================================

import os
import sys
import json
import hashlib
import logging
import asyncio
import tempfile
import re
from datetime import datetime
from dataclasses import dataclass, field, asdict
from typing import Optional, Dict, List, Any, Tuple, Union, Callable
from pathlib import Path
from enum import Enum
import subprocess
import glob
import shutil
import difflib
import importlib.util
import random
import threading
from urllib.parse import quote
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FuturesTimeoutError

# Для стабільної роботи CUDA-детермінізму (використовується перед torch-операціями).
os.environ.setdefault("CUBLAS_WORKSPACE_CONFIG", ":4096:8")

# Налаштування логування
logging.basicConfig(
    level=logging.INFO,
    format='[%(asctime)s] [%(name)s] %(levelname)s: %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger('AudioDramaBot')
# Зменшуємо шум від сторонніх пакетів, які спамлять warning-ами в Colab
logging.getLogger('sox').setLevel(logging.ERROR)

NARRATOR_NAME = "Диктор"
NARRATOR_ALIASES = {
    "диктор",
    "диктора",
    "диктору",
    "диктором",
    "дикторе",
    "оповідач",
    "оповідача",
    "рассказчик",
    "рассказчика",
    "narrator",
    "diktor",
    "opovidach",
}

LANGUAGE_ALIASES: Dict[str, str] = {
    "ru": "russian",
    "rus": "russian",
    "russian": "russian",
    "рус": "russian",
    "русский": "russian",
    "русская": "russian",
    "рос": "russian",
    "російська": "russian",
    "російською": "russian",
    "uk": "ukrainian",
    "ua": "ukrainian",
    "ukr": "ukrainian",
    "ukrainian": "ukrainian",
    "українська": "ukrainian",
    "українською": "ukrainian",
    "укр": "ukrainian",
    "en": "english",
    "eng": "english",
    "english": "english",
    "англ": "english",
    "английский": "english",
    "англійська": "english",
    "англійською": "english",
}

SUPPORTED_PRONUNCIATION_LANGUAGES = {"russian", "ukrainian", "english"}
MALE_NAME_EXCEPTIONS = {
    "никита",
    "илья",
    "илья",
    "кузьма",
    "лука",
}


def is_narrator_alias(name: Optional[str]) -> bool:
    """Перевіряє чи ім'я є варіантом оповідача/диктора."""
    if not name:
        return False
    normalized = str(name).strip().lower()
    normalized = normalized.strip(" \t\n\r\"'`«»[](){}:;,.!?-")
    if normalized in NARRATOR_ALIASES:
        return True
    normalized_compact = re.sub(r"[^a-zа-яіїєґё]", "", normalized)
    return normalized_compact in NARRATOR_ALIASES


def canonical_character_name(name: Optional[str]) -> str:
    """Повертає канонічне ім'я персонажа."""
    if is_narrator_alias(name):
        return NARRATOR_NAME
    value = (name or "").strip()
    value = value.strip(" \t\n\r\"'`«»[](){}:;,.!?-")
    return value


def normalize_gender(gender: Optional[str], default: str = "male") -> str:
    """Нормалізує стать до male/female/neutral."""
    if not gender:
        return default
    normalized = str(gender).strip().lower()
    if normalized in {"male", "man", "чоловік", "чоловіча", "чоловічий", "мужской", "мужской голос"}:
        return "male"
    if normalized in {"female", "woman", "жінка", "жіноча", "жіночий", "женский", "женский голос"}:
        return "female"
    if normalized in {"neutral", "нейтральний", "нейтральна", "нейтрально"}:
        return "neutral"
    return default


def extract_gender_from_text(text: Optional[str]) -> Optional[str]:
    """Витягує запитувану стать голосу з текстової команди."""
    normalized = (text or "").strip().lower()
    if not normalized:
        return None
    if re.search(r"(жіноч|female|woman|женск)", normalized):
        return "female"
    if re.search(r"(чоловіч|male|man|мужск)", normalized):
        return "male"
    if re.search(r"(нейтрал|neutral)", normalized):
        return "neutral"
    return None


def build_gender_aware_prompt(character: "Character") -> str:
    """Формує prompt для TTS з явно зафіксованою статтю голосу."""
    gender = normalize_gender(character.gender, "male")
    base_description = (character.voice_description or "").strip()
    if character.role == "narrator":
        # Для диктора не використовуємо суперечливий текст від LLM, щоби уникати фліпів статі.
        if gender == "female":
            base_description = "Голос Диктора рівний, впевнений, з чіткою дикцією та стабільним жіночим тембром."
        elif gender == "neutral":
            base_description = "Голос Диктора рівний, впевнений, з чіткою дикцією та нейтральним тембром."
        else:
            base_description = "Голос Диктора рівний, впевнений, з чіткою дикцією та стабільним чоловічим тембром."
    elif not base_description:
        base_description = f"Голос {character.name} природний, виразний і добре артикульований."
    if gender == "female":
        gender_hint = (
            "ОБОВ'ЯЗКОВО жіночий голос, без чоловічого тембру. "
            "Female voice only. Женский голос. Pitch: medium-high."
        )
    elif gender == "neutral":
        gender_hint = (
            "ОБОВ'ЯЗКОВО нейтральний голос, без явного чоловічого або жіночого тембру. "
            "Neutral voice. Pitch: medium."
        )
    else:
        gender_hint = (
            "ОБОВ'ЯЗКОВО чоловічий голос, без жіночого тембру. "
            "Male voice only. Мужской голос. Pitch: low-medium."
        )
    return f"{base_description} {gender_hint}".strip()


def normalize_language_code(value: Optional[str], default: str = "russian") -> str:
    """Нормалізує код/назву мови до russian/ukrainian/english."""
    if not value:
        return default
    token = str(value).strip().lower()
    token = re.sub(r"[^a-zа-яіїєґ]+", "", token)
    normalized = LANGUAGE_ALIASES.get(token)
    if normalized:
        return normalized
    if token in SUPPORTED_PRONUNCIATION_LANGUAGES:
        return token
    return default


def language_human_name(language: str) -> str:
    """Людська назва мови для промптів/повідомлень."""
    lang = normalize_language_code(language, "russian")
    if lang == "ukrainian":
        return "українська"
    if lang == "english":
        return "english"
    return "русский"


LANGUAGE_SHORT_CODES: Dict[str, str] = {
    "russian": "ru",
    "ukrainian": "uk",
    "english": "en",
}

NUMERIC_POINT_WORD: Dict[str, str] = {
    "russian": "целых",
    "ukrainian": "цілих",
    "english": "point",
}

PERCENT_WORD: Dict[str, str] = {
    "russian": "процентов",
    "ukrainian": "відсотків",
    "english": "percent",
}

DIGIT_WORDS: Dict[str, Tuple[str, ...]] = {
    "russian": ("ноль", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"),
    "ukrainian": ("нуль", "один", "два", "три", "чотири", "п'ять", "шість", "сім", "вісім", "дев'ять"),
    "english": ("zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine"),
}

STRESS_VOWELS: Dict[str, str] = {
    "russian": "аеёиоуыэюя",
    "ukrainian": "аеєиіїоуюя",
}

RUSSIAN_YO_HINTS: Dict[str, str] = {
    "ежик": "ёжик",
    "ежика": "ёжика",
    "ежики": "ёжики",
    "елка": "ёлка",
    "елки": "ёлки",
    "елочку": "ёлочку",
    "еще": "ещё",
    "все": "всё",
    "все-таки": "всё-таки",
    "трех": "трёх",
    "четырех": "четырёх",
    "берет": "берёт",
    "ее": "её",
}

NUMBER_TOKEN_PATTERN = re.compile(r"(?<![\w])([+-]?\d+(?:[.,]\d+)?)(%?)(?![\w])")
STRESS_WORD_PATTERN = re.compile(r"[А-Яа-яЁёІіЇїЄєҐґ']+")

ALLOWED_EMOTIONS = {
    "neutral",
    "happy",
    "sad",
    "angry",
    "excited",
    "calm",
    "dramatic",
    "whisper",
    "shout",
}

EMOTION_ALIASES: Dict[str, str] = {
    "neutral": "neutral",
    "нейтрально": "neutral",
    "нейтральний": "neutral",
    "нейтральный": "neutral",
    "спокійно": "calm",
    "спокойно": "calm",
    "calm": "calm",
    "happy": "happy",
    "радісно": "happy",
    "радостно": "happy",
    "sad": "sad",
    "сумно": "sad",
    "грустно": "sad",
    "angry": "angry",
    "злісно": "angry",
    "зло": "angry",
    "сердито": "angry",
    "excited": "excited",
    "схвильовано": "excited",
    "взволнованно": "excited",
    "dramatic": "dramatic",
    "драматично": "dramatic",
    "драматичною": "dramatic",
    "whisper": "whisper",
    "шепотом": "whisper",
    "пошепки": "whisper",
    "шепотом": "whisper",
    "шепіт": "whisper",
    "shout": "shout",
    "крик": "shout",
    "кричить": "shout",
    "кричит": "shout",
    "громко": "shout",
}

EMOTION_SPEED_MAP: Dict[str, float] = {
    "neutral": 1.0,
    "calm": 0.97,
    "happy": 1.01,
    "sad": 0.96,
    "angry": 1.01,
    "excited": 1.03,
    "dramatic": 0.99,
    "whisper": 0.94,
    "shout": 1.02,
}


def language_short_code(language: str) -> str:
    """Повертає короткий код мови для ASR/TTS API."""
    lang = normalize_language_code(language, "russian")
    return LANGUAGE_SHORT_CODES.get(lang, "ru")


def _digit_to_word(digit: str, language: str) -> str:
    """Перетворює одну цифру у слово (fallback)."""
    lang = normalize_language_code(language, "russian")
    try:
        index = int(digit)
    except Exception:
        return digit
    words = DIGIT_WORDS.get(lang) or DIGIT_WORDS["russian"]
    if index < 0 or index >= len(words):
        return digit
    return words[index]


def _fallback_number_to_words(raw_number: str, language: str) -> str:
    """Fallback-конвертація числа у слова поцифрово."""
    lang = normalize_language_code(language, "russian")
    number = (raw_number or "").strip()
    sign_prefix = ""
    if number.startswith("-"):
        sign_prefix = "минус " if lang == "russian" else "мінус " if lang == "ukrainian" else "minus "
        number = number[1:]
    elif number.startswith("+"):
        number = number[1:]

    normalized = number.replace(",", ".")
    if "." in normalized:
        int_part, frac_part = normalized.split(".", 1)
        int_words = " ".join(_digit_to_word(ch, lang) for ch in int_part if ch.isdigit()) or _digit_to_word("0", lang)
        frac_words = " ".join(_digit_to_word(ch, lang) for ch in frac_part if ch.isdigit()) or _digit_to_word("0", lang)
        point_word = NUMERIC_POINT_WORD.get(lang, "point")
        return f"{sign_prefix}{int_words} {point_word} {frac_words}".strip()

    digits_words = " ".join(_digit_to_word(ch, lang) for ch in normalized if ch.isdigit())
    if not digits_words:
        return raw_number
    return f"{sign_prefix}{digits_words}".strip()


def _num_to_words(raw_number: str, language: str) -> str:
    """Конвертує число у слова через num2words, fallback поцифрово."""
    lang = normalize_language_code(language, "russian")
    normalized = (raw_number or "").strip().replace(",", ".")
    if not normalized:
        return raw_number

    if num2words is None:
        return _fallback_number_to_words(raw_number, lang)

    lang_code = LANGUAGE_SHORT_CODES.get(lang, "ru")
    try:
        if "." in normalized:
            int_part, frac_part = normalized.split(".", 1)
            int_value = int(int_part or "0")
            int_words = str(num2words(int_value, lang=lang_code))
            frac_words = " ".join(_digit_to_word(ch, lang) for ch in frac_part if ch.isdigit()) or _digit_to_word("0", lang)
            point_word = NUMERIC_POINT_WORD.get(lang, "point")
            return f"{int_words} {point_word} {frac_words}".strip()
        return str(num2words(int(normalized), lang=lang_code))
    except Exception:
        return _fallback_number_to_words(raw_number, lang)


def expand_numbers_for_tts(text: str, language: str) -> str:
    """Розкриває числа словами для стабільнішого TTS."""
    if not text or not re.search(r"\d", text):
        return text
    lang = normalize_language_code(language, "russian")

    def _replace(match: re.Match) -> str:
        raw_number = match.group(1)
        percent_mark = match.group(2) or ""
        words = _num_to_words(raw_number, lang)
        if percent_mark:
            percent_word = PERCENT_WORD.get(lang, "percent")
            words = f"{words} {percent_word}".strip()
        return words

    return NUMBER_TOKEN_PATTERN.sub(_replace, text)


def apply_inline_stress_markers(text: str, language: str) -> str:
    """Підтримує stress-нотацію в стилі рАса/расА -> ра́са/раса́."""
    if not text:
        return text
    lang = normalize_language_code(language, "russian")
    vowels = STRESS_VOWELS.get(lang)
    if not vowels:
        return text

    def _replace_word(match: re.Match) -> str:
        word = match.group(0)
        if len(word) < 2:
            return word
        if len(word) > 1 and word[0].isupper() and word[1:].islower():
            return word
        if word.isupper():
            return word

        stressed_positions = [i for i, ch in enumerate(word) if ch.lower() in vowels and ch.isupper()]
        if len(stressed_positions) != 1:
            return word

        pos = stressed_positions[0]
        lowered = word.lower()
        chars = list(lowered)
        chars[pos] = chars[pos] + "́"  # combining acute accent
        return "".join(chars)

    return STRESS_WORD_PATTERN.sub(_replace_word, text)


def ensure_terminal_punctuation(text: str, role: str = "supporting") -> str:
    """Гарантує природне завершення репліки, щоб уникнути обриву інтонації."""
    value = (text or "").strip()
    if not value:
        return value
    if value.endswith((".", "!", "?", "…")):
        return value
    if role == "narrator":
        return value + "."
    return value + "."


def infer_emotion_from_text(text: str) -> str:
    """Евристично визначає емоцію, якщо LLM не повернула коректну."""
    value = (text or "").strip().lower()
    if not value:
        return "neutral"
    if "..." in value or "…" in value:
        return "sad"
    if "!" in value and "?" in value:
        return "dramatic"
    if "!" in value:
        return "excited"
    return "neutral"


def normalize_emotion_label(value: Optional[str], fallback_text: str = "") -> str:
    """Нормалізує emotion до одного з дозволених токенів."""
    token = (value or "").strip().lower()
    token = token.replace("ё", "е")
    token = re.sub(r"[^a-zа-яіїєґ_ -]", "", token)
    token = re.sub(r"\s+", " ", token).strip()
    if token in ALLOWED_EMOTIONS:
        return token
    if token in EMOTION_ALIASES:
        return EMOTION_ALIASES[token]
    if token:
        for part in token.split():
            if part in ALLOWED_EMOTIONS:
                return part
            if part in EMOTION_ALIASES:
                return EMOTION_ALIASES[part]
    return infer_emotion_from_text(fallback_text)


def emotion_speed_multiplier(emotion: str, role: str = "supporting") -> float:
    """Повертає множник швидкості TTS за емоцією."""
    normalized = normalize_emotion_label(emotion)
    speed = EMOTION_SPEED_MAP.get(normalized, 1.0)
    if role == "narrator":
        speed = min(speed, 1.0)
    return max(0.84, min(1.15, float(speed)))


def emotion_prompt_hint(emotion: str, language: str = "russian") -> str:
    """Короткий prompt-hint для стилю озвучки."""
    emo = normalize_emotion_label(emotion)
    lang = normalize_language_code(language, "russian")
    if lang == "ukrainian":
        return f"Емоція репліки: {emo}. Подача має бути природною, без різкого обриву на фіналі."
    if lang == "english":
        return f"Line emotion: {emo}. Keep natural cadence and avoid abrupt cutoff at sentence ending."
    return f"Эмоция реплики: {emo}. Подача естественная, без резкого обрыва на финале фразы."


def localized_voice_description(
    character_name: str,
    role: str = "supporting",
    gender: str = "male",
    language: str = "russian"
) -> str:
    """Повертає дефолтний опис голосу мовою поточного тексту."""
    lang = normalize_language_code(language, "russian")
    gender_norm = normalize_gender(gender, "male")
    
    if lang == "ukrainian":
        if role == "narrator":
            gender_hint = "жіночим тембром" if gender_norm == "female" else "нейтральним тембром" if gender_norm == "neutral" else "чоловічим тембром"
            return (
                f"Голос {character_name} рівний, виразний і спокійний, з чіткою дикцією, "
                f"природними паузами, {gender_hint}, стабільною гучністю та впевненим ритмом "
                "для довгих оповідних фрагментів."
            )
        return (
            f"Голос {character_name} природний, емоційно гнучкий і добре артикульований, "
            "з виразними інтонаціями, чіткими приголосними, контрольованою динамікою, "
            "живим тембром та стабільною подачею протягом усієї сцени."
        )
    
    if lang == "english":
        if role == "narrator":
            gender_hint = "female timbre" if gender_norm == "female" else "neutral timbre" if gender_norm == "neutral" else "male timbre"
            return (
                f"{character_name}'s voice is steady, expressive, and calm, with clear diction, "
                f"natural pauses, a {gender_hint}, stable loudness, and confident pacing for long narration."
            )
        return (
            f"{character_name}'s voice is natural, emotionally flexible, and well articulated, with expressive "
            "intonation, clear consonants, controlled dynamics, and a stable delivery across the whole scene."
        )
    
    # russian default
    if role == "narrator":
        gender_hint = "женским тембром" if gender_norm == "female" else "нейтральным тембром" if gender_norm == "neutral" else "мужским тембром"
        return (
            f"Голос {character_name} ровный, выразительный и спокойный, с четкой дикцией, "
            f"естественными паузами, {gender_hint}, стабильной громкостью и уверенным ритмом "
            "для длительного повествования."
        )
    return (
        f"Голос {character_name} естественный, эмоционально гибкий и хорошо артикулированный, "
        "с выразительными интонациями, четкими согласными, контролируемой динамикой и "
        "стабильной подачей в течение всей сцены."
    )


def detect_text_language(text: str, default: str = "russian") -> str:
    """Евристично визначає мову тексту (english/ukrainian/russian)."""
    if not text:
        return default
    
    sample = text[:5000]
    latin = len(re.findall(r"[A-Za-z]", sample))
    cyrillic = len(re.findall(r"[А-Яа-яІіЇїЄєҐґЁёЪъЫыЭэ]", sample))
    uk_specific = len(re.findall(r"[ІіЇїЄєҐґ]", sample))
    ru_specific = len(re.findall(r"[ЁёЪъЫыЭэ]", sample))
    
    if latin > cyrillic * 1.3 and latin > 40:
        return "english"
    if cyrillic > 0:
        if uk_specific > ru_specific:
            return "ukrainian"
        return "russian"
    return default


def infer_gender_from_name(name: str) -> Optional[str]:
    """Евристично оцінює стать персонажа за іменем."""
    normalized = re.sub(r"[^a-zа-яіїєґё]", "", (name or "").strip().lower())
    if not normalized:
        return None
    normalized_alt = phonetic_name_key(normalized) or normalized
    if normalized in MALE_NAME_EXCEPTIONS:
        return "male"
    if normalized in {"елена", "elena", "helen", "ольга", "наталья", "maria", "мария", "valkyrie", "valkyria", "valkyriea"}:
        return "female"
    if normalized_alt in {"валкирия", "валкирие", "валькирия"}:
        return "female"
    
    female_suffixes = (
        "а", "я", "ия", "ія", "ея",
        "a", "ia", "ya", "ina", "ena", "anna", "ella", "ette", "ine", "ira"
    )
    male_suffixes = (
        "ов", "ев", "ин", "ич", "ыч", "ий", "ой", "ей",
        "son", "sen", "er", "or", "ard", "ik"
    )
    
    if any(normalized.endswith(suffix) for suffix in female_suffixes):
        return "female"
    if any(normalized_alt.endswith(suffix) for suffix in female_suffixes):
        return "female"
    if any(normalized.endswith(suffix) for suffix in male_suffixes):
        return "male"
    if any(normalized_alt.endswith(suffix) for suffix in male_suffixes):
        return "male"
    return None


def resolve_character_gender(name: str, role: str, requested_gender: Optional[str]) -> str:
    """Нормалізує та коригує стать персонажа."""
    role_norm = (role or "").strip().lower()
    if role_norm == "narrator" or is_narrator_alias(name):
        return "male"
    
    requested = normalize_gender(requested_gender, "neutral")
    inferred = infer_gender_from_name(name)
    
    if requested in {"male", "female"} and inferred:
        # Якщо ім'я явно жіноче, а модель віддала male, виправляємо.
        if requested == "male" and inferred == "female":
            return "female"
        return requested
    
    if requested == "neutral" and inferred:
        return inferred
    
    if requested in {"male", "female"}:
        return requested
    return inferred or "male"


def normalize_name_for_match(name: str) -> str:
    """Нормалізує ім'я для пошуку дублікатів."""
    value = (name or "").strip().lower()
    value = value.replace("ё", "е").replace("й", "и").replace("і", "и").replace("ї", "и")
    value = value.replace("ы", "и").replace("ъ", "").replace("ь", "")
    value = value.replace("ґ", "г").replace("’", "'").replace("`", "'")
    value = value.strip(" \t\n\r\"'`«»[](){}:;,.!?-")
    value = re.sub(r"[^a-zа-яіїєґ0-9]", "", value)
    return value


def _latin_to_cyrillic_approx(value: str) -> str:
    """Груба транслітерація латиниці для кращого merge імен (shnir ~ шнир)."""
    if not value:
        return value
    s = value.lower()
    combos = [
        ("shch", "щ"),
        ("sch", "щ"),
        ("zh", "ж"),
        ("ch", "ч"),
        ("sh", "ш"),
        ("yu", "ю"),
        ("ya", "я"),
        ("yo", "е"),
        ("ye", "е"),
        ("kh", "х"),
        ("ts", "ц"),
        ("iy", "ий"),
    ]
    for src, dst in combos:
        s = s.replace(src, dst)
    table = str.maketrans({
        "a": "а",
        "b": "б",
        "v": "в",
        "g": "г",
        "d": "д",
        "e": "е",
        "z": "з",
        "i": "и",
        "j": "й",
        "k": "к",
        "l": "л",
        "m": "м",
        "n": "н",
        "o": "о",
        "p": "п",
        "r": "р",
        "s": "с",
        "t": "т",
        "u": "у",
        "f": "ф",
        "h": "х",
        "c": "к",
        "q": "к",
        "w": "в",
        "x": "кс",
        "y": "и",
    })
    return s.translate(table)


def phonetic_name_key(name: str) -> str:
    """Ключ для порівняння імен незалежно від дрібних орфографічних/трансліт відмінностей."""
    normalized = normalize_name_for_match(name)
    if not normalized:
        return normalized
    if re.search(r"[a-z]", normalized):
        normalized = _latin_to_cyrillic_approx(normalized)
        normalized = normalize_name_for_match(normalized)
    normalized = normalized.replace("кк", "к").replace("нн", "н").replace("лл", "л")
    return normalized


def are_names_equivalent(left: str, right: str) -> bool:
    """Евристично визначає, що два імені - це той самий герой."""
    if not left or not right:
        return False
    if canonical_character_name(left) == canonical_character_name(right):
        return True
    a = phonetic_name_key(left)
    b = phonetic_name_key(right)
    if not a or not b:
        return False
    if a == b:
        return True
    if a.startswith(b) or b.startswith(a):
        if min(len(a), len(b)) >= 5 and abs(len(a) - len(b)) <= 1:
            return True
    ratio = difflib.SequenceMatcher(a=a, b=b).ratio()
    if ratio >= 0.93 and abs(len(a) - len(b)) <= 1 and a[0] == b[0]:
        return True
    return False


def resolve_name_alias(candidate_name: str, existing_names: List[str]) -> str:
    """Повертає вже наявне ім'я персонажа, якщо candidate є його варіантом."""
    candidate = canonical_character_name(candidate_name) or candidate_name
    if is_narrator_alias(candidate):
        return NARRATOR_NAME
    for existing in existing_names:
        if are_names_equivalent(candidate, existing):
            return existing
    return candidate


class ResourceLimitPauseError(RuntimeError):
    """Помилка, яка сигналізує про паузу через обмеження ресурсів (GPU/Colab)."""


class MissingReferenceVoiceError(RuntimeError):
    """Помилка strict-mode: у персонажа є референс, але voice prompt недоступний."""

    def __init__(
        self,
        character_name: str,
        scene_id: Optional[int] = None,
        order_in_scene: Optional[int] = None,
        reference_audio_path: str = "",
        reason: str = "voice_prompt_missing",
        message: Optional[str] = None
    ):
        self.character_name = character_name or ""
        self.scene_id = scene_id
        self.order_in_scene = order_in_scene
        self.reference_audio_path = reference_audio_path or ""
        self.reason = reason or "voice_prompt_missing"
        super().__init__(
            message
            or f"Missing reference voice prompt for '{self.character_name}' ({self.reason})"
        )


def is_resource_limit_error(exc: Exception) -> bool:
    """Евристично визначає помилки, що схожі на ліміт ресурсів Colab/GPU."""
    message = str(exc or "").lower()
    markers = [
        "cuda out of memory",
        "cublas_status_alloc_failed",
        "no cuda gpus are available",
        "cuda error",
        "out of memory",
        "resource exhausted",
        "quota",
        "limit reached",
        "insufficient resources",
        "device-side assert",
        "runtime disconnected",
        "t4",
    ]
    return any(marker in message for marker in markers)

# =============================================================================
# AUTO-INSTALL DEPENDENCIES
# =============================================================================

DEPENDENCY_IMPORT_PROBES: Dict[str, Tuple[str, ...]] = {
    "python-telegram-bot": ("telegram",),
    "google-genai": ("google.genai",),
    "google-generativeai": ("google.generativeai",),
    "python-docx": ("docx",),
    "faster-whisper": ("faster_whisper",),
}


def _dependency_import_candidates(package: str) -> List[str]:
    """Повертає список модулів для import-check конкретного пакету."""
    base_package = re.split(r"[<>=!~]", (package or "").strip(), maxsplit=1)[0].strip()
    if not base_package:
        return []
    candidates: List[str] = list(DEPENDENCY_IMPORT_PROBES.get(base_package, ()))
    fallback = base_package.replace("-", "_")
    if fallback and fallback not in candidates:
        candidates.append(fallback)
    return candidates


def _is_dependency_installed(package: str) -> bool:
    """Повертає True, якщо хоча б один import-probe успішний."""
    for module_name in _dependency_import_candidates(package):
        try:
            if importlib.util.find_spec(module_name) is None:
                continue
            return True
        except Exception:
            continue
    return False


def install_dependencies() -> bool:
    """Автоматичне встановлення всіх необхідних залежностей."""
    dependencies = [
        'python-telegram-bot',
        'google-genai',
        'google-generativeai',
        'pypdf',
        'ebooklib',
        'python-docx',
        'pypandoc',
        'aiohttp',
        'numpy',
        'scipy',
        'librosa',
        'soundfile',
        'pydub',
        'num2words',
        'openai',
        'faster-whisper',
        'nest_asyncio',
    ]
    
    print("\n" + "="*60, flush=True)
    print("[AudioDramaBot] AUTO-INSTALLING DEPENDENCIES", flush=True)
    print("="*60, flush=True)
    
    installed = []
    failed = []
    
    for package in dependencies:
        if _is_dependency_installed(package):
            print(f"[AudioDramaBot]   ✓ {package} already installed", flush=True)
            continue
        try:
            print(f"[AudioDramaBot]   Installing {package}...", flush=True)
            subprocess.check_call(
                [sys.executable, '-m', 'pip', 'install', '--disable-pip-version-check', '-q', package],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL
            )
            if _is_dependency_installed(package):
                print(f"[AudioDramaBot]   ✓ {package} installed", flush=True)
                installed.append(package)
            else:
                raise RuntimeError("import check failed after install")
        except Exception as e:
            print(f"[AudioDramaBot]   ✗ Failed to install {package}: {e}", flush=True)
            failed.append(package)
    
    print("\n[AudioDramaBot] Installation Summary:", flush=True)
    if installed:
        print(f"  Installed: {', '.join(installed)}", flush=True)
    if failed:
        print(f"  Failed: {', '.join(failed)}", flush=True)
    print("="*60 + "\n", flush=True)
    
    return len(failed) == 0


def ensure_colab_system_tools() -> None:
    """Встановлює системні інструменти в Colab за потреби (sox, ffmpeg)."""
    if not ('google.colab' in sys.modules or os.path.exists('/content')):
        return
    
    required_tools = ['sox', 'ffmpeg']
    missing_tools = [tool for tool in required_tools if shutil.which(tool) is None]
    
    if not missing_tools:
        return
    
    logger.info(f"[AudioDramaBot] Installing missing system tools: {', '.join(missing_tools)}")
    try:
        subprocess.check_call(
            ['apt-get', 'update', '-qq'],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL
        )
        subprocess.check_call(
            ['apt-get', 'install', '-y', '-qq'] + missing_tools,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL
        )
        logger.info("[AudioDramaBot] System tools installed")
    except Exception as e:
        logger.warning(f"[AudioDramaBot] Failed to install system tools ({', '.join(missing_tools)}): {e}")


# Автоматичне встановлення залежностей
install_dependencies()
ensure_colab_system_tools()

# =============================================================================
# ОСНОВНІ ІМПОРТИ (після встановлення залежностей)
# =============================================================================

try:
    import numpy as np
except ImportError:
    np = None
    logger.warning("numpy not available")

try:
    import aiohttp
except ImportError:
    aiohttp = None
    logger.warning("aiohttp not available")

try:
    from num2words import num2words
except ImportError:
    num2words = None
    logger.warning("num2words not available")

try:
    from openai import OpenAI
except ImportError:
    OpenAI = None
    logger.warning("openai package not available")

try:
    from faster_whisper import WhisperModel as FasterWhisperModel
except ImportError:
    FasterWhisperModel = None
    logger.warning("faster-whisper package not available")

# =============================================================================
# КОРИСТУВАЦЬКІ ТИПИ
# =============================================================================

# Type aliases для кращої читабельності
AudioData = np.ndarray
FilePath = str
URL = str
JsonDict = Dict[str, Any]

# =============================================================================
# SECTION 1.5: INTERNAL TTS FALLBACKS (NO EXTERNAL PROJECT FILES REQUIRED)
# =============================================================================

def _apply_deterministic_seed(seed_value: Any) -> Optional[int]:
    """Застосовує deterministic seed для random/numpy/torch."""
    try:
        seed = int(seed_value)
    except Exception:
        return None
    if seed < 0:
        seed = abs(seed)

    try:
        random.seed(seed)
    except Exception:
        pass

    try:
        if np is not None:
            np.random.seed(seed % (2**32 - 1))
    except Exception:
        pass

    try:
        import torch
        torch.manual_seed(seed)
        if torch.cuda.is_available():
            torch.cuda.manual_seed(seed)
            torch.cuda.manual_seed_all(seed)
        if hasattr(torch, "use_deterministic_algorithms"):
            torch.use_deterministic_algorithms(True, warn_only=True)
        if hasattr(torch, "backends") and hasattr(torch.backends, "cudnn"):
            torch.backends.cudnn.deterministic = True
            torch.backends.cudnn.benchmark = False
    except Exception:
        pass

    return seed


def _normalize_audio_array(audio: Any) -> np.ndarray:
    """Нормалізує довільний audio output у 1D numpy float32."""
    if audio is None:
        return np.array([], dtype=np.float32)
    
    if hasattr(audio, "cpu"):
        audio = audio.cpu()
    if hasattr(audio, "numpy"):
        audio = audio.numpy()
    
    arr = np.array(audio, dtype=np.float32).flatten()
    if arr.size == 0:
        return arr
    
    max_abs = float(np.max(np.abs(arr)))
    if max_abs > 1.0:
        arr = arr / max_abs
    
    return arr


def _change_audio_speed(audio: np.ndarray, speed: float) -> np.ndarray:
    """Швидка зміна швидкості без важких залежностей."""
    if audio is None or audio.size == 0:
        return np.array([], dtype=np.float32)
    if speed <= 0 or abs(speed - 1.0) < 1e-6:
        return audio
    
    src_idx = np.arange(len(audio), dtype=np.float32)
    dst_idx = np.arange(0.0, len(audio), speed, dtype=np.float32)
    if dst_idx.size == 0:
        return audio
    
    sped = np.interp(dst_idx, src_idx, audio).astype(np.float32)
    return sped


class InternalVoiceDesignEngine:
    """Вбудований VoiceDesign engine на qwen-tts."""
    
    def __init__(self, model_name: str = "Qwen/Qwen3-TTS-12Hz-1.7B-VoiceDesign"):
        self._model_name = model_name
        self._model = None
        self._sample_rate = 24000
    
    def load_model(self, device: Optional[str] = None) -> bool:
        try:
            import torch
            from qwen_tts import Qwen3TTSModel
            
            if device is None:
                device = "cuda" if torch.cuda.is_available() else "cpu"
            
            kwargs: Dict[str, Any] = {}
            if device == "cuda":
                kwargs["device_map"] = "auto"
                kwargs["torch_dtype"] = torch.float16
            else:
                kwargs["device_map"] = "cpu"
                kwargs["torch_dtype"] = torch.float32
            
            logger.info(f"[InternalVoiceDesign] Loading model {self._model_name} on {device}...")
            self._model = Qwen3TTSModel.from_pretrained(self._model_name, **kwargs)
            logger.info("[InternalVoiceDesign] Model loaded")
            return True
        except Exception as e:
            logger.error(f"[InternalVoiceDesign] Model load failed: {e}")
            return False
    
    def is_loaded(self) -> bool:
        return self._model is not None
    
    def set_model(self, model) -> None:
        self._model = model
    
    def generate(
        self,
        text: str,
        voice_config: Dict[str, Any],
        language: str = "russian"
    ) -> Tuple[Optional[np.ndarray], int]:
        if not self.is_loaded():
            return None, 0
        
        prompt = voice_config.get("prompt", "A natural voice")
        speed = float(voice_config.get("speed", 1.0) or 1.0)
        _apply_deterministic_seed(voice_config.get("seed"))
        
        try:
            result = None
            if hasattr(self._model, "generate_voice_design"):
                result = self._model.generate_voice_design(text, prompt, language)
            elif hasattr(self._model, "generate"):
                try:
                    result = self._model.generate(text=text, prompt=prompt, language=language)
                except TypeError:
                    result = self._model.generate(text, prompt, language)
            elif hasattr(self._model, "synthesize"):
                result = self._model.synthesize(text, prompt, language)
            
            if result is None:
                return None, 0
            
            if isinstance(result, tuple):
                audio, sr = result
            else:
                audio, sr = result, self._sample_rate
            
            audio_arr = _normalize_audio_array(audio)
            if audio_arr.size == 0:
                return None, 0
            
            if speed != 1.0:
                audio_arr = _change_audio_speed(audio_arr, speed)
            
            return audio_arr, int(sr or self._sample_rate)
        except Exception as e:
            logger.error(f"[InternalVoiceDesign] Generate failed: {e}")
            return None, 0


class InternalVoiceCloneEngine:
    """Вбудований VoiceClone engine на qwen-tts."""
    
    def __init__(
        self,
        model_name: str = "Qwen/Qwen3-TTS-12Hz-1.7B-Base",
        cache_dir: Optional[str] = None
    ):
        self._model_name = model_name
        self._model = None
        self._sample_rate = 24000
        self._prompt_cache: Dict[str, Dict[str, Any]] = {}
        self._cache_dir = Path(cache_dir) if cache_dir else None
        if self._cache_dir is not None:
            self._cache_dir.mkdir(parents=True, exist_ok=True)
    
    def load_model(self, device: Optional[str] = None) -> bool:
        try:
            import torch
            from qwen_tts import Qwen3TTSModel
            
            if device is None:
                device = "cuda" if torch.cuda.is_available() else "cpu"
            
            kwargs: Dict[str, Any] = {}
            if device == "cuda":
                kwargs["device_map"] = "auto"
                kwargs["torch_dtype"] = torch.float16
            else:
                kwargs["device_map"] = "cpu"
                kwargs["torch_dtype"] = torch.float32
            
            logger.info(f"[InternalVoiceClone] Loading model {self._model_name} on {device}...")
            self._model = Qwen3TTSModel.from_pretrained(self._model_name, **kwargs)
            logger.info("[InternalVoiceClone] Model loaded")
            return True
        except Exception as e:
            logger.error(f"[InternalVoiceClone] Model load failed: {e}")
            return False
    
    def is_loaded(self) -> bool:
        return self._model is not None
    
    def set_model(self, model) -> None:
        self._model = model
    
    def _build_cache_key(self, reference_audio_path: str, reference_text: str) -> str:
        payload = f"{reference_audio_path}|{reference_text}".encode("utf-8")
        return hashlib.md5(payload).hexdigest()
    
    def create_voice_prompt(
        self,
        reference_audio_path: str,
        reference_text: str,
        use_cache: bool = True,
        validate: bool = True
    ) -> Tuple[Optional[Dict[str, Any]], bool]:
        if not self.is_loaded():
            return None, False
        if not reference_audio_path or not os.path.exists(reference_audio_path):
            return None, False
        
        cache_key = self._build_cache_key(reference_audio_path, reference_text)
        if use_cache and cache_key in self._prompt_cache:
            return self._prompt_cache[cache_key], False
        
        try:
            prompt = None
            if hasattr(self._model, "create_voice_clone_prompt"):
                prompt = self._model.create_voice_clone_prompt(
                    ref_audio=reference_audio_path,
                    ref_text=reference_text,
                    x_vector_only_mode=False
                )
            elif hasattr(self._model, "create_voice_prompt"):
                prompt = self._model.create_voice_prompt(reference_audio_path, reference_text)
            
            if prompt is None:
                return None, False
            
            if use_cache:
                self._prompt_cache[cache_key] = prompt
            
            return prompt, True
        except Exception as e:
            logger.error(f"[InternalVoiceClone] create_voice_prompt failed: {e}")
            return None, False
    
    def generate(
        self,
        text: str,
        voice_config: Dict[str, Any],
        language: str = "russian"
    ) -> Tuple[Optional[np.ndarray], int]:
        if not self.is_loaded():
            return None, 0
        
        speed = float(voice_config.get("speed", 1.0) or 1.0)
        instruct = voice_config.get("instruct")
        voice_prompt = voice_config.get("voice_prompt")
        _apply_deterministic_seed(voice_config.get("seed"))
        
        if voice_prompt is None:
            reference_audio = voice_config.get("reference_audio")
            reference_text = voice_config.get("reference_text")
            if reference_audio and reference_text:
                voice_prompt, _ = self.create_voice_prompt(
                    reference_audio_path=reference_audio,
                    reference_text=reference_text,
                    use_cache=True,
                    validate=True
                )
        
        if voice_prompt is None:
            return None, 0
        
        try:
            result = None
            if hasattr(self._model, "generate_voice_clone"):
                result = self._model.generate_voice_clone(
                    text=text,
                    voice_clone_prompt=voice_prompt,
                    instruct=instruct
                )
            elif hasattr(self._model, "generate"):
                result = self._model.generate(
                    text=text,
                    voice_prompt=voice_prompt,
                    language=language
                )
            
            if result is None:
                return None, 0
            
            if isinstance(result, tuple):
                audio, sr = result
            else:
                audio, sr = result, self._sample_rate
            
            audio_arr = _normalize_audio_array(audio)
            if audio_arr.size == 0:
                return None, 0
            
            if speed != 1.0:
                audio_arr = _change_audio_speed(audio_arr, speed)
            
            return audio_arr, int(sr or self._sample_rate)
        except Exception as e:
            logger.error(f"[InternalVoiceClone] Generate failed: {e}")
            return None, 0


class SimpleVoiceLibraryManager:
    """Легка бібліотека reference-голосів без зовнішніх модулів."""
    
    LIBRARY_FILE = "voice_library_light.json"
    
    def __init__(self, clone_engine: InternalVoiceCloneEngine, drive_path: str):
        self._engine_clone = clone_engine
        self._drive_path = drive_path
        self._library: Dict[str, Dict[str, Any]] = {}
        os.makedirs(self._drive_path, exist_ok=True)
    
    def _library_path(self) -> str:
        return os.path.join(self._drive_path, self.LIBRARY_FILE)
    
    def load_library(self) -> bool:
        path = self._library_path()
        if not os.path.exists(path):
            self._library = {}
            return True
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            self._library = data.get("entries", {})
            return True
        except Exception as e:
            logger.warning(f"[SimpleVoiceLibrary] load failed: {e}")
            self._library = {}
            return False
    
    def save_library(self) -> bool:
        path = self._library_path()
        try:
            payload = {
                "version": "1.0.0",
                "updated_at": datetime.now().isoformat(),
                "entries": self._library
            }
            with open(path, "w", encoding="utf-8") as f:
                json.dump(payload, f, ensure_ascii=False, indent=2)
            return True
        except Exception as e:
            logger.warning(f"[SimpleVoiceLibrary] save failed: {e}")
            return False
    
    def register_reference(
        self,
        character_name: str,
        reference_audio_path: str,
        reference_text: str,
        voice_preset: str = "custom",
        gender: str = "male",
        language: str = "russian"
    ) -> bool:
        if not character_name or not reference_audio_path:
            return False
        
        rel_path = reference_audio_path
        try:
            rel_path = os.path.relpath(reference_audio_path, self._drive_path)
        except Exception:
            pass
        
        self._library[character_name] = {
            "reference_audio_path": rel_path,
            "reference_text": reference_text,
            "voice_preset": voice_preset,
            "gender": gender,
            "language": language,
            "updated_at": datetime.now().isoformat()
        }
        return self.save_library()
    
    def get_voice(
        self,
        character_name: str,
        voice_preset: str,
        gender: str = "male",
        language: str = "russian"
    ) -> Tuple[Optional[Dict[str, Any]], bool]:
        entry = self._library.get(character_name)
        if not entry:
            return None, False

        entry_gender = normalize_gender(entry.get("gender"), "male")
        requested_gender = normalize_gender(gender, "male")
        if entry_gender != requested_gender:
            logger.warning(
                f"[SimpleVoiceLibrary] Skip '{character_name}' due gender mismatch: "
                f"entry={entry_gender}, requested={requested_gender}"
            )
            return None, False

        entry_language = normalize_language_code(entry.get("language"), "russian")
        requested_language = normalize_language_code(language, "russian")
        if entry_language != requested_language:
            logger.warning(
                f"[SimpleVoiceLibrary] Skip '{character_name}' due language mismatch: "
                f"entry={entry_language}, requested={requested_language}"
            )
            return None, False
        
        ref_audio_path = entry.get("reference_audio_path")
        ref_text = entry.get("reference_text")
        if not ref_audio_path or not ref_text:
            return None, False
        
        full_path = ref_audio_path
        if not os.path.isabs(ref_audio_path):
            full_path = os.path.join(self._drive_path, ref_audio_path)
        
        if not os.path.exists(full_path):
            return None, False
        
        voice_prompt, _ = self._engine_clone.create_voice_prompt(
            reference_audio_path=full_path,
            reference_text=ref_text,
            use_cache=True,
            validate=True
        )
        return voice_prompt, False

# =============================================================================
# SECTION 2: API CLIENTS
# =============================================================================

class GeminiClient:
    """Клієнт для Gemini API - переробка тексту на сценарій."""
    
    MODEL_NAME = os.getenv("GEMINI_MODEL_NAME", "gemini-3-flash-preview")
    
    def __init__(self, api_key: str):
        self.api_key = api_key
        self._model = None
        self._client = None
        self._backend = "none"
        self._initialized = False
    
    def _ensure_initialized(self) -> bool:
        """Ініціалізує Gemini модель при першому використанні."""
        if self._initialized:
            return self._backend != "none"

        modern_error: Optional[Exception] = None
        try:
            try:
                from google import genai as modern_genai
            except Exception:
                import google.genai as modern_genai
            self._client = modern_genai.Client(api_key=self.api_key)
            self._backend = "google.genai"
            self._initialized = True
            logger.info(f"[GeminiClient] Initialized with model {self.MODEL_NAME} via google.genai")
            return True
        except Exception as e:
            modern_error = e

        try:
            import google.generativeai as genai
            genai.configure(api_key=self.api_key)
            self._model = genai.GenerativeModel(self.MODEL_NAME)
            self._backend = "google.generativeai"
            self._initialized = True
            if modern_error is not None:
                logger.warning(
                    "[GeminiClient] google.genai init failed, fallback to google.generativeai: "
                    f"{modern_error}"
                )
            logger.info(f"[GeminiClient] Initialized with model {self.MODEL_NAME} via google.generativeai")
            return True
        except Exception as e:
            logger.error(f"[GeminiClient] Initialization failed: {e}")
            if modern_error is not None:
                logger.error(f"[GeminiClient] google.genai initialization error: {modern_error}")
            self._initialized = True
            return False

    def _extract_response_text(self, response: Any) -> str:
        """Уніфікує витяг тексту з відповідей різних Gemini SDK."""
        text_value = getattr(response, "text", None)
        if isinstance(text_value, str) and text_value.strip():
            return text_value
        return str(text_value or response or "")

    def _generate_content_sync(
        self,
        prompt: str,
        generation_config: Optional[Dict[str, Any]] = None
    ) -> str:
        """Уніфікований sync виклик Gemini для обох SDK."""
        if self._backend == "google.genai" and self._client is not None:
            config_value: Any = None
            if generation_config:
                config_payload = dict(generation_config)
                try:
                    from google.genai import types as genai_types
                    config_value = genai_types.GenerateContentConfig(**config_payload)
                except Exception:
                    config_value = config_payload
            response = self._client.models.generate_content(
                model=self.MODEL_NAME,
                contents=prompt,
                config=config_value
            )
            return self._extract_response_text(response)

        if self._model is None:
            raise RuntimeError("Gemini model is not initialized")
        response = self._model.generate_content(
            prompt,
            generation_config=generation_config if generation_config else None
        )
        return self._extract_response_text(response)
    
    async def convert_to_scenario(
        self,
        text: str,
        title: str = "Untitled",
        language: Optional[str] = None
    ) -> 'Scenario':
        """Перетворює текст у структурований сценарій."""
        if not self._ensure_initialized():
            raise RuntimeError("Gemini client not initialized")
        
        scenario_language = normalize_language_code(language or detect_text_language(text, "russian"), "russian")
        prompt = self._build_scenario_prompt(text, title, scenario_language)
        
        try:
            response_text = await asyncio.to_thread(
                self._generate_content_sync,
                prompt,
                generation_config={
                    "temperature": 0.7,
                    "max_output_tokens": 8192,
                }
            )
            
            scenario_data = self._parse_scenario_response(
                response_text,
                title,
                scenario_language
            )
            return Scenario(**scenario_data)
        except Exception as e:
            logger.error(f"[GeminiClient] Scenario conversion failed: {e}")
            raise
    
    async def extract_characters(self, text: str, language: Optional[str] = None) -> List['Character']:
        """Витягує персонажів з тексту."""
        if not self._ensure_initialized():
            raise RuntimeError("Gemini client not initialized")
        
        out_language = normalize_language_code(language or detect_text_language(text, "russian"), "russian")
        out_language_name = language_human_name(out_language)
        
        prompt = f"""Проаналізуй текст та витягни всіх персонажів.
Для кожного персонажа надай:
- ім'я (name)
- роль (role): narrator, protagonist, supporting
- стать (gender): male, female, neutral
- вік (age_range): child, young, middle, elderly
- риси характеру (personality_traits): список 3-5 рис
- опис голосу (voice_description): МІНІМУМ 20 слів, детальний опис голосу
Якщо є оповідач, його name повинен бути строго "{NARRATOR_NAME}".
Для "{NARRATOR_NAME}" став gender="male" за замовченням.
Уважно визначай стать інших персонажів за іменем і контекстом, не плутай male/female.
Поле voice_description пиши мовою: {out_language_name}.

Текст:
{text[:5000]}

Відповідь у форматі JSON:
{{
    "characters": [
        {{
            "name": "ім'я",
            "role": "role",
            "gender": "gender",
            "age_range": "age_range",
            "personality_traits": ["trait1", "trait2"],
            "voice_description": "детальний опис голосу мінімум 20 слів..."
        }}
    ]
        }}"""
        
        try:
            response_text = await asyncio.to_thread(
                self._generate_content_sync,
                prompt
            )
            
            data = self._extract_json(response_text)
            characters: List[Character] = []
            for raw in data.get("characters", []):
                raw_name = canonical_character_name(raw.get("name", ""))
                if not raw_name:
                    continue
                role = (raw.get("role", "supporting") or "supporting").strip().lower()
                if raw_name == NARRATOR_NAME:
                    role = "narrator"
                if role not in {"narrator", "protagonist", "supporting", "antagonist", "minor"}:
                    role = "supporting"
                voice_description = (raw.get("voice_description", "") or "").strip()
                if not voice_description:
                    voice_description = localized_voice_description(
                        character_name=raw_name,
                        role=role,
                        gender=raw.get("gender", "male"),
                        language=out_language
                    )
                desc_gender = extract_gender_from_text(voice_description)
                gender = resolve_character_gender(raw_name, role, desc_gender or raw.get("gender"))
                characters.append(Character(
                    name=raw_name,
                    role=role,
                    gender=gender,
                    age_range=raw.get("age_range", "middle"),
                    personality_traits=raw.get("personality_traits", []),
                    voice_description=voice_description
                ))
            return characters
        except Exception as e:
            err_text = str(e or "")
            lowered = err_text.lower()
            if "429" in lowered or "quota" in lowered or "rate limit" in lowered:
                logger.warning(
                    "[GeminiClient] Character extraction quota/rate-limited. "
                    "Falling back to heuristic character inference from scenario dialogues."
                )
            else:
                logger.error(f"[GeminiClient] Character extraction failed: {e}")
            return []
    
    async def describe_sounds(
        self,
        scenario: 'Scenario',
        language: Optional[str] = None
    ) -> List['SoundCue']:
        """Генерує описи фонових звуків для сцен."""
        if not self._ensure_initialized():
            raise RuntimeError("Gemini client not initialized")
        
        out_language = normalize_language_code(language or "russian", "russian")
        out_language_name = language_human_name(out_language)
        
        sounds = []
        for scene in scenario.scenes:
            prompt = f"""Для сцени опиши фонові звуки.

Сцена: {scene.title}
Опис: {scene.setting}
Опис description пиши мовою: {out_language_name}.
keywords поверни англійською, короткі 1-3 слова для пошуку на Pexels.

Відповідь у форматі JSON:
{{
    "sounds": [
        {{
            "description": "опис звуку",
            "keywords": ["keyword1", "keyword2"],
            "duration_estimate": 30.0,
            "volume": 0.3
        }}
    ]
}}"""
            
            try:
                response_text = await asyncio.to_thread(
                    self._generate_content_sync,
                    prompt
                )
                
                data = self._extract_json(response_text)
                for i, sound_data in enumerate(data.get("sounds", [])):
                    sounds.append(SoundCue(
                        scene_id=scene.id,
                        description=sound_data.get("description", ""),
                        keywords=sound_data.get("keywords", []),
                        start_time=0.0,
                        duration=sound_data.get("duration_estimate", 30.0),
                        volume=sound_data.get("volume", 0.3)
                    ))
            except Exception as e:
                logger.warning(f"[GeminiClient] Sound description failed for scene {scene.id}: {e}")
        
        return sounds
    
    def _build_scenario_prompt(self, text: str, title: str, language: str = "russian") -> str:
        """Будує prompt для генерації сценарію."""
        output_language = language_human_name(language)
        return f"""Перетвори текст у сценарій аудіоп'єси.

Правила:
1. Розділи на сцени (кожна сцена - окрема локація/час)
2. Для кожної репліки вкажи персонажа та емоцію
3. Додай опис сеттингу для кожної сцени
4. Збережи весь оригінальний текст
5. Якщо є оповідач, використовуй ім'я персонажа строго "{NARRATOR_NAME}"
6. Для реплік персонажів не використовуй приписки типу "він сказав", "вона сказала", "he said", "she said"
7. Текст {NARRATOR_NAME} має бути повністю переробленим описом дії/атмосфери без фраз "он сказал/она сказала/він сказав/he said"
8. Усі тексти в полях scene.title, scene.setting і dialogue_lines.text мають бути мовою: {output_language}
9. emotion поверни ТІЛЬКИ з набору: neutral, happy, sad, angry, excited, calm, dramatic, whisper, shout
10. Емоцію визначай із контексту конкретної сцени, відносин між героями та розвитку події, а не випадково
11. Для кожної репліки обов'язково став осмислену emotion; не зловживай neutral
12. Кожна репліка повинна завершуватись природним розділовим знаком (. ! ? …), не обривай фрази
13. УСІ цифри та числа в текстових полях обов'язково пиши словами цією ж мовою (напр. "666" -> "шістсот шістдесят шість")

Текст:
{text[:10000]}

Відповідь у форматі JSON:
{{
    "title": "{title}",
    "author": "автор якщо відомо",
    "scenes": [
        {{
            "id": 1,
            "title": "назва сцени",
            "setting": "опис місця дії",
            "dialogue_lines": [
                {{
                    "character_name": "ім'я",
                    "text": "текст репліки",
                    "emotion": "емоція",
                    "order_in_scene": 1
                }}
            ]
        }}
    ],
    "total_duration_estimate": 300.0
}}"""
    
    def _parse_scenario_response(
        self,
        response_text: str,
        title: str,
        language: str = "russian"
    ) -> Dict:
        """Парсить відповідь Gemini у структуру сценарію."""
        scenario_language = normalize_language_code(language, "russian")
        data = self._extract_json(response_text)
        
        if not data:
            # Fallback - створюємо простий сценарій
            fallback_text = expand_numbers_for_tts(response_text[:500], scenario_language)
            return {
                "title": title,
                "author": "Unknown",
                "scenes": [Scene(
                    id=1,
                    title="Main Scene",
                    setting="Default setting",
                    dialogue_lines=[DialogueLine(
                        character_name=NARRATOR_NAME,
                        text=fallback_text,
                        emotion="neutral",
                        scene_id=1,
                        order_in_scene=1
                    )],
                    sound_cues=[]
                )],
                "characters": [],
                "total_duration_estimate": 60.0
            }
        
        # Перетворюємо dialogue_lines у об'єкти
        scenes = []
        for scene_data in data.get("scenes", []):
            dialogue_lines = []
            for dl_data in scene_data.get("dialogue_lines", []):
                character_name = canonical_character_name(dl_data.get("character_name", NARRATOR_NAME))
                if not character_name:
                    character_name = NARRATOR_NAME
                cleaned_text = self._clean_dialogue_text(
                    dl_data.get("text", ""),
                    character_name
                )
                cleaned_text = expand_numbers_for_tts(cleaned_text, scenario_language)
                normalized_emotion = normalize_emotion_label(
                    dl_data.get("emotion", "neutral"),
                    cleaned_text
                )
                role_hint = "narrator" if is_narrator_alias(character_name) else "supporting"
                cleaned_text = ensure_terminal_punctuation(cleaned_text, role_hint)
                dialogue_lines.append(DialogueLine(
                    character_name=character_name,
                    text=cleaned_text,
                    emotion=normalized_emotion,
                    scene_id=scene_data.get("id", 1),
                    order_in_scene=dl_data.get("order_in_scene", 1)
                ))
            
            scenes.append(Scene(
                id=scene_data.get("id", 1),
                title=expand_numbers_for_tts(scene_data.get("title", "Scene"), scenario_language),
                setting=expand_numbers_for_tts(scene_data.get("setting", ""), scenario_language),
                dialogue_lines=dialogue_lines,
                sound_cues=[]
            ))
        
        return {
            "title": expand_numbers_for_tts(data.get("title", title), scenario_language),
            "author": data.get("author", "Unknown"),
            "scenes": scenes,
            "characters": [],
            "total_duration_estimate": data.get("total_duration_estimate", 300.0)
        }

    def _clean_dialogue_text(self, text: str, character_name: str) -> str:
        """Підчищає авторські приписки у репліках і дикторських вставках."""
        original = (text or "").strip()
        if not original:
            return original
        cleaned = original.strip("«»\"")
        
        # Для персонажів прибираємо хвіст "..., сказал X / he said".
        if character_name != NARRATOR_NAME:
            trailing_patterns = [
                r"[,–—-]\s*(он|она|він|вона)\s+(сказал[аи]?|спросил[аи]?|ответил[аи]?|відповів|відповіла|запитав|запитала)\b.*$",
                r"[,–—-]\s*(he|she)\s+(said|asked|replied)\b.*$",
                r"[,–—-]\s*(сказал[аи]?|спросил[аи]?|ответил[аи]?|відповів|відповіла|запитав|запитала)\s+[A-ZА-ЯЁІЇЄҐ][\w-]*\b.*$",
            ]
            for pattern in trailing_patterns:
                cleaned = re.sub(pattern, "", cleaned, flags=re.IGNORECASE).strip()
        else:
            # Для Диктора прибираємо прямі згадки "хтось сказав/asked".
            narrator_noise_patterns = [
                r"\b(он|она|він|вона)\s+(сказал[аи]?|спросил[аи]?|ответил[аи]?|відповів|відповіла|запитав|запитала)\b",
                r"\b(he|she)\s+(said|asked|replied)\b",
                r"\b(сказал[аи]?|спросил[аи]?|ответил[аи]?|відповів|відповіла|запитав|запитала)\s+[A-ZА-ЯЁІЇЄҐ][\w-]*\b",
            ]
            for pattern in narrator_noise_patterns:
                cleaned = re.sub(pattern, "", cleaned, flags=re.IGNORECASE)
            cleaned = re.sub(r"\s{2,}", " ", cleaned).strip(" ,;:-")
            cleaned = ensure_terminal_punctuation(cleaned, "narrator")

        return cleaned or original
    
    def _extract_json(self, text: str) -> Optional[Dict]:
        """Витягує JSON з тексту відповіді."""
        # Спроба знайти JSON у markdown код блоці
        json_match = re.search(r'```(?:json)?\s*([\s\S]*?)\s*```', text)
        if json_match:
            try:
                return json.loads(json_match.group(1))
            except json.JSONDecodeError:
                pass
        
        # Спроба знайти JSON безпосередньо
        try:
            # Знаходимо першу { і останню }
            start = text.find('{')
            end = text.rfind('}') + 1
            if start != -1 and end > start:
                return json.loads(text[start:end])
        except json.JSONDecodeError:
            pass
        
        return None


class PexelsClient:
    """Клієнт для Pexels API - пошук та завантаження фонових звуків."""
    
    API_BASE_VIDEOS = "https://api.pexels.com/videos"
    
    def __init__(self, api_key: str):
        self.api_key = api_key
        self._session: Optional[aiohttp.ClientSession] = None
    
    async def _get_session(self) -> aiohttp.ClientSession:
        """Отримує або створює HTTP сесію."""
        if self._session is None or self._session.closed:
            self._session = aiohttp.ClientSession(
                headers={"Authorization": self.api_key}
            )
        return self._session
    
    async def close(self) -> None:
        """Закриває HTTP сесію."""
        if self._session and not self._session.closed:
            await self._session.close()
    
    async def search_sounds(self, query: str, duration_min: float = 10.0) -> List['SoundResult']:
        """Шукає відео/аудио матеріали на Pexels."""
        session = await self._get_session()
        
        # Pexels має відео API, використовуємо його для пошуку фонових сцен
        url = f"{self.API_BASE_VIDEOS}/search"
        params = {
            "query": query,
            "per_page": 10
        }
        
        results = []
        try:
            async with session.get(url, params=params) as response:
                if response.status == 200:
                    data = await response.json()
                    
                    for video in data.get("videos", []):
                        # Витягуємо аудіо з відео файлів
                        for video_file in video.get("video_files", []):
                            if video_file.get("file_type") in ["video/mp4", "video/webm"]:
                                results.append(SoundResult(
                                    id=str(video.get("id", "")),
                                    url=video_file.get("link", ""),
                                    description=video.get("url", ""),
                                    duration=float(video.get("duration", 0)),
                                    source="pexels"
                                ))
                                break  # Беремо тільки один файл з кожного відео
        except Exception as e:
            logger.error(f"[PexelsClient] Search failed: {e}")
        
        return results
    
    async def download_sound(self, url: str, output_path: str) -> Optional[str]:
        """Завантажує аудіо файл."""
        session = await self._get_session()
        
        try:
            async with session.get(url) as response:
                if response.status == 200:
                    content = await response.read()
                    
                    with open(output_path, 'wb') as f:
                        f.write(content)
                    
                    logger.info(f"[PexelsClient] Downloaded: {output_path}")
                    return output_path
        except Exception as e:
            logger.error(f"[PexelsClient] Download failed: {e}")
        
        return None
    
    async def get_video_audio(self, video_url: str, output_path: str) -> Optional[str]:
        """Витягує аудіо з відео файлу Pexels."""
        # Спочатку завантажуємо відео
        video_path = output_path.replace('.wav', '.mp4')
        
        downloaded = await self.download_sound(video_url, video_path)
        if not downloaded:
            return None
        
        # Витягуємо аудіо через ffmpeg/pydub
        try:
            from pydub import AudioSegment
            
            audio = AudioSegment.from_file(video_path)
            audio.export(output_path, format="wav")
            
            # Видаляємо тимчасовий відео файл
            os.remove(video_path)
            
            logger.info(f"[PexelsClient] Extracted audio: {output_path}")
            return output_path
        except Exception as e:
            logger.error(f"[PexelsClient] Audio extraction failed: {e}")
            return None


# =============================================================================
# SECTION 3: DATA STRUCTURES
# =============================================================================

class CharacterRole(Enum):
    """Ролі персонажів."""
    NARRATOR = "narrator"
    PROTAGONIST = "protagonist"
    SUPPORTING = "supporting"
    ANTAGONIST = "antagonist"
    MINOR = "minor"


class EmotionType(Enum):
    """Типи емоцій для реплік."""
    NEUTRAL = "neutral"
    HAPPY = "happy"
    SAD = "sad"
    ANGRY = "angry"
    EXCITED = "excited"
    CALM = "calm"
    DRAMATIC = "dramatic"
    WHISPER = "whisper"
    SHOUT = "shout"


@dataclass
class Character:
    """Персонаж аудіоп'єси."""
    name: str
    role: str = "supporting"  # narrator, protagonist, supporting, antagonist, minor
    voice_description: str = ""  # мінімум 20 слів
    gender: str = "male"
    age_range: str = "middle"  # child, young, middle, elderly
    personality_traits: List[str] = field(default_factory=list)
    seed: Optional[int] = None
    reference_audio_path: Optional[str] = None
    reference_text: Optional[str] = None
    voice_prompt: Optional[Dict] = None
    gender_locked: bool = False
    
    def __post_init__(self):
        # Валідація опису голосу
        if self.voice_description:
            word_count = len(self.voice_description.split())
            if word_count < 20:
                logger.warning(
                    f"[Character] Voice description for '{self.name}' has only {word_count} words "
                    f"(recommended: 20+)"
                )
        
        # Генеруємо seed якщо не задано
        if self.seed is None:
            self.seed = self._generate_seed()
    
    def _generate_seed(self) -> int:
        """Генерує seed на основі імені персонажа."""
        hash_obj = hashlib.md5(self.name.encode('utf-8'))
        return int(hash_obj.hexdigest(), 16) % (2**32)
    
    def to_dict(self) -> Dict:
        return {
            "name": self.name,
            "role": self.role,
            "voice_description": self.voice_description,
            "gender": self.gender,
            "age_range": self.age_range,
            "personality_traits": self.personality_traits,
            "seed": self.seed,
            "reference_audio_path": self.reference_audio_path,
            "reference_text": self.reference_text,
            "gender_locked": self.gender_locked
        }
    
    @classmethod
    def from_dict(cls, data: Dict) -> 'Character':
        return cls(
            name=data.get("name", "Unknown"),
            role=data.get("role", "supporting"),
            voice_description=data.get("voice_description", ""),
            gender=data.get("gender", "male"),
            age_range=data.get("age_range", "middle"),
            personality_traits=data.get("personality_traits", []),
            seed=data.get("seed"),
            reference_audio_path=data.get("reference_audio_path"),
            reference_text=data.get("reference_text"),
            gender_locked=bool(data.get("gender_locked", False))
        )


@dataclass
class DialogueLine:
    """Репліка персонажа."""
    character_name: str
    text: str
    emotion: str = "neutral"
    scene_id: int = 1
    order_in_scene: int = 1
    audio_path: Optional[str] = None
    duration: float = 0.0
    
    def to_dict(self) -> Dict:
        return {
            "character_name": self.character_name,
            "text": self.text,
            "emotion": self.emotion,
            "scene_id": self.scene_id,
            "order_in_scene": self.order_in_scene,
            "audio_path": self.audio_path,
            "duration": self.duration
        }


@dataclass
class SoundCue:
    """Опис фонового звуку."""
    scene_id: int
    description: str
    keywords: List[str] = field(default_factory=list)
    start_time: float = 0.0
    duration: float = 30.0
    volume: float = 0.3
    local_path: Optional[str] = None
    fade_in: float = 1.0
    fade_out: float = 1.0
    loop: bool = False
    
    def to_dict(self) -> Dict:
        return {
            "scene_id": self.scene_id,
            "description": self.description,
            "keywords": self.keywords,
            "start_time": self.start_time,
            "duration": self.duration,
            "volume": self.volume,
            "local_path": self.local_path,
            "fade_in": self.fade_in,
            "fade_out": self.fade_out,
            "loop": self.loop
        }


@dataclass
class Scene:
    """Сцена аудіоп'єси."""
    id: int
    title: str
    setting: str
    dialogue_lines: List[DialogueLine] = field(default_factory=list)
    sound_cues: List[SoundCue] = field(default_factory=list)
    duration_estimate: float = 0.0
    
    def to_dict(self) -> Dict:
        return {
            "id": self.id,
            "title": self.title,
            "setting": self.setting,
            "dialogue_lines": [dl.to_dict() for dl in self.dialogue_lines],
            "sound_cues": [sc.to_dict() for sc in self.sound_cues],
            "duration_estimate": self.duration_estimate
        }
    
    def get_total_duration(self) -> float:
        """Розраховує загальну тривалість сцени."""
        dialogue_duration = sum(dl.duration for dl in self.dialogue_lines)
        return max(dialogue_duration, self.duration_estimate)


@dataclass
class Scenario:
    """Повний сценарій аудіоп'єси."""
    title: str
    author: str = "Unknown"
    scenes: List[Scene] = field(default_factory=list)
    characters: List[Character] = field(default_factory=list)
    total_duration_estimate: float = 0.0
    created_at: str = field(default_factory=lambda: datetime.now().isoformat())
    
    def to_dict(self) -> Dict:
        return {
            "title": self.title,
            "author": self.author,
            "scenes": [s.to_dict() for s in self.scenes],
            "characters": [c.to_dict() for c in self.characters],
            "total_duration_estimate": self.total_duration_estimate,
            "created_at": self.created_at
        }
    
    @classmethod
    def from_dict(cls, data: Dict) -> 'Scenario':
        scenes = [Scene(**s) for s in data.get("scenes", [])]
        characters = [Character.from_dict(c) for c in data.get("characters", [])]
        
        return cls(
            title=data.get("title", "Untitled"),
            author=data.get("author", "Unknown"),
            scenes=scenes,
            characters=characters,
            total_duration_estimate=data.get("total_duration_estimate", 0.0),
            created_at=data.get("created_at", datetime.now().isoformat())
        )
    
    def get_character(self, name: str) -> Optional[Character]:
        """Знаходить персонажа за ім'ям."""
        if not name:
            return None
        candidate = canonical_character_name(name) or name
        if is_narrator_alias(candidate):
            candidate = NARRATOR_NAME
        for char in self.characters:
            if char.name.lower() == candidate.lower():
                return char
        for char in self.characters:
            if are_names_equivalent(candidate, char.name):
                return char
        return None
    
    def get_scene(self, scene_id: int) -> Optional[Scene]:
        """Знаходить сцену за ID."""
        for scene in self.scenes:
            if scene.id == scene_id:
                return scene
        return None


@dataclass
class SoundResult:
    """Результат пошуку звуку."""
    id: str
    url: str
    description: str
    duration: float
    source: str = "pexels"
    local_path: Optional[str] = None


@dataclass
class MixingConfig:
    """Конфігурація мікшування аудіо."""
    voice_volume: float = 1.0
    background_volume: float = 0.3
    voice_fade_in: float = 0.1
    voice_fade_out: float = 0.1
    background_fade_in: float = 1.0
    background_fade_out: float = 1.0
    pause_between_lines: float = 0.5
    speaker_change_pause: float = 0.12
    narrator_transition_pause: float = 0.22
    pause_between_scenes: float = 2.0
    chunk_fade_in: float = 0.02
    chunk_fade_out: float = 0.08
    tail_silence: float = 0.08
    narrator_tail_silence: float = 0.18
    sample_rate: int = 24000
    
    def to_dict(self) -> Dict:
        return asdict(self)


@dataclass
class Checkpoint:
    """Контрольна точка генерації."""
    timestamp: str
    duration: float
    audio_path: str
    scene_id: int
    message: str = ""
    
    def to_dict(self) -> Dict:
        return asdict(self)


@dataclass
class OrchestratorConfig:
    """Конфігурація головного координатора."""
    one_script_mode: bool = True
    drive_base_path: str = "/content/drive/MyDrive/audio_drama"
    cache_dir: str = "/content/cache/audio_drama"
    preview_voice_duration: float = 3.0
    test_voice_duration: float = 15.0
    checkpoint_1min: bool = True
    checkpoint_30min: bool = True
    auto_send_telegram: bool = True
    language: str = "russian"
    prefer_internal_tts: bool = True
    allow_external_tts_fallback: bool = False
    lazy_load_design: bool = True
    lazy_load_clone: bool = True
    unload_design_before_clone: bool = True
    quality_guard_enabled: bool = True
    quality_guard_max_attempts: int = 2
    quality_guard_similarity_threshold: float = 0.78
    quality_guard_min_tail_silence_ms: float = 120.0
    quality_guard_use_openai_asr: bool = False
    quality_guard_asr_model: str = "gpt-4o-mini-transcribe"
    quality_guard_use_local_asr: bool = True
    quality_guard_local_asr_model: str = "small"
    quality_guard_local_asr_device: str = "auto"
    quality_guard_local_asr_compute_type: str = "auto"
    strict_reference_mode_enabled: bool = True
    pre_voice_reference_stage_enabled: bool = True
    pre_voice_reference_timeout_sec: int = 0
    require_qwen_dual_models: bool = True
    qwen_precheck_on_init: bool = False
    qwen_auto_install_on_load: bool = True
    
    def to_dict(self) -> Dict:
        return asdict(self)


# =============================================================================
# SECTION 4: FILE PARSERS
# =============================================================================

class FileParser:
    """Парсинг різних форматів файлів."""
    
    SUPPORTED_EXTENSIONS = ['.txt', '.pdf', '.epub', '.fb2', '.doc', '.docx']
    
    @staticmethod
    def parse_txt(file_path: str) -> str:
        """Парсить TXT файл."""
        encodings = ['utf-8', 'windows-1251', 'cp1251', 'koi8-r', 'iso-8859-5']
        
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    content = f.read()
                logger.info(f"[FileParser] TXT parsed with encoding: {encoding}")
                return content
            except UnicodeDecodeError:
                continue
        
        raise ValueError(f"Could not decode TXT file: {file_path}")
    
    @staticmethod
    def parse_pdf(file_path: str) -> str:
        """Парсить PDF файл."""
        try:
            from pypdf import PdfReader
            
            reader = PdfReader(file_path)
            text_parts = []
            
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            
            content = "\n\n".join(text_parts)
            logger.info(f"[FileParser] PDF parsed: {len(reader.pages)} pages")
            return content
        except ImportError:
            logger.error("[FileParser] pypdf not installed")
            raise ImportError("Install pypdf: pip install pypdf")
        except Exception as e:
            logger.error(f"[FileParser] PDF parsing failed: {e}")
            raise
    
    @staticmethod
    def parse_epub(file_path: str) -> str:
        """Парсить EPUB файл."""
        try:
            from ebooklib import epub
            
            book = epub.read_epub(file_path)
            text_parts = []
            
            for item in book.get_items():
                if item.get_type() == 9:  # ITEM_DOCUMENT
                    content = item.get_content()
                    # Видаляємо HTML теги
                    text = re.sub(r'<[^>]+>', '', content.decode('utf-8', errors='ignore'))
                    text_parts.append(text)
            
            content = "\n\n".join(text_parts)
            logger.info(f"[FileParser] EPUB parsed: {len(text_parts)} chapters")
            return content
        except ImportError:
            logger.error("[FileParser] ebooklib not installed")
            raise ImportError("Install ebooklib: pip install ebooklib")
        except Exception as e:
            logger.error(f"[FileParser] EPUB parsing failed: {e}")
            raise
    
    @staticmethod
    def parse_fb2(file_path: str) -> str:
        """Парсить FB2 файл."""
        try:
            # FB2 це XML формат
            import xml.etree.ElementTree as ET
            
            tree = ET.parse(file_path)
            root = tree.getroot()
            
            # FB2 namespace
            ns = {'fb2': 'http://www.gribuser.ru/xml/fictionbook/2.0'}
            
            text_parts = []
            
            # Знаходимо всі параграфи
            for p in root.findall('.//fb2:p', ns):
                if p.text:
                    text_parts.append(p.text)
            
            content = "\n".join(text_parts)
            logger.info(f"[FileParser] FB2 parsed: {len(text_parts)} paragraphs")
            return content
        except Exception as e:
            logger.error(f"[FileParser] FB2 parsing failed: {e}")
            # Fallback - читаємо як звичайний текст
            return FileParser.parse_txt(file_path)
    
    @staticmethod
    def parse_docx(file_path: str) -> str:
        """Парсить DOCX файл."""
        try:
            from docx import Document
            
            doc = Document(file_path)
            text_parts = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text_parts.append(paragraph.text)
            
            content = "\n".join(text_parts)
            logger.info(f"[FileParser] DOCX parsed: {len(doc.paragraphs)} paragraphs")
            return content
        except ImportError:
            logger.error("[FileParser] python-docx not installed")
            raise ImportError("Install python-docx: pip install python-docx")
        except Exception as e:
            logger.error(f"[FileParser] DOCX parsing failed: {e}")
            raise
    
    @staticmethod
    def parse_doc(file_path: str) -> str:
        """Парсить DOC файл (старий формат)."""
        try:
            # Спроба через pypandoc
            import pypandoc
            
            content = pypandoc.convert_file(file_path, 'plain')
            logger.info(f"[FileParser] DOC parsed via pypandoc")
            return content
        except ImportError:
            logger.warning("[FileParser] pypandoc not installed, trying alternative")
        except Exception as e:
            logger.warning(f"[FileParser] pypandoc failed: {e}")
        
        # Fallback - спроба як docx
        try:
            return FileParser.parse_docx(file_path)
        except Exception:
            pass
        
        # Останній варіант - як текст
        return FileParser.parse_txt(file_path)
    
    @staticmethod
    def auto_detect_and_parse(file_path: str) -> str:
        """Автоматично визначає формат та парсить файл."""
        ext = Path(file_path).suffix.lower()
        
        if ext == '.txt':
            return FileParser.parse_txt(file_path)
        elif ext == '.pdf':
            return FileParser.parse_pdf(file_path)
        elif ext == '.epub':
            return FileParser.parse_epub(file_path)
        elif ext == '.fb2':
            return FileParser.parse_fb2(file_path)
        elif ext == '.docx':
            return FileParser.parse_docx(file_path)
        elif ext == '.doc':
            return FileParser.parse_doc(file_path)
        else:
            # Спроба як текст
            logger.warning(f"[FileParser] Unknown extension {ext}, trying as text")
            return FileParser.parse_txt(file_path)
    
    @staticmethod
    def is_supported(file_path: str) -> bool:
        """Перевіряє чи підтримується формат файлу."""
        ext = Path(file_path).suffix.lower()
        return ext in FileParser.SUPPORTED_EXTENSIONS


# =============================================================================
# SECTION 4.5: PRONUNCIATION MANAGEMENT
# =============================================================================

class PronunciationManager:
    """Словники вимови/наголосів для різних мов."""
    
    DEFAULT_RULES: Dict[str, Dict[str, str]] = {
        "russian": {
            "ежик": "ёжик",
            "елка": "ёлка",
            "ёлка": "ёлка",
            "все": "всё",
            "ее": "её",
            "еще": "ещё",
            "все-таки": "всё-таки",
            "трех": "трёх",
            "четырех": "четырёх",
            "берет": "берёт",
            "руки": "рУки",
            "воровство": "воровствО",
            "нубу": "нУбу",
            "нуба": "нУба",
        },
        "ukrainian": {
            "пєса": "п'єса",
            "обєкт": "об'єкт",
            "інтервю": "інтерв'ю",
            "медіа контент": "медіаконтент",
        },
        "english": {},
    }
    
    def __init__(self, storage_path: str):
        self.storage_path = storage_path
        self._rules: Dict[str, Dict[str, str]] = {
            lang: dict(values) for lang, values in self.DEFAULT_RULES.items()
        }
        self.load()
    
    def load(self) -> None:
        """Завантажує словники з файла."""
        if not os.path.exists(self.storage_path):
            self.save()
            return
        try:
            with open(self.storage_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            languages = data.get("languages", {})
            for lang in SUPPORTED_PRONUNCIATION_LANGUAGES:
                custom = languages.get(lang, {})
                merged = dict(self.DEFAULT_RULES.get(lang, {}))
                for source, target in custom.items():
                    if source and target:
                        merged[str(source)] = str(target)
                self._rules[lang] = merged
        except Exception as e:
            logger.warning(f"[Pronunciation] Failed to load dictionary: {e}")
    
    def save(self) -> None:
        """Зберігає словники у файл."""
        try:
            os.makedirs(os.path.dirname(self.storage_path), exist_ok=True)
            payload = {
                "version": 1,
                "updated_at": datetime.now().isoformat(),
                "languages": self._rules,
            }
            with open(self.storage_path, "w", encoding="utf-8") as f:
                json.dump(payload, f, ensure_ascii=False, indent=2)
        except Exception as e:
            logger.warning(f"[Pronunciation] Failed to save dictionary: {e}")
    
    def list_rules(self, language: str) -> Dict[str, str]:
        """Повертає правила для мови."""
        lang = normalize_language_code(language, "russian")
        return dict(self._rules.get(lang, {}))
    
    def upsert_rule(self, language: str, source: str, target: str) -> bool:
        """Додає або оновлює правило."""
        lang = normalize_language_code(language, "russian")
        source = (source or "").strip()
        target = (target or "").strip()
        if not source or not target:
            return False
        if lang not in self._rules:
            self._rules[lang] = {}
        self._rules[lang][source] = target
        self.save()
        return True
    
    def remove_rule(self, language: str, source: str) -> bool:
        """Видаляє правило."""
        lang = normalize_language_code(language, "russian")
        source = (source or "").strip()
        if not source:
            return False
        if source not in self._rules.get(lang, {}):
            return False
        self._rules[lang].pop(source, None)
        self.save()
        return True
    
    def apply(self, text: str, language: str) -> str:
        """Застосовує правила вимови до тексту."""
        if not text:
            return text
        lang = normalize_language_code(language, "russian")
        result = text
        rules = self._rules.get(lang, {})
        for source, target in sorted(rules.items(), key=lambda it: len(it[0]), reverse=True):
            if not source or not target:
                continue
            pattern = rf"(?<![\w']){re.escape(source)}(?![\w'])"
            result = re.sub(
                pattern,
                lambda match: self._match_case(match.group(0), target),
                result,
                flags=re.IGNORECASE
            )
        
        if lang == "russian":
            result = self._auto_restore_yo(result)
        
        return result
    
    def _auto_restore_yo(self, text: str) -> str:
        """Автовідновлення частини форм з ё."""
        result = text
        for source, target in RUSSIAN_YO_HINTS.items():
            pattern = rf"(?<![\w']){re.escape(source)}(?![\w'])"
            result = re.sub(
                pattern,
                lambda match: self._match_case(match.group(0), target),
                result,
                flags=re.IGNORECASE
            )
        return result
    
    @staticmethod
    def _match_case(original: str, replacement: str) -> str:
        """Наближено зберігає регістр слова при заміні."""
        if original.isupper():
            return replacement.upper()
        if len(original) > 1 and original[0].isupper() and original[1:].islower():
            return replacement[0].upper() + replacement[1:]
        return replacement


# =============================================================================
# SECTION 4.6: AUDIO QUALITY GUARD
# =============================================================================

@dataclass
class DialogueQualityIssue:
    """Окрема проблема якості озвучки."""
    code: str
    severity: str = "critical"
    message: str = ""
    expected: str = ""
    actual: str = ""
    suggestion: str = ""

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


@dataclass
class DialogueQualityReport:
    """Звіт перевірки якості репліки."""
    passed: bool
    attempt: int
    language: str
    asr_provider: str = "none"
    transcript: str = ""
    similarity: float = 1.0
    tail_dbfs: float = -120.0
    trailing_silence_ms: float = 0.0
    issues: List[DialogueQualityIssue] = field(default_factory=list)
    corrected_text: str = ""

    def to_dict(self) -> Dict[str, Any]:
        payload = asdict(self)
        payload["issues"] = [issue.to_dict() for issue in self.issues]
        return payload


class AudioQualityGuard:
    """Перевіряє якість синтезу й пропонує автокорекції перед retry."""

    def __init__(
        self,
        enabled: bool = True,
        max_attempts: int = 2,
        similarity_threshold: float = 0.78,
        min_tail_silence_ms: float = 120.0,
        use_openai_asr: bool = False,
        openai_asr_model: str = "gpt-4o-mini-transcribe",
        openai_api_key: Optional[str] = None,
        use_local_asr: bool = True,
        local_asr_model: str = "small",
        local_asr_device: str = "auto",
        local_asr_compute_type: str = "auto"
    ):
        self.enabled = bool(enabled)
        self.max_attempts = max(1, int(max_attempts or 1))
        self.similarity_threshold = float(similarity_threshold or 0.78)
        self.min_tail_silence_ms = float(min_tail_silence_ms or 120.0)
        self.use_openai_asr = bool(use_openai_asr)
        self.openai_asr_model = (openai_asr_model or "gpt-4o-mini-transcribe").strip()
        self.openai_api_key = (openai_api_key or os.getenv("OPENAI_API_KEY", "")).strip()
        self.use_local_asr = bool(use_local_asr)
        self.local_asr_model = (local_asr_model or "small").strip() or "small"
        self.local_asr_device = (local_asr_device or "auto").strip().lower()
        self.local_asr_compute_type = (local_asr_compute_type or "auto").strip().lower()
        self._openai_client = None
        self._openai_failed = False
        self._warned_missing_openai = False
        self._faster_whisper_model = None
        self._faster_whisper_failed = False
        self._warned_missing_local_asr = False

    def _normalize_for_similarity(self, text: str) -> str:
        value = (text or "").lower()
        value = value.replace("ё", "е").replace("́", "")
        value = re.sub(r"[^a-zа-яіїєґ0-9\s]", " ", value)
        value = re.sub(r"\s+", " ", value).strip()
        return value

    def _get_openai_client(self):
        if not self.use_openai_asr or self._openai_failed:
            return None
        if not self.openai_api_key:
            if not self._warned_missing_openai:
                logger.info("[QualityGuard] OPENAI_API_KEY not set, ASR checks are skipped")
                self._warned_missing_openai = True
            return None
        if OpenAI is None:
            if not self._warned_missing_openai:
                logger.warning("[QualityGuard] openai package missing, ASR checks are skipped")
                self._warned_missing_openai = True
            return None
        if self._openai_client is None:
            try:
                self._openai_client = OpenAI(api_key=self.openai_api_key)
            except Exception as e:
                logger.warning(f"[QualityGuard] Failed to initialize OpenAI client: {e}")
                self._openai_failed = True
                return None
        return self._openai_client

    def _resolve_local_asr_runtime(self) -> Tuple[str, str]:
        """Повертає (device, compute_type) для faster-whisper."""
        device = self.local_asr_device if self.local_asr_device in {"auto", "cpu", "cuda"} else "auto"
        compute_type = self.local_asr_compute_type if self.local_asr_compute_type else "auto"

        if device == "auto":
            has_cuda = False
            try:
                import torch
                has_cuda = bool(torch.cuda.is_available())
            except Exception:
                has_cuda = False
            device = "cuda" if has_cuda else "cpu"

        if compute_type == "auto":
            compute_type = "int8_float16" if device == "cuda" else "int8"

        return device, compute_type

    def _get_local_asr_model(self):
        if not self.use_local_asr or self._faster_whisper_failed:
            return None
        if FasterWhisperModel is None:
            if not self._warned_missing_local_asr:
                logger.warning("[QualityGuard] faster-whisper not installed, local ASR disabled")
                self._warned_missing_local_asr = True
            self._faster_whisper_failed = True
            return None

        if self._faster_whisper_model is not None:
            return self._faster_whisper_model

        device, compute_type = self._resolve_local_asr_runtime()
        compute_candidates: List[str] = [compute_type]
        if device == "cuda":
            compute_candidates.extend(["float16", "int8_float16", "int8"])
        else:
            compute_candidates.extend(["int8", "float32"])

        # Унікалізуємо порядок кандидатів.
        unique_candidates: List[str] = []
        for candidate in compute_candidates:
            token = (candidate or "").strip().lower()
            if token and token not in unique_candidates:
                unique_candidates.append(token)

        last_error: Optional[Exception] = None
        for candidate in unique_candidates:
            try:
                self._faster_whisper_model = FasterWhisperModel(
                    self.local_asr_model,
                    device=device,
                    compute_type=candidate,
                    cpu_threads=max(1, min(8, (os.cpu_count() or 4)))
                )
                logger.info(
                    f"[QualityGuard] Local ASR ready: faster-whisper model={self.local_asr_model}, "
                    f"device={device}, compute_type={candidate}"
                )
                return self._faster_whisper_model
            except Exception as e:
                last_error = e
                continue

        self._faster_whisper_failed = True
        logger.warning(
            f"[QualityGuard] Failed to initialize local ASR model '{self.local_asr_model}' "
            f"(device={device}, compute={compute_type}): {last_error}"
        )
        return None

    def _transcribe_with_local_asr(self, audio_path: str, language: str) -> Tuple[str, str]:
        model = self._get_local_asr_model()
        if model is None:
            return "", "local_asr_unavailable"
        if not audio_path or not os.path.exists(audio_path):
            return "", "none"

        lang = normalize_language_code(language, "russian")
        lang_code = language_short_code(lang)
        if lang_code not in {"ru", "uk", "en"}:
            lang_code = None

        try:
            segments, _ = model.transcribe(
                audio_path,
                task="transcribe",
                language=lang_code,
                beam_size=1,
                best_of=1,
                temperature=0.0,
                vad_filter=True,
                condition_on_previous_text=False
            )
            parts: List[str] = []
            for segment in segments:
                text = (getattr(segment, "text", "") or "").strip()
                if text:
                    parts.append(text)
            transcript = " ".join(parts).strip()
            if transcript:
                return transcript, "faster_whisper"
            return "", "faster_whisper_empty"
        except Exception as e:
            logger.warning(f"[QualityGuard] Local ASR failed: {e}")
            return "", "faster_whisper_error"

    def _transcribe_with_openai(self, audio_path: str, language: str) -> Tuple[str, str]:
        client = self._get_openai_client()
        if client is None:
            return "", "openai_unavailable"
        if not audio_path or not os.path.exists(audio_path):
            return "", "none"

        lang = normalize_language_code(language, "russian")
        lang_code = language_short_code(lang)
        try:
            with open(audio_path, "rb") as f:
                response = client.audio.transcriptions.create(
                    model=self.openai_asr_model,
                    file=f,
                    response_format="text",
                    language=lang_code
                )
            if isinstance(response, str):
                return response.strip(), "openai"
            text = getattr(response, "text", "") or str(response or "")
            return text.strip(), "openai"
        except Exception as e:
            logger.warning(f"[QualityGuard] OpenAI ASR failed: {e}")
            return "", "openai_error"

    def _transcribe(self, audio_path: str, language: str) -> Tuple[str, str]:
        if not audio_path or not os.path.exists(audio_path):
            return "", "none"

        lang = normalize_language_code(language, "russian")
        last_provider = "none"

        if self.use_local_asr:
            transcript, provider = self._transcribe_with_local_asr(audio_path, lang)
            last_provider = provider or last_provider
            if transcript:
                return transcript, provider

        if self.use_openai_asr:
            transcript, provider = self._transcribe_with_openai(audio_path, lang)
            last_provider = provider or last_provider
            if transcript:
                return transcript, provider

        return "", last_provider

    def _analyze_audio_tail(self, audio_path: str) -> Tuple[float, float]:
        """Повертає (tail_dbfs, trailing_silence_ms)."""
        if not audio_path or not os.path.exists(audio_path):
            return -120.0, 0.0
        try:
            from scipy.io import wavfile
            sr, data = wavfile.read(audio_path)
            if sr <= 0:
                return -120.0, 0.0

            arr = np.array(data, dtype=np.float32)
            if arr.ndim > 1:
                arr = np.mean(arr, axis=1)
            if arr.size == 0:
                return -120.0, 0.0

            peak = float(np.max(np.abs(arr)))
            if peak > 0:
                arr = arr / peak

            tail_window = max(1, int(sr * 0.08))
            tail = arr[-tail_window:]
            tail_rms = float(np.sqrt(np.mean(np.square(tail)))) if tail.size else 0.0
            tail_dbfs = float(20.0 * np.log10(max(tail_rms, 1e-6)))

            silence_threshold = 0.01
            idx = arr.size - 1
            while idx >= 0 and abs(arr[idx]) < silence_threshold:
                idx -= 1
            trailing_samples = max(0, (arr.size - 1) - idx)
            trailing_silence_ms = (trailing_samples * 1000.0) / float(sr)
            return tail_dbfs, trailing_silence_ms
        except Exception as e:
            logger.warning(f"[QualityGuard] Tail analysis failed for {audio_path}: {e}")
            return -120.0, 0.0

    def _detect_yo_mismatches(self, expected_text: str, transcript: str, language: str) -> List[Tuple[str, str]]:
        lang = normalize_language_code(language, "russian")
        if lang != "russian" or not expected_text or not transcript:
            return []

        expected_words = {
            word.lower()
            for word in re.findall(r"[а-яё]+", expected_text.lower())
            if "ё" in word
        }
        transcript_norm = transcript.lower().replace("ё", "е")

        mismatches: List[Tuple[str, str]] = []
        for yo_word in sorted(expected_words):
            e_word = yo_word.replace("ё", "е")
            if e_word == yo_word:
                continue
            pattern = rf"(?<![а-яё]){re.escape(e_word)}(?![а-яё])"
            if re.search(pattern, transcript_norm):
                mismatches.append((e_word, yo_word))
        return mismatches

    def build_retry_text(self, original_text: str, language: str, issues: List[DialogueQualityIssue]) -> str:
        """Формує текст для retry на базі виявлених проблем."""
        if not original_text:
            return original_text

        lang = normalize_language_code(language, "russian")
        result = str(original_text)
        codes = {issue.code for issue in issues}

        for issue in issues:
            if issue.code != "yo_mismatch":
                continue
            source = (issue.actual or "").strip()
            target = (issue.expected or "").strip()
            if not source or not target:
                continue
            pattern = rf"(?<![\w']){re.escape(source)}(?![\w'])"
            result = re.sub(
                pattern,
                lambda match: PronunciationManager._match_case(match.group(0), target),
                result,
                flags=re.IGNORECASE
            )

        if re.search(r"\d", result) or "numbers" in codes:
            result = expand_numbers_for_tts(result, lang)

        result = apply_inline_stress_markers(result, lang)

        if "abrupt_ending" in codes:
            trimmed = result.rstrip()
            if trimmed.endswith((".", "!", "?")):
                result = f"{trimmed} .."
            elif not trimmed.endswith(("..", "...", "…")):
                result = f"{trimmed} ..."
            else:
                result = trimmed

        return result

    def learn_pronunciation_rules(
        self,
        pronunciation_manager: Optional[PronunciationManager],
        language: str,
        issues: List[DialogueQualityIssue]
    ) -> int:
        """Навчає словник вимови на стабільних yo-замінниках."""
        if pronunciation_manager is None:
            return 0

        learned = 0
        for issue in issues:
            if issue.code != "yo_mismatch":
                continue
            source = (issue.actual or "").strip()
            target = (issue.expected or "").strip()
            if not source or not target:
                continue
            if pronunciation_manager.upsert_rule(language, source, target):
                learned += 1
        return learned

    def analyze_dialogue(
        self,
        audio_path: str,
        expected_text: str,
        language: str,
        attempt: int = 1
    ) -> DialogueQualityReport:
        """Запускає комплексну перевірку озвученої репліки."""
        lang = normalize_language_code(language, "russian")
        transcript, provider = self._transcribe(audio_path, lang) if self.enabled else ("", "none")
        tail_dbfs, trailing_silence_ms = self._analyze_audio_tail(audio_path)

        issues: List[DialogueQualityIssue] = []
        similarity = 1.0

        if transcript:
            expected_norm = self._normalize_for_similarity(expected_text)
            transcript_norm = self._normalize_for_similarity(transcript)
            if expected_norm and transcript_norm:
                similarity = difflib.SequenceMatcher(None, expected_norm, transcript_norm).ratio()
            if similarity < self.similarity_threshold:
                issues.append(DialogueQualityIssue(
                    code="asr_mismatch",
                    message=f"Низька схожість ASR та очікуваного тексту ({similarity:.2f})",
                    expected=expected_text[:220],
                    actual=transcript[:220],
                    suggestion="Перегенерувати репліку з агресивнішою нормалізацією"
                ))

            if re.search(r"\d", expected_text) and similarity < 0.95:
                issues.append(DialogueQualityIssue(
                    code="numbers",
                    message="Ймовірно некоректна вимова чисел",
                    expected=expected_text[:220],
                    actual=transcript[:220],
                    suggestion="Розгорнути числа словами перед TTS"
                ))

            yo_mismatches = self._detect_yo_mismatches(expected_text, transcript, lang)
            for source, target in yo_mismatches[:8]:
                issues.append(DialogueQualityIssue(
                    code="yo_mismatch",
                    message=f"Ймовірна заміна 'ё' -> 'е': {source} -> {target}",
                    expected=target,
                    actual=source,
                    suggestion=f"Додати правило вимови {source}={target}"
                ))

        expected_tail = (expected_text or "").strip().endswith((".", "!", "?"))
        if expected_tail and trailing_silence_ms < self.min_tail_silence_ms and tail_dbfs > -33.0:
            issues.append(DialogueQualityIssue(
                code="abrupt_ending",
                message=(
                    "Ймовірно різкий обрив на кінці фрази "
                    f"(tail={tail_dbfs:.1f} dBFS, silence={trailing_silence_ms:.1f} ms)"
                ),
                expected=f">={self.min_tail_silence_ms:.0f} ms тиші",
                actual=f"{trailing_silence_ms:.1f} ms",
                suggestion="Додати паузу в кінець репліки та перегенерувати"
            ))

        corrected_text = self.build_retry_text(expected_text, lang, issues)
        return DialogueQualityReport(
            passed=(len(issues) == 0),
            attempt=max(1, int(attempt or 1)),
            language=lang,
            asr_provider=provider,
            transcript=transcript,
            similarity=float(similarity),
            tail_dbfs=float(tail_dbfs),
            trailing_silence_ms=float(trailing_silence_ms),
            issues=issues,
            corrected_text=corrected_text
        )


# =============================================================================
# SECTION 5: VOICE MANAGEMENT
# =============================================================================

class VoiceManager:
    """Управління голосами персонажів."""
    
    def __init__(
        self,
        design_engine=None,
        clone_engine=None,
        library_manager=None,
        cache=None,
        create_from_reference_fn=None,
        text_preprocessor: Optional[Callable[[str, str], str]] = None
    ):
        self._design_engine = design_engine
        self._clone_engine = clone_engine
        self._library_manager = library_manager
        self._cache = cache
        self._create_from_reference_fn = create_from_reference_fn
        self._text_preprocessor = text_preprocessor
        self._voice_registry: Dict[str, Character] = {}
        self._test_voices: Dict[str, List[str]] = {}  # character_name -> list of test audio paths
    
    def set_engines(
        self,
        design_engine,
        clone_engine,
        library_manager,
        cache=None,
        create_from_reference_fn=None,
        text_preprocessor: Optional[Callable[[str, str], str]] = None
    ) -> None:
        """Встановлює TTS двигуни."""
        self._design_engine = design_engine
        self._clone_engine = clone_engine
        self._library_manager = library_manager
        self._cache = cache
        if create_from_reference_fn is not None:
            self._create_from_reference_fn = create_from_reference_fn
        if text_preprocessor is not None:
            self._text_preprocessor = text_preprocessor
        logger.info("[VoiceManager] Engines set")
    
    async def create_test_voice(
        self,
        character: Character,
        duration_sec: float = 15.0,
        output_dir: str = "/content/test_voices",
        language: str = "russian"
    ) -> str:
        """Створює тестовий голос для персонажа (15+ сек)."""
        if self._design_engine is None:
            raise RuntimeError("Design engine not initialized")
        
        os.makedirs(output_dir, exist_ok=True)
        
        # Генеруємо текст для тестового голосу
        tts_language = normalize_language_code(language, "russian")
        test_text = self._generate_test_text(character, duration_sec, tts_language)
        if self._text_preprocessor:
            try:
                test_text = self._text_preprocessor(test_text, tts_language)
            except Exception as prep_err:
                logger.warning(f"[VoiceManager] Text preprocessing failed for {character.name}: {prep_err}")
        character.reference_text = test_text
        desc_gender = extract_gender_from_text(character.voice_description)
        if character.gender_locked:
            character.gender = normalize_gender(character.gender, "male")
        else:
            character.gender = resolve_character_gender(character.name, character.role, desc_gender or character.gender)
        
        # Формуємо voice config
        voice_config = {
            "seed": character.seed,
            "prompt": build_gender_aware_prompt(character),
            "speed": 1.0,
            "gender": character.gender
        }
        
        try:
            logger.info(f"[VoiceManager] Creating test voice for '{character.name}'...")
            
            audio, sr = await asyncio.to_thread(
                self._design_engine.generate,
                test_text,
                voice_config,
                tts_language
            )
            
            if audio is None:
                raise RuntimeError("Audio generation returned None")
            
            # Зберігаємо тестовий голос
            safe_name = character.name.replace(' ', '_').lower()
            version = len(self._test_voices.get(character.name, [])) + 1
            output_path = os.path.join(output_dir, f"{safe_name}_v{version}.wav")
            
            self._save_audio(audio, sr, output_path)
            if not os.path.exists(output_path) or os.path.getsize(output_path) == 0:
                raise RuntimeError(f"Voice file was not saved correctly: {output_path}")
            
            # Реєструємо тестовий голос
            if character.name not in self._test_voices:
                self._test_voices[character.name] = []
            self._test_voices[character.name].append(output_path)
            
            # Оновлюємо reference для персонажа
            character.reference_audio_path = output_path
            character.voice_prompt = await self._build_voice_prompt_from_reference(character)
            
            logger.info(f"[VoiceManager] Test voice created: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"[VoiceManager] Test voice creation failed: {e}")
            raise
    
    async def regenerate_voice(
        self,
        character_name: str,
        modification: str = "",
        gender_override: Optional[str] = None,
        output_dir: str = "/content/test_voices",
        language: str = "russian",
        duration_sec: Optional[float] = None
    ) -> str:
        """Перегенерує голос персонажа з модифікацією."""
        _, character = self._resolve_character(character_name)
        if character is None:
            raise ValueError(f"Character '{character_name}' not found")
        
        if gender_override:
            character.gender = normalize_gender(gender_override, character.gender or "male")
            character.gender_locked = True
        
        # Додаємо модифікацію до опису голосу
        if modification:
            original_description = character.voice_description
            character.voice_description = f"{original_description}. {modification}".strip(". ")
        
        # Створюємо новий тестовий голос
        target_duration = float(duration_sec) if duration_sec and duration_sec > 0 else 15.0
        return await self.create_test_voice(
            character,
            duration_sec=target_duration,
            output_dir=output_dir,
            language=language
        )
    
    async def get_voice_reference(self, character_name: str, language: str = "russian") -> Optional[Dict]:
        """Отримує voice reference для персонажа."""
        _, character = self._resolve_character(character_name)
        if character is None:
            return None
        
        if character.voice_prompt:
            return character.voice_prompt
        
        if character.reference_audio_path:
            character.voice_prompt = await self._build_voice_prompt_from_reference(character)
            if character.voice_prompt:
                return character.voice_prompt
            logger.warning(
                f"[VoiceManager] Character '{character.name}' has reference file but prompt build failed. "
                "Skip library fallback to avoid voice mix."
            )
            return None

        if self._library_manager:
            if character.role == "narrator" or is_narrator_alias(character.name):
                logger.info(
                    f"[VoiceManager] Skip narrator library fallback for '{character.name}' "
                    "to prevent accidental voice drift. Use explicit reference file."
                )
                return None
            # Fallback на бібліотеку, якщо локального reference ще немає
            voice_prompt, _ = await asyncio.to_thread(
                self._library_manager.get_voice,
                character.name,
                "custom",
                character.gender,
                normalize_language_code(language, "russian")
            )
            if voice_prompt:
                character.voice_prompt = voice_prompt
                return voice_prompt
        
        return None
    
    def register_character(self, character: Character) -> None:
        """Реєструє персонажа в системі."""
        character.name = canonical_character_name(character.name) or character.name
        if character.role == "narrator" or is_narrator_alias(character.name):
            character.name = NARRATOR_NAME
            character.role = "narrator"
        desc_gender = extract_gender_from_text(character.voice_description)
        if character.gender_locked:
            character.gender = normalize_gender(character.gender, "male")
        else:
            character.gender = resolve_character_gender(character.name, character.role, desc_gender or character.gender)
        self._voice_registry[character.name] = character
        logger.info(f"[VoiceManager] Registered character: {character.name}")
    
    def get_character(self, name: str) -> Optional[Character]:
        """Отримує персонажа за ім'ям."""
        _, character = self._resolve_character(name)
        return character
    
    def list_characters(self) -> List[Character]:
        """Повертає список всіх персонажів."""
        return list(self._voice_registry.values())
    
    def save_voice_to_library(self, character: Character, audio_path: str) -> bool:
        """Зберігає голос у бібліотеку."""
        if self._create_from_reference_fn is None:
            logger.warning("[VoiceManager] Library manager not available")
            return False
        
        try:
            guessed_language = detect_text_language(
                (character.reference_text or character.voice_description or ""),
                "russian"
            )
            ref_text = character.reference_text or self._generate_test_text(character, 15.0, guessed_language)
            reference_language = detect_text_language(ref_text, "russian")
            voice_prompt, error_message = self._create_from_reference_fn(
                character_name=character.name,
                reference_audio_path=audio_path,
                reference_text=ref_text,
                voice_preset="custom",
                gender=character.gender,
                language=reference_language
            )
            if voice_prompt:
                character.voice_prompt = voice_prompt
            if error_message:
                logger.warning(f"[VoiceManager] Voice library save warning for {character.name}: {error_message}")
            logger.info(f"[VoiceManager] Voice saved to library: {character.name}")
            return voice_prompt is not None
        except Exception as e:
            logger.error(f"[VoiceManager] Failed to save voice: {e}")
            return False

    def _resolve_character(self, requested_name: str) -> Tuple[Optional[str], Optional[Character]]:
        """Шукає персонажа без урахування регістру і пробілів."""
        if not requested_name:
            return None, None
        requested_name = canonical_character_name(requested_name) or requested_name
        normalized = requested_name.strip().lower()
        if is_narrator_alias(normalized):
            normalized = NARRATOR_NAME.lower()
        for name, character in self._voice_registry.items():
            if name.strip().lower() == normalized:
                return name, character
        for name, character in self._voice_registry.items():
            if are_names_equivalent(requested_name, name):
                return name, character
        return None, None

    async def _build_voice_prompt_from_reference(self, character: Character) -> Optional[Dict]:
        """Будує voice_prompt з останнього затвердженого reference-а персонажа."""
        if not character.reference_audio_path or not os.path.exists(character.reference_audio_path):
            return None
        if self._clone_engine is None:
            return None
        
        guessed_language = detect_text_language(
            (character.reference_text or character.voice_description or ""),
            "russian"
        )
        reference_text = character.reference_text or self._generate_test_text(character, 15.0, guessed_language)
        
        try:
            voice_prompt = None
            if hasattr(self._clone_engine, "create_voice_prompt"):
                voice_prompt, _ = await asyncio.to_thread(
                    self._clone_engine.create_voice_prompt,
                    character.reference_audio_path,
                    reference_text,
                    True,
                    True
                )
            
            if voice_prompt:
                await asyncio.to_thread(
                    self.save_voice_to_library,
                    character,
                    character.reference_audio_path
                )
                return voice_prompt
            logger.warning(
                f"[VoiceManager] Failed to create voice_prompt from direct reference for {character.name}. "
                "Library fallback disabled to avoid voice mix."
            )
        except Exception as e:
            logger.warning(f"[VoiceManager] Failed to build voice reference for {character.name}: {e}")
        
        return None
    
    def _generate_test_text(
        self,
        character: Character,
        duration_sec: float,
        language: str = "russian"
    ) -> str:
        """Генерує текст для тестового голосу мовою поточного проєкту."""
        word_count = int(duration_sec * 2.5)
        lang = normalize_language_code(language, "russian")
        gender = normalize_gender(character.gender, "male")
        
        if lang == "ukrainian":
            narrator_intro = "Я ваш диктор" if gender == "male" else "Я ваша дикторка" if gender == "female" else "Я ваш оповідач"
            gender_assertion = (
                "Мій голос жіночий, м'який і впевнений." if gender == "female"
                else "Мій голос нейтральний, рівний і спокійний." if gender == "neutral"
                else "Мій голос чоловічий, упевнений і рівний."
            )
            base_texts = {
                "narrator": f"{gender_assertion} Ласкаво просимо у світ аудіоп'єси. {narrator_intro}, і сьогодні я поведу вас цією історією.",
                "protagonist": f"{gender_assertion} Вітаю, мене звати {{name}}, і я головний герой цієї історії. Я готовий провести вас крізь усі події.",
                "supporting": f"{gender_assertion} Доброго дня, я {{name}}, один із персонажів цієї історії. Мій голос має виразні інтонації та чітку дикцію.",
                "antagonist": f"{gender_assertion} Що ж, я {{name}}, і я тут, щоб змінити перебіг подій. Мій голос звучить напружено й виразно.",
                "default": f"{gender_assertion} Це тестовий голос для персонажа на ім'я {{name}}. Голос має бути стабільним і впізнаваним."
            }
        elif lang == "english":
            narrator_intro = "I am your narrator" if gender == "male" else "I am your female narrator" if gender == "female" else "I am your narrator"
            gender_assertion = (
                "My voice is female, clear, and steady." if gender == "female"
                else "My voice is neutral, balanced, and calm." if gender == "neutral"
                else "My voice is male, clear, and steady."
            )
            base_texts = {
                "narrator": f"{gender_assertion} Welcome to this audio drama. {narrator_intro}, and I will guide you through the story.",
                "protagonist": f"{gender_assertion} Hello, my name is {{name}}, and I am the protagonist of this story. I am ready for every challenge ahead.",
                "supporting": f"{gender_assertion} Hi, I am {{name}}, one of the characters in this story. My voice should stay expressive and consistent.",
                "antagonist": f"{gender_assertion} So, I am {{name}}, and I am here to change everything. My voice carries tension and confidence.",
                "default": f"{gender_assertion} This is a test voice for the character {{name}}. The voice must remain recognizable across long narration."
            }
        else:
            narrator_intro = "Я ваш рассказчик" if gender == "male" else "Я ваша рассказчица" if gender == "female" else "Я ваш диктор"
            gender_assertion = (
                "Мой голос женский, мягкий и уверенный." if gender == "female"
                else "Мой голос нейтральный, ровный и спокойный." if gender == "neutral"
                else "Мой голос мужской, уверенный и ровный."
            )
            base_texts = {
                "narrator": f"{gender_assertion} Добро пожаловать в мир аудиопьесы. {narrator_intro}, и сегодня я проведу вас по этой истории.",
                "protagonist": f"{gender_assertion} Привет, меня зовут {{name}}, и я главный герой этой истории. Я готов к любым испытаниям впереди.",
                "supporting": f"{gender_assertion} Здравствуйте, я {{name}}, один из персонажей этой истории. Мой голос должен звучать выразительно и стабильно.",
                "antagonist": f"{gender_assertion} Что ж, я {{name}}, и я здесь, чтобы изменить ход событий. Мой голос звучит напряженно и уверенно.",
                "default": f"{gender_assertion} Это тестовый голос для персонажа по имени {{name}}. Голос должен быть узнаваемым и стабильным."
            }
        
        role = character.role if character.role in base_texts else "default"
        text = base_texts[role].format(name=character.name)
        
        while len(text.split()) < word_count:
            text += " " + text
        
        return " ".join(text.split()[:word_count])
    
    def _save_audio(self, audio: np.ndarray, sr: int, path: str) -> None:
        """Зберігає аудіо у файл."""
        try:
            from scipy.io import wavfile
            
            # Нормалізація
            if np.max(np.abs(audio)) > 1.0:
                audio = audio / np.max(np.abs(audio))
            
            audio_int16 = (audio * 32767).astype(np.int16)
            wavfile.write(path, sr, audio_int16)
        except ImportError:
            logger.error("[VoiceManager] scipy not installed")
            raise


# =============================================================================
# SECTION 6: AUDIO PROCESSING
# =============================================================================

class AudioMixer:
    """Мікшування аудіо доріжок."""
    
    def __init__(self, sample_rate: int = 24000):
        self.sample_rate = sample_rate
    
    def mix_voice_with_background(
        self,
        voice_path: str,
        background_path: str,
        config: MixingConfig,
        output_path: str
    ) -> str:
        """Мікшує голос з фоновим звуком."""
        try:
            from pydub import AudioSegment
            
            # Завантажуємо аудіо
            voice = AudioSegment.from_file(voice_path)
            background = AudioSegment.from_file(background_path)
            
            # Нормалізуємо гучність
            voice = voice - (voice.dBFS - (-20))  # Цільовий dBFS для голосу
            
            # Підлаштовуємо фоновий звук під довжину голосу
            if len(background) < len(voice):
                # Зациклюємо фон якщо коротший
                loops_needed = (len(voice) // len(background)) + 1
                background = background * loops_needed
            
            background = background[:len(voice)]
            
            # Знижуємо гучність фону
            background = background - (background.dBFS - (-30))  # Тихіший фон
            background = background - 10  # Ще тихіше (dB)
            
            # Додаємо fade ефекти
            voice = voice.fade_in(int(config.voice_fade_in * 1000)).fade_out(
                int(config.voice_fade_out * 1000)
            )
            background = background.fade_in(int(config.background_fade_in * 1000)).fade_out(
                int(config.background_fade_out * 1000)
            )
            
            # Мікшуємо
            mixed = voice.overlay(background)
            
            # Експортуємо
            mixed.export(output_path, format="wav")
            
            logger.info(f"[AudioMixer] Mixed audio saved: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"[AudioMixer] Mixing failed: {e}")
            raise
    
    def add_fade_effects(
        self,
        audio_path: str,
        fade_in: float,
        fade_out: float,
        output_path: Optional[str] = None
    ) -> str:
        """Додає fade in/out ефекти."""
        try:
            from pydub import AudioSegment
            
            audio = AudioSegment.from_file(audio_path)
            
            if fade_in > 0:
                audio = audio.fade_in(int(fade_in * 1000))
            if fade_out > 0:
                audio = audio.fade_out(int(fade_out * 1000))
            
            output = output_path or audio_path
            audio.export(output, format="wav")
            
            return output
        except Exception as e:
            logger.error(f"[AudioMixer] Fade effects failed: {e}")
            raise
    
    def normalize_audio(self, audio_path: str, target_db: float = -20.0) -> str:
        """Нормалізує аудіо до цільового рівня."""
        try:
            from pydub import AudioSegment
            
            audio = AudioSegment.from_file(audio_path)
            change_in_dBFS = target_db - audio.dBFS
            normalized = audio.apply_gain(change_in_dBFS)
            
            normalized.export(audio_path, format="wav")
            return audio_path
        except Exception as e:
            logger.error(f"[AudioMixer] Normalization failed: {e}")
            raise
    
    def add_silence(
        self,
        audio_path: str,
        silence_before: float = 0.0,
        silence_after: float = 0.0,
        output_path: Optional[str] = None
    ) -> str:
        """Додає тишу до аудіо."""
        try:
            from pydub import AudioSegment
            
            audio = AudioSegment.from_file(audio_path)
            
            if silence_before > 0:
                silence = AudioSegment.silent(duration=int(silence_before * 1000))
                audio = silence + audio
            
            if silence_after > 0:
                silence = AudioSegment.silent(duration=int(silence_after * 1000))
                audio = audio + silence
            
            output = output_path or audio_path
            audio.export(output, format="wav")
            
            return output
        except Exception as e:
            logger.error(f"[AudioMixer] Adding silence failed: {e}")
            raise


class AudioAssembler:
    """Збірка фінального аудіофайлу."""
    
    def __init__(self, sample_rate: int = 24000):
        self.sample_rate = sample_rate
        self._mixer = AudioMixer(sample_rate)
    
    def assemble_from_scenes(
        self,
        scenes: List[Scene],
        audio_dir: str,
        config: MixingConfig,
        output_path: str
    ) -> str:
        """Збирає фінальне аудіо з сцен."""
        try:
            from pydub import AudioSegment
            
            final_audio = AudioSegment.silent(duration=0)
            
            for scene in scenes:
                scene_audio = self._assemble_scene(scene, audio_dir, config)
                
                if scene_audio:
                    final_audio = final_audio + scene_audio
                    
                    # Пауза між сценами
                    if config.pause_between_scenes > 0:
                        silence = AudioSegment.silent(
                            duration=int(config.pause_between_scenes * 1000)
                        )
                        final_audio = final_audio + silence
            
            # Експортуємо фінальне аудіо
            final_audio.export(output_path, format="wav")
            
            duration_sec = len(final_audio) / 1000.0
            logger.info(f"[AudioAssembler] Final audio assembled: {output_path} ({duration_sec:.1f}s)")
            
            return output_path
            
        except Exception as e:
            logger.error(f"[AudioAssembler] Assembly failed: {e}")
            raise

    def _is_narrator_dialogue(self, dialogue: Optional[DialogueLine]) -> bool:
        if dialogue is None:
            return False
        return is_narrator_alias(dialogue.character_name) or canonical_character_name(dialogue.character_name) == NARRATOR_NAME

    def _apply_dialogue_smoothing(self, voice_audio: Any, dialogue: DialogueLine, config: MixingConfig, audio_module: Any) -> Any:
        """Легка пост-обробка chunk: fade + контроль хвоста, щоб уникати різких обривів."""
        audio = voice_audio
        if audio is None or len(audio) <= 0:
            return audio

        fade_in_ms = max(0, int(float(config.chunk_fade_in or 0.0) * 1000))
        fade_out_ms = max(0, int(float(config.chunk_fade_out or 0.0) * 1000))
        max_fade = max(fade_in_ms, fade_out_ms)
        if max_fade > 0 and len(audio) > (max_fade * 2 + 30):
            if fade_in_ms > 0:
                audio = audio.fade_in(fade_in_ms)
            if fade_out_ms > 0:
                audio = audio.fade_out(fade_out_ms)

        tail_sec = float(config.narrator_tail_silence if self._is_narrator_dialogue(dialogue) else config.tail_silence)
        tail_ms = max(0, int(tail_sec * 1000))
        if tail_ms > 0:
            audio = audio + audio_module.silent(duration=tail_ms)
        return audio

    def _resolve_dialogue_pause_sec(
        self,
        current_dialogue: DialogueLine,
        next_dialogue: Optional[DialogueLine],
        config: MixingConfig
    ) -> float:
        """Розраховує паузу між репліками з урахуванням зміни спікера."""
        if next_dialogue is None:
            return 0.0
        pause_sec = float(config.pause_between_lines or 0.0)
        current_name = canonical_character_name(current_dialogue.character_name) or current_dialogue.character_name
        next_name = canonical_character_name(next_dialogue.character_name) or next_dialogue.character_name
        speaker_changed = not are_names_equivalent(current_name, next_name)
        if speaker_changed:
            pause_sec += float(config.speaker_change_pause or 0.0)
        if self._is_narrator_dialogue(current_dialogue) != self._is_narrator_dialogue(next_dialogue):
            pause_sec += float(config.narrator_transition_pause or 0.0)
        return max(0.0, pause_sec)
    
    def _assemble_scene(
        self,
        scene: Scene,
        audio_dir: str,
        config: MixingConfig
    ) -> Optional[Any]:
        """Збирає аудіо для однієї сцени."""
        try:
            from pydub import AudioSegment
            
            scene_audio = AudioSegment.silent(duration=0)
            
            for index, dialogue in enumerate(scene.dialogue_lines):
                # Знаходимо файл з аудіо репліки
                if dialogue.audio_path and os.path.exists(dialogue.audio_path):
                    voice_audio = AudioSegment.from_file(dialogue.audio_path)
                    voice_audio = self._apply_dialogue_smoothing(voice_audio, dialogue, config, AudioSegment)
                    scene_audio = scene_audio + voice_audio
                    
                    # Пауза між репліками
                    next_dialogue = scene.dialogue_lines[index + 1] if (index + 1) < len(scene.dialogue_lines) else None
                    pause_sec = self._resolve_dialogue_pause_sec(dialogue, next_dialogue, config)
                    if pause_sec > 0:
                        silence = AudioSegment.silent(
                            duration=int(pause_sec * 1000)
                        )
                        scene_audio = scene_audio + silence
            
            # Додаємо фонові звуки якщо є
            for sound_cue in scene.sound_cues:
                if sound_cue.local_path and os.path.exists(sound_cue.local_path):
                    bg_audio = AudioSegment.from_file(sound_cue.local_path)
                    
                    # Підлаштовуємо гучність
                    bg_audio = bg_audio - int((1 - sound_cue.volume) * 20)
                    
                    # Оверлей на сцену
                    if len(bg_audio) < len(scene_audio):
                        # Зациклюємо фон
                        loops = (len(scene_audio) // len(bg_audio)) + 1
                        bg_audio = bg_audio * loops
                    
                    bg_audio = bg_audio[:len(scene_audio)]
                    scene_audio = scene_audio.overlay(bg_audio)
            
            return scene_audio
            
        except Exception as e:
            logger.error(f"[AudioAssembler] Scene assembly failed: {e}")
            return None
    
    def add_pauses(
        self,
        audio_chunks: List[str],
        pause_duration: float,
        output_path: str
    ) -> str:
        """З'єднує аудіо частини з паузами."""
        try:
            from pydub import AudioSegment
            
            final_audio = AudioSegment.silent(duration=0)
            silence = AudioSegment.silent(duration=int(pause_duration * 1000))
            
            for i, chunk_path in enumerate(audio_chunks):
                if os.path.exists(chunk_path):
                    chunk = AudioSegment.from_file(chunk_path)
                    final_audio = final_audio + chunk
                    
                    if i < len(audio_chunks) - 1:
                        final_audio = final_audio + silence
            
            final_audio.export(output_path, format="wav")
            
            return output_path
            
        except Exception as e:
            logger.error(f"[AudioAssembler] Adding pauses failed: {e}")
            raise
    
    def get_duration(self, audio_path: str) -> float:
        """Повертає тривалість аудіо у секундах."""
        try:
            from pydub import AudioSegment
            audio = AudioSegment.from_file(audio_path)
            return len(audio) / 1000.0
        except Exception as e:
            logger.error(f"[AudioAssembler] Get duration failed: {e}")
            return 0.0


# =============================================================================
# SECTION 7: TELEGRAM BOT HANDLERS
# =============================================================================

class AudioDramaBot:
    """Telegram Bot для управління процесом створення аудіоп'єс."""
    
    def __init__(self, token: str, chat_id: str):
        self.token = token
        self.chat_id = chat_id
        self._bot = None
        self._application = None
        self._bound_thread_id: Optional[int] = None
        self._bound_loop_id: Optional[int] = None
        self._orchestrator: Optional['AudioDramaOrchestrator'] = None
        self._current_scenario: Optional[Scenario] = None
        self._is_processing = False
        self._progress_message_id: Optional[int] = None
        self._progress_chat_id: Optional[Union[int, str]] = None
        self._callback_character_map: Dict[str, str] = {}
        self._voice_card_messages: Dict[str, Tuple[Union[int, str], int, str]] = {}
        self._pending_reference_character: Optional[str] = None
        self._pending_reference_md: Dict[str, str] = {}
    
    def set_orchestrator(self, orchestrator: 'AudioDramaOrchestrator') -> None:
        """Встановлює посилання на оркестратор."""
        self._orchestrator = orchestrator
    
    async def initialize(self) -> bool:
        """Ініціалізує бота."""
        try:
            from telegram import Bot
            from telegram.ext import ApplicationBuilder
            
            builder = ApplicationBuilder().token(self.token)
            for timeout_attr in ("connect_timeout", "read_timeout", "write_timeout", "pool_timeout"):
                if hasattr(builder, timeout_attr):
                    builder = getattr(builder, timeout_attr)(30.0)
            if hasattr(builder, "concurrent_updates"):
                builder = builder.concurrent_updates(True)
            self._application = builder.build()
            self._bot = self._application.bot
            
            # Реєструємо handlers
            self._register_handlers()
            
            # Перевіряємо з'єднання
            me = await asyncio.wait_for(self._bot.get_me(), timeout=15)
            logger.info(f"[AudioDramaBot] Initialized: @{me.username}")
            self._bound_thread_id = threading.get_ident()
            self._bound_loop_id = id(asyncio.get_running_loop())
            
            return True
        except asyncio.TimeoutError:
            logger.error("[AudioDramaBot] Initialization timed out while calling Telegram get_me()")
            return False
        except Exception as e:
            logger.error(f"[AudioDramaBot] Initialization failed: {e}")
            return False

    def is_bound_to_current_loop(self) -> bool:
        """Повертає True, якщо bot-інстанс створено в поточному asyncio loop/thread."""
        try:
            loop = asyncio.get_running_loop()
        except RuntimeError:
            return False
        return (
            self._bound_thread_id == threading.get_ident() and
            self._bound_loop_id == id(loop)
        )
    
    def _register_handlers(self) -> None:
        """Реєструє обробники команд."""
        from telegram.ext import CommandHandler, MessageHandler, filters, CallbackQueryHandler
        
        # Команди
        self._application.add_handler(CommandHandler("start", self.handle_start))
        self._application.add_handler(CommandHandler("help", self.handle_help))
        self._application.add_handler(CommandHandler("menu", self.handle_help))
        self._application.add_handler(CommandHandler("stop", self.handle_stop))
        self._application.add_handler(CommandHandler("cancel", self.handle_cancel))
        self._application.add_handler(CommandHandler("back", self.handle_back))
        self._application.add_handler(CommandHandler("resume", self.handle_resume))
        self._application.add_handler(CommandHandler("render", self.handle_render))
        self._application.add_handler(CommandHandler("status", self.handle_status))
        self._application.add_handler(CommandHandler("voices", self.handle_voices))
        
        # Обробка документів
        self._application.add_handler(MessageHandler(filters.Document.ALL, self.handle_document))
        # Окремо: аудіо/voice як референс голосу персонажа
        self._application.add_handler(MessageHandler(filters.AUDIO | filters.VOICE, self.handle_audio_reference))
        
        # Обробка текстових повідомлень (для команд перегенерації)
        self._application.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, self.handle_text))
        
        # Обробка inline кнопок
        self._application.add_handler(CallbackQueryHandler(self.handle_callback))
        if hasattr(self._application, "add_error_handler"):
            self._application.add_error_handler(self.handle_error)

    def _is_transient_telegram_error(self, error: Optional[Exception]) -> bool:
        """Визначає тимчасові мережеві помилки Telegram/httpx."""
        if error is None:
            return False
        message = str(error).lower()
        markers = [
            "timed out",
            "connecttimeout",
            "readtimeout",
            "networkerror",
            "httpx.readerror",
            "httpx.connecttimeout",
            "httpcore.connecttimeout",
            "httpcore.readtimeout",
        ]
        return any(marker in message for marker in markers)

    async def handle_error(self, update, context) -> None:
        """Глобальний error-handler для PTB, щоб не втрачати винятки в логах."""
        error = getattr(context, "error", None)
        if self._is_transient_telegram_error(error):
            logger.warning(f"[AudioDramaBot] Transient Telegram network error: {error}")
            return
        if error is None:
            logger.error("[AudioDramaBot] Unknown Telegram error (context.error is None)")
            return
        logger.error(
            f"[AudioDramaBot] Unhandled Telegram exception: {error}",
            exc_info=(type(error), error, error.__traceback__)
        )
    
    async def handle_start(self, update, context) -> None:
        """Обробляє команду /start."""
        welcome_message = """🎭 **Audio Drama Bot**

Привіт! Я допоможу створити аудіоп'єсу з твого тексту.

**Підтримувані формати:**
📄 TXT, PDF, EPUB, FB2, DOC, DOCX

**Як користуватися:**
1. Надішли мені файл з текстом
2. Я перетворю його у сценарій
3. Створю голоси для персонажів
4. Згенерую аудіоп'єсу

**Команди:**
/start - почати роботу
/help - коротке меню дій по поточному етапу
/stop - зупинити генерацію
/cancel - скасувати поточний підетап (upload/очікування)
/back - крок назад (вийти з режиму завантаження голосу)
/resume - продовжити останню paused-генерацію
/render - швидко перезібрати фінальне аудіо з кешем реплік
/status - статус поточної роботи
/voices - список голосів
слово ru ежик=ёжик - додати правило вимови
покажи словник ru - показати словник вимови

💡 Після аналізу персонажів є пауза: можна одразу завантажити свої референси голосів,
а потім натиснути '▶️ Продовжити'.

Надішли файл, щоб почати! 📁"""
        
        await update.message.reply_text(welcome_message, parse_mode='Markdown')

    def _get_control_stage(self) -> str:
        """Повертає поточний етап взаємодії з ботом."""
        if self._orchestrator and self._orchestrator.is_waiting_for_pre_voice_reference():
            if self._pending_reference_character:
                return "prevoice_upload"
            return "prevoice_wait"
        if self._is_processing:
            return "processing"
        if self._current_scenario and self._orchestrator:
            approved, total, pending = self._orchestrator.get_voice_approval_summary()
            if total > 0 and pending:
                return "voice_approval"
            if total > 0 and approved == total:
                return "ready_to_render"
        return "idle"

    def _build_context_help_text(self) -> str:
        """Повертає контекстну довідку по доступних діях."""
        stage = self._get_control_stage()
        if stage == "prevoice_upload":
            target = self._pending_reference_character or "персонажа"
            return (
                f"📌 Зараз активний режим завантаження голосу для '{target}'.\n"
                "Надішли audio/voice (або .md, потім audio).\n"
                "Команди: 'назад', 'скасувати' (пропустити свої голоси), 'продовжуй', '/resume', '/stop'."
            )
        if stage == "prevoice_wait":
            return (
                "📌 Етап pre-voice: можна завантажити свої референси перед тестовими голосами.\n"
                "Команди: 'продовжуй' або /resume, 'пропусти' або /cancel, /stop (пауза), /voices."
            )
        if stage == "voice_approval":
            return (
                "📌 Етап затвердження голосів.\n"
                "Команди: 'затверди голос [ім'я]', 'зміни голос [ім'я]', "
                "'додай [характеристика] [ім'я]', /voices, /resume."
            )
        if stage == "ready_to_render":
            return (
                "📌 Усі голоси готові.\n"
                "Команди: /render (швидка перезбірка фіналу), /voices, /status."
            )
        if stage == "processing":
            return (
                "📌 Зараз триває генерація.\n"
                "Команди: /status, /stop, /resume."
            )
        return (
            "📌 Готовий до нового файлу.\n"
            "Надішли TXT/PDF/EPUB/FB2/DOC/DOCX або скористайся /status, /voices."
        )

    def _clear_pending_reference_selection(self, clear_md: bool = False) -> int:
        """Скидає режим завантаження референсу для конкретного героя."""
        cleared = 0
        if self._pending_reference_character:
            key = self._pending_ref_key(self._pending_reference_character)
            if clear_md and key in self._pending_reference_md:
                self._pending_reference_md.pop(key, None)
            self._pending_reference_character = None
            cleared += 1
        return cleared

    async def handle_help(self, update, context) -> None:
        """Показує контекстне меню/довідку."""
        await self._reply_text(update, self._build_context_help_text())

    async def handle_cancel(self, update, context) -> None:
        """Скасовує поточний підетап без втрати всієї сесії."""
        cleared = self._clear_pending_reference_selection(clear_md=False)
        if cleared:
            await self._reply_text(
                update,
                "↩️ Скасовано режим завантаження голосу для обраного персонажа.\n"
                "Можеш обрати іншого героя або продовжити генерацію."
            )
            return

        if self._orchestrator and self._orchestrator.is_waiting_for_pre_voice_reference():
            await self._orchestrator.handle_pre_voice_reference_decision(True)
            await self._reply_text(
                update,
                "⏭ Пропускаю етап завантаження своїх референсів.\n"
                "Запускаю генерацію тестових голосів."
            )
            return

        if self._is_processing:
            await self.handle_stop(update, context)
            return

        await self._reply_text(update, "ℹ️ Немає активної дії для скасування.")

    async def handle_back(self, update, context) -> None:
        """Крок назад у UI-потоці (переважно для upload-режиму)."""
        cleared = self._clear_pending_reference_selection(clear_md=False)
        if cleared:
            await self._reply_text(
                update,
                "↩️ Повернувся назад. Режим завантаження конкретного героя вимкнено.\n"
                f"{self._build_context_help_text()}"
            )
            return
        await self._reply_text(update, self._build_context_help_text())

    async def handle_render(self, update, context) -> None:
        """Інкрементальна перезбірка фінального аудіо з використанням кешу реплік."""
        if self._is_processing:
            await self._reply_text(update, "⏳ Уже триває генерація. Зачекай завершення або натисни /stop.")
            return
        if not self._orchestrator or not self._current_scenario:
            await self._reply_text(update, "❌ Немає активного сценарію для перезбірки.")
            return
        if not self._orchestrator._current_work_dir:
            await self._reply_text(update, "❌ Не знайдено робочу папку поточної сесії.")
            return
        if self._orchestrator.is_waiting_for_pre_voice_reference():
            await self._reply_text(
                update,
                "⌛ Зараз активний pre-voice етап. Спершу заверши його через /resume, /cancel (пропуск) або /stop."
            )
            return
        approved, total, pending = self._orchestrator.get_voice_approval_summary()
        if total > 0 and pending:
            await self._reply_text(
                update,
                "⌛ Не всі голоси затверджені. Спочатку затверди/перегенеруй голоси, потім запускай /render."
            )
            return

        self._is_processing = True
        try:
            dirty_names = []
            if hasattr(self._orchestrator, "get_dirty_characters"):
                dirty_names = self._orchestrator.get_dirty_characters()
            if dirty_names:
                preview = ", ".join(dirty_names[:6])
                if len(dirty_names) > 6:
                    preview += f", ... (+{len(dirty_names) - 6})"
                await self._reply_text(
                    update,
                    "🔁 Запускаю інкрементальну перезбірку.\n"
                    f"Будуть перегенеровані лише репліки змінених голосів: {preview}."
                )
            else:
                await self._reply_text(
                    update,
                    "🔁 Запускаю швидку перезбірку фінального аудіо з повторним використанням готових реплік."
                )

            result = await self._orchestrator.generate_audio_drama(self._current_scenario)
            if result:
                stats = {}
                if hasattr(self._orchestrator, "get_last_generation_stats"):
                    stats = self._orchestrator.get_last_generation_stats()
                reused = int(stats.get("reused_dialogues", 0)) if isinstance(stats, dict) else 0
                regenerated = int(stats.get("regenerated_dialogues", 0)) if isinstance(stats, dict) else 0
                total = int(stats.get("total_dialogues", 0)) if isinstance(stats, dict) else 0
                await self._reply_text(
                    update,
                    (
                        f"✅ Перезбірку завершено.\n"
                        f"📁 {result}\n"
                        f"{self._orchestrator.format_drive_hint(result)}\n"
                        f"♻️ Reuse реплік: {reused}/{total}\n"
                        f"🔄 Перегенеровано: {regenerated}"
                    ).strip()
                )
            else:
                await self._reply_text(
                    update,
                    "ℹ️ Перезбірка не завершилась (можлива пауза або stop). Скористайся /resume або /status."
                )
        except Exception as e:
            logger.error(f"[AudioDramaBot] Render failed: {e}")
            await self._reply_text(update, f"❌ Помилка перезбірки: {str(e)}")
        finally:
            self._is_processing = False
    
    async def handle_stop(self, update, context) -> None:
        """Обробляє команду /stop."""
        if self._is_processing:
            self._is_processing = False
            if self._orchestrator:
                if self._orchestrator.is_waiting_for_pre_voice_reference():
                    await self._orchestrator.handle_pre_voice_reference_decision(False)
                    await self._reply_text(update, "⏹ Зупинено на етапі завантаження референсів.")
                    return
                await self._orchestrator.handle_checkpoint_decision(False)
            await self._reply_text(update, "⏹ Генерацію зупинено.")
        else:
            await self._reply_text(update, "Немає активної генерації.")

    async def handle_resume(self, update, context) -> None:
        """Обробляє команду /resume."""
        if self._is_processing:
            if self._orchestrator and self._orchestrator.is_waiting_for_pre_voice_reference():
                await self._orchestrator.handle_pre_voice_reference_decision(True)
                await self._reply_text(update, "▶️ Продовжую. Запускаю генерацію тестових голосів...")
                return
            await self._reply_text(update, "⏳ Зараз вже виконується генерація.")
            return
        if not self._orchestrator:
            await self._reply_text(update, "❌ Оркестратор не ініціалізовано.")
            return
        
        self._is_processing = True
        try:
            await self._reply_text(update, "▶️ Спроба відновити paused-генерацію...")
            result = await self._orchestrator.resume_from_last_pause()
            self._current_scenario = self._orchestrator._current_scenario
            if result:
                drive_hint = self._orchestrator.format_drive_hint(result)
                await self._reply_text(
                    update,
                    f"✅ Resume завершено. Фінальний файл:\n📁 {result}\n{drive_hint}".strip()
                )
            else:
                await self._reply_text(
                    update,
                    "ℹ️ Немає паузи для відновлення або генерацію не вдалося продовжити."
                )
        except Exception as e:
            logger.error(f"[AudioDramaBot] Resume failed: {e}")
            await self._reply_text(update, f"❌ Помилка resume: {str(e)}")
        finally:
            self._is_processing = False
    
    async def handle_status(self, update, context) -> None:
        """Обробляє команду /status."""
        if self._current_scenario:
            current_lang = self._orchestrator.get_current_language() if self._orchestrator else "unknown"
            status = f"""📊 **Статус**

🎭 П'єса: {self._current_scenario.title}
📝 Сцен: {len(self._current_scenario.scenes)}
👥 Персонажів: {len(self._current_scenario.characters)}
⏱ Очікувана тривалість: {self._current_scenario.total_duration_estimate:.0f} сек
🌐 Мова озвучки: {current_lang}

{'🔄 Обробка...' if self._is_processing else '✅ Готово до роботи'}"""
        else:
            status = "📊 Статус: Очікую файл з текстом."
        
        await update.message.reply_text(status, parse_mode='Markdown')
    
    async def handle_voices(self, update, context) -> None:
        """Обробляє команду /voices."""
        if not self._current_scenario or not self._current_scenario.characters:
            await update.message.reply_text("👥 Голосів поки немає. Надішли файл для створення п'єси.")
            return
        
        voices_info = "👥 **Голоси персонажів:**\n\n"
        approval_map: Dict[str, bool] = {}
        if self._orchestrator:
            approval_map = dict(getattr(self._orchestrator, "_voice_approvals", {}))
        for char in self._current_scenario.characters:
            test_count = len(self._orchestrator._voice_manager._test_voices.get(char.name, []))
            approval_flag = approval_map.get(char.name)
            approval_status = "✅" if approval_flag else "⌛"
            voices_info += f"🎭 **{char.name}** ({char.role})\n"
            voices_info += f"   Стать: {char.gender}, Вік: {char.age_range}\n"
            voices_info += f"   Тестових голосів: {test_count}\n"
            voices_info += f"   Затвердження: {approval_status}\n\n"
        
        await update.message.reply_text(voices_info, parse_mode='Markdown')
    
    def _is_audio_document(self, document) -> bool:
        """Перевіряє, що document є аудіофайлом-референсом."""
        if not document:
            return False
        mime = (getattr(document, "mime_type", "") or "").lower()
        if mime.startswith("audio/"):
            return True
        ext = Path(getattr(document, "file_name", "") or "").suffix.lower()
        return ext in {".wav", ".mp3", ".m4a", ".ogg", ".opus", ".flac", ".aac", ".webm"}
    
    def _pending_ref_key(self, character_name: Optional[str]) -> str:
        return (canonical_character_name(character_name) or character_name or "").strip().lower()
    
    async def handle_audio_reference(self, update, context) -> None:
        """Обробляє завантажене audio/voice повідомлення як референс голосу."""
        if not update or not getattr(update, "message", None):
            return
        
        message = update.message
        source_text = (message.caption or message.text or "").strip()
        if message.audio:
            tg_file = await message.audio.get_file()
            file_name = message.audio.file_name or f"audio_reference_{int(datetime.now().timestamp())}.mp3"
        elif message.voice:
            tg_file = await message.voice.get_file()
            file_name = f"voice_reference_{int(datetime.now().timestamp())}.ogg"
        else:
            await message.reply_text("❌ Не вдалося прочитати аудіо-повідомлення.")
            return
        
        await self._handle_reference_audio_upload(update, tg_file, file_name, source_text)
    
    async def _handle_reference_audio_upload(self, update, tg_file, file_name: str, source_text: str = "") -> None:
        """Завантажує і реєструє користувацький референс голосу персонажа."""
        if not self._orchestrator:
            await update.message.reply_text("❌ Оркестратор не ініціалізовано.")
            return
        
        character_name = self._extract_character_name_from_text(source_text)
        if not character_name and self._pending_reference_character:
            character_name = self._pending_reference_character
        if not character_name and self._orchestrator:
            # Якщо ім'я не вказано, беремо поточного pending героя (послідовне затвердження).
            _, _, pending = self._orchestrator.get_voice_approval_summary()
            if pending:
                character_name = pending[0]
        
        if not character_name:
            keyboard = self._get_character_keyboard() if self._current_scenario else None
            await update.message.reply_text(
                "❌ Не визначив персонажа. Надішли аудіо з підписом, наприклад: "
                "'голос Марка' або 'референс Диктора'.",
                reply_markup=keyboard
            )
            return
        
        await update.message.reply_text(f"🎧 Отримав референс для '{character_name}'. Зберігаю та оновлюю голос...")
        
        suffix = Path(file_name or "reference.wav").suffix or ".wav"
        tmp_path = ""
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                await tg_file.download_to_drive(tmp.name)
                tmp_path = tmp.name
            
            description_override = self._extract_voice_modification(source_text, character_name)
            pending_key = self._pending_ref_key(character_name)
            if pending_key in self._pending_reference_md:
                description_override = self._pending_reference_md.pop(pending_key)
            elif source_text and len(source_text.split()) >= 8 and not description_override:
                # Якщо в caption одразу передали опис голосу природною мовою.
                description_override = source_text
            gender_override = extract_gender_from_text(source_text)
            saved_audio_path, saved_md_path = await self._orchestrator.import_reference_voice(
                character_name=character_name,
                source_audio_path=tmp_path,
                description_override=description_override,
                gender_override=gender_override
            )
            
            if not saved_audio_path:
                await update.message.reply_text("❌ Не вдалося зберегти референсний голос.")
                return
            
            resolved, approved_count, total_count, all_approved = self._orchestrator.update_voice_approval(
                character_name,
                approved=True
            )
            progress_text = self._orchestrator.get_voice_approval_progress_text()
            drive_hint = self._orchestrator.format_drive_hint(saved_audio_path)
            md_info = f"\n📝 Опис: {saved_md_path}" if saved_md_path else ""
            
            sent_audio = await self._reply_audio(
                update,
                saved_audio_path,
                caption=f"🎤 Користувацький референс: {resolved or character_name}"
            )
            if sent_audio:
                self.track_voice_card_message(resolved or character_name, sent_audio)
                await self.mark_voice_card_approved(resolved or character_name, progress_text)
            
            message = (
                f"✅ Референс голосу збережено для '{resolved or character_name}'.\n"
                f"📁 {saved_audio_path}{md_info}\n{drive_hint}\n"
                f"{progress_text}"
            )
            if all_approved:
                message += "\n🎬 Усі голоси затверджені."
            if self._orchestrator and self._orchestrator.is_waiting_for_pre_voice_reference():
                ready_count, total_count_refs, pending_names = self._orchestrator.get_reference_upload_summary()
                pending_preview = ", ".join(pending_names[:6]) if pending_names else "-"
                if len(pending_names) > 6:
                    pending_preview += f", ... (+{len(pending_names) - 6})"
                message += (
                    f"\n\n📦 Готові референси: {ready_count}/{total_count_refs}\n"
                    f"⌛ Ще без власного референсу: {pending_preview}\n"
                    "Після завершення натисни '▶️ Продовжити' у повідомленні pre-stage."
                )
            await update.message.reply_text(message.strip())
            self._pending_reference_character = None
        except Exception as e:
            logger.error(f"[AudioDramaBot] Voice reference upload failed: {e}")
            await update.message.reply_text(f"❌ Помилка імпорту референсу: {str(e)}")
        finally:
            try:
                if tmp_path and os.path.exists(tmp_path):
                    os.remove(tmp_path)
            except Exception:
                pass
    
    async def handle_document(self, update, context) -> None:
        """Обробляє завантажений документ."""
        document = update.message.document
        doc_name = document.file_name or f"doc_{int(datetime.now().timestamp())}"
        doc_ext = Path(doc_name).suffix.lower()
        
        # Аудіо-документи обробляємо як користувацькі референси голосів.
        if self._is_audio_document(document):
            try:
                tg_file = await document.get_file()
                source_text = (update.message.caption or update.message.text or "").strip()
                file_name = doc_name if doc_ext else f"audio_reference_{int(datetime.now().timestamp())}.wav"
                await self._handle_reference_audio_upload(update, tg_file, file_name, source_text)
            except Exception as e:
                logger.error(f"[AudioDramaBot] Audio-document reference failed: {e}")
                await self._reply_text(update, f"❌ Помилка обробки аудіо-референсу: {str(e)}")
            return
        
        # Markdown-документ як опис голосу для наступного завантаженого аудіо.
        if doc_ext == ".md":
            source_text = (update.message.caption or update.message.text or "").strip()
            character_name = self._extract_character_name_from_text(source_text) or self._pending_reference_character
            if not character_name and self._orchestrator:
                _, _, pending = self._orchestrator.get_voice_approval_summary()
                if pending:
                    character_name = pending[0]
            if not character_name:
                keyboard = self._get_character_keyboard() if self._current_scenario else None
                await self._reply_text(
                    update,
                    "❌ Не визначив для кого цей .md опис.\n"
                    "Натисни кнопку '📥 Свій голос' на картці персонажа або додай ім'я в підпис.",
                    reply_markup=keyboard
                )
                return
            tmp_md = ""
            try:
                md_file = await document.get_file()
                with tempfile.NamedTemporaryFile(delete=False, suffix=".md") as tmp:
                    await md_file.download_to_drive(tmp.name)
                    tmp_md = tmp.name
                with open(tmp_md, "r", encoding="utf-8", errors="ignore") as f:
                    md_text = f.read().strip()
                if not md_text:
                    await self._reply_text(update, "⚠️ .md файл порожній.")
                    return
                key = self._pending_ref_key(character_name)
                self._pending_reference_md[key] = md_text
                self._pending_reference_character = character_name
                await self._reply_text(
                    update,
                    f"📝 Опис голосу для '{character_name}' збережено.\n"
                    "Тепер надішли аудіо/voice референс."
                )
            except Exception as e:
                logger.error(f"[AudioDramaBot] Markdown reference parse failed: {e}")
                await self._reply_text(update, f"❌ Не вдалося обробити .md: {str(e)}")
            finally:
                try:
                    if tmp_md and os.path.exists(tmp_md):
                        os.remove(tmp_md)
                except Exception:
                    pass
            return
        
        if self._is_processing:
            if self._orchestrator and self._orchestrator.is_waiting_for_pre_voice_reference():
                await self._reply_text(
                    update,
                    "⏳ Зараз активний pre-voice етап поточного твору.\n"
                    "Цей файл поки не запускаю.\n"
                    "Варіанти: /resume (продовжити поточний), /cancel (пропустити свої голоси), /back (вийти з режиму upload), "
                    "/stop (зупинити поточний), після цього надішли файл знову."
                )
            else:
                await self._reply_text(
                    update,
                    "⏳ Вже обробляю інший запит.\n"
                    "Використай /status або /stop, і лише потім надішли новий файл."
                )
            return
        
        # Перевіряємо формат
        if not FileParser.is_supported(document.file_name):
            await self._reply_text(
                update,
                f"❌ Формат '{document.file_name}' не підтримується.\n"
                f"Підтримувані: {', '.join(FileParser.SUPPORTED_EXTENSIONS)}"
            )
            return
        
        await self._reply_text(update, f"📄 Отримано файл: {document.file_name}\n🔄 Починаю обробку...")
        
        self._is_processing = True
        
        try:
            # Завантажуємо файл
            file = await document.get_file()
            
            with tempfile.NamedTemporaryFile(delete=False, suffix=Path(document.file_name).suffix) as tmp:
                await file.download_to_drive(tmp.name)
                tmp_path = tmp.name
            
            # Обробляємо через оркестратор
            if self._orchestrator:
                result = await self._orchestrator.process_uploaded_file(
                    tmp_path,
                    update.message.from_user.id,
                    source_name=document.file_name
                )
                self._current_scenario = self._orchestrator._current_scenario
                
                if result:
                    drive_hint = ""
                    if self._orchestrator:
                        drive_hint = self._orchestrator.format_drive_hint(result)
                    await self._reply_text(
                        update,
                        f"✅ Аудіоп'єсу створено!\n📁 Файл: {result}\n{drive_hint}".strip()
                    )
                    if os.path.exists(result):
                        already_sent = bool(getattr(self._orchestrator, "_last_final_audio_sent_ok", False))
                        if not already_sent:
                            sent_final = await self._reply_audio(
                                update,
                                result,
                                caption="🎭 Фінальний аудіофайл"
                            )
                            if sent_final and self._orchestrator:
                                self._orchestrator._last_final_audio_sent_ok = True
                    await self._reply_text(update, "⏸ Поточну генерацію завершено. Надішли новий файл для наступної аудіоп'єси.")
                else:
                    paused_hint_sent = False
                    if self._orchestrator and self._orchestrator._current_work_dir:
                        state = self._orchestrator._load_generation_state(self._orchestrator._current_work_dir)
                        status = (state.get("status") or "").strip().lower()
                        reason = (state.get("reason") or "").strip().lower()
                        if status.startswith("paused"):
                            paused_hint_sent = True
                            if reason == "pre_voice_reference_upload_pending":
                                await self._reply_text(
                                    update,
                                    "⏸ Генерацію поставлено на паузу перед тестовими голосами.\n"
                                    "Завантаж референси і натисни '▶️ Продовжити' або використай /resume."
                                )
                            elif reason == "reference_voice_prompt_missing":
                                extra = state.get("extra", {}) if isinstance(state.get("extra", {}), dict) else {}
                                char_name = (extra.get("character_name") or "").strip()
                                ref_path = (extra.get("reference_audio_path") or "").strip()
                                ref_hint = f"\n📁 {ref_path}" if ref_path else ""
                                who = f"для героя '{char_name}' " if char_name else ""
                                await self._reply_text(
                                    update,
                                    "⏸ Генерацію поставлено на паузу (strict reference mode).\n"
                                    f"Не зібрався voice prompt {who}. "
                                    "Онови референс/перегенеруй голос і запусти /resume."
                                    f"{ref_hint}"
                                )
                            elif reason == "voice_approval_pending":
                                await self._reply_text(
                                    update,
                                    "⏸ Генерація на паузі: очікується затвердження голосів.\n"
                                    "Після затвердження використай /resume."
                                )
                            elif reason == "test_voice_generation_failed":
                                extra = state.get("extra", {}) if isinstance(state.get("extra", {}), dict) else {}
                                missing = extra.get("missing_characters", [])
                                missing_preview = ", ".join(missing[:8]) if isinstance(missing, list) else ""
                                if isinstance(missing, list) and len(missing) > 8:
                                    missing_preview += f", ... (+{len(missing) - 8})"
                                missing_line = f"\n⌛ Без аудіо: {missing_preview}" if missing_preview else ""
                                await self._reply_text(
                                    update,
                                    "⏸ Генерацію поставлено на паузу: не вдалося створити тестові голоси.\n"
                                    "Перевір Qwen3-TTS runtime (qwen_tts + transformers) або завантаж свої референси."
                                    f"{missing_line}\n"
                                    "Після виправлення запусти /resume."
                                )
                            else:
                                await self._reply_text(
                                    update,
                                    "⏸ Генерацію поставлено на паузу.\n"
                                    "Використай /resume для продовження."
                                )
                    if not paused_hint_sent:
                        await self._reply_text(update, "❌ Не вдалося створити аудіоп'єсу.")
            else:
                await self._reply_text(update, "❌ Оркестратор не ініціалізовано.")
            
            # Видаляємо тимчасовий файл
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
            
        except Exception as e:
            logger.error(f"[AudioDramaBot] Document processing failed: {e}")
            await self._reply_text(update, f"❌ Помилка обробки: {str(e)}")
        
        finally:
            self._is_processing = False
    
    async def handle_text(self, update, context) -> None:
        """Обробляє текстові повідомлення (команди перегенерації)."""
        original_text = (update.message.text or "").strip()
        text = original_text.lower()

        if self._is_help_command(text):
            await self.handle_help(update, context)
            return

        if self._is_cancel_command(text):
            await self.handle_cancel(update, context)
            return

        if self._is_back_command(text):
            await self.handle_back(update, context)
            return
        
        if self._is_quick_voice_ok(text):
            await self._handle_quick_voice_ok(update)
            return
        
        if self._is_continue_command(text):
            if self._orchestrator and self._orchestrator.is_waiting_for_pre_voice_reference():
                await self._orchestrator.handle_pre_voice_reference_decision(True)
                await self._reply_text(update, "▶️ Продовжую. Запускаю генерацію тестових голосів...")
                return
            await self.handle_resume(update, context)
            return

        if self._is_render_command(text):
            await self.handle_render(update, context)
            return
        
        # Команди словника вимови
        if self._is_pronunciation_add_command(original_text):
            await self._handle_pronunciation_add(update, original_text)
        
        elif self._is_pronunciation_remove_command(original_text):
            await self._handle_pronunciation_remove(update, original_text)
        
        elif self._is_pronunciation_list_command(original_text):
            await self._handle_pronunciation_list(update, original_text)
        
        elif (("затвер" in text or "approve" in text) and "голос" in text):
            await self._handle_voice_approval(update, original_text)
        
        # Команда перегенерації голосу
        elif "зміни голос" in text or "перегенеруй" in text:
            await self._handle_voice_regeneration(update, original_text)
        
        # Команда модифікації голосу
        elif (("додай" in text or "добав" in text) and ("голос" in text or "голосу" in text)):
            await self._handle_voice_modification(update, original_text)
        
        else:
            await update.message.reply_text(
                "Не розпізнав команду.\n"
                f"{self._build_context_help_text()}\n\n"
                "Швидкі приклади:\n"
                "• пропусти (без своїх голосів)\n"
                "• зміни голос [ім'я]\n"
                "• додай [характеристику] [ім'я]\n"
                "• затверди голос [ім'я]\n"
                "• слово ru ежик=ёжик\n"
                "• покажи словник ru"
            )

    def _is_quick_voice_ok(self, text: str) -> bool:
        token = re.sub(r"[^a-zа-яіїєґ]+", "", (text or "").strip().lower())
        return token in {"ок", "ok", "okay", "прийнято", "приймаю", "accept", "yes", "ага"}

    def _is_continue_command(self, text: str) -> bool:
        lowered = (text or "").strip().lower()
        normalized = re.sub(r"\s+", " ", lowered)
        if any(
            phrase in normalized for phrase in [
                "продовжити генерацію",
                "продовжити з паузи",
                "пропусти свої голоси",
                "пропустити свої голоси",
                "без своїх голосів",
                "continue generation",
                "continue from pause",
                "skip own voices",
                "skip voice upload",
            ]
        ):
            return True
        compact = re.sub(r"[^a-zа-яіїєґ]+", "", normalized)
        return compact in {
            "продовжуй",
            "продовжи",
            "продовжити",
            "продолжай",
            "далі",
            "дальше",
            "continue",
            "resume",
            "пропусти",
            "пропустити",
            "skip",
            "go",
            "goon",
            "вперед",
        }

    def _is_cancel_command(self, text: str) -> bool:
        compact = re.sub(r"[^a-zа-яіїєґ]+", "", (text or "").strip().lower())
        return compact in {
            "скасувати",
            "відміни",
            "відміна",
            "отмена",
            "cancel",
            "скас",
        }

    def _is_back_command(self, text: str) -> bool:
        compact = re.sub(r"[^a-zа-яіїєґ]+", "", (text or "").strip().lower())
        return compact in {"назад", "back", "повернись", "повернутись", "goback"}

    def _is_help_command(self, text: str) -> bool:
        compact = re.sub(r"[^a-zа-яіїєґ]+", "", (text or "").strip().lower())
        return compact in {"help", "menu", "меню", "допомога", "шо робити", "щоробити", "команди"}

    def _is_render_command(self, text: str) -> bool:
        lowered = (text or "").strip().lower()
        if "перезбери" in lowered and ("аудіо" in lowered or "финал" in lowered or "фінал" in lowered):
            return True
        if "перегенеруй" in lowered and ("фінал" in lowered or "финал" in lowered):
            return True
        compact = re.sub(r"[^a-zа-яіїєґ]+", "", lowered)
        return compact in {"render", "rebuildaudio", "перезбериаудіо", "перезбірка"}

    async def _handle_quick_voice_ok(self, update) -> None:
        if not self._orchestrator:
            await self._reply_text(update, "❌ Оркестратор не ініціалізовано.")
            return
        resolved, approved_count, total_count, all_approved = self._orchestrator.approve_next_pending_voice()
        if not resolved:
            await self._reply_text(update, "ℹ️ Немає голосів, що очікують затвердження.")
            return
        status_text = self._orchestrator.get_voice_approval_progress_text()
        message = f"✅ Голос '{resolved}' затверджено.\n{status_text}"
        if all_approved:
            message += "\n🎬 Усі голоси затверджені."
        await self.mark_voice_card_approved(resolved, status_text)
        await self._reply_text(update, message)

    def _is_pronunciation_add_command(self, text: str) -> bool:
        lowered = (text or "").strip().lower()
        return ("=" in lowered) and any(
            marker in lowered for marker in [
                "слово",
                "словник",
                "словарь",
                "правило",
                "ударение",
                "наголос",
            ]
        )

    def _is_pronunciation_remove_command(self, text: str) -> bool:
        lowered = (text or "").strip().lower()
        return any(
            lowered.startswith(prefix) for prefix in [
                "видали правило",
                "видалити правило",
                "удали правило",
                "delete rule",
                "remove rule",
                "видали слово",
                "удали слово",
            ]
        )

    def _is_pronunciation_list_command(self, text: str) -> bool:
        lowered = (text or "").strip().lower()
        return any(
            lowered.startswith(prefix) for prefix in [
                "покажи словник",
                "покажи словарь",
                "show dictionary",
                "словник",
                "словарь",
            ]
        ) and ("=" not in lowered)

    def _extract_language_from_command(self, text: str) -> str:
        lowered = (text or "").strip().lower()
        for raw in re.split(r"[\s,;:()\[\]{}]+", lowered):
            token = re.sub(r"[^a-zа-яіїєґ]+", "", raw)
            if not token:
                continue
            if token in LANGUAGE_ALIASES:
                return normalize_language_code(token, "russian")
        if self._orchestrator:
            return self._orchestrator.get_current_language()
        return "russian"

    def _remove_language_token(self, text: str) -> str:
        parts = []
        removed = False
        for token in (text or "").split():
            clean = re.sub(r"[^a-zа-яіїєґ]+", "", token.lower())
            if not removed and clean in LANGUAGE_ALIASES:
                removed = True
                continue
            parts.append(token)
        return " ".join(parts).strip()

    async def _handle_pronunciation_add(self, update, text: str) -> None:
        if not self._orchestrator:
            await self._reply_text(update, "❌ Оркестратор не ініціалізовано.")
            return
        
        payload = text.strip()
        lowered = payload.lower()
        for marker in ["слово", "словник", "словарь", "правило", "ударение", "наголос"]:
            if lowered.startswith(marker):
                payload = payload[len(marker):].strip()
                break
        
        if "=" not in payload:
            await self._reply_text(update, "❌ Формат: слово [ru|uk|en] source=target")
            return
        
        left, right = payload.split("=", 1)
        language = self._extract_language_from_command(left)
        source = self._remove_language_token(left).strip(" \"'")
        target = right.strip().strip(" \"'")
        
        if not source or not target:
            await self._reply_text(update, "❌ Порожнє source або target.")
            return
        
        if self._orchestrator.add_pronunciation_rule(language, source, target):
            dictionary_path = self._orchestrator.get_pronunciation_dictionary_path()
            dictionary_hint = self._orchestrator.format_drive_hint(dictionary_path) if dictionary_path else ""
            await self._reply_text(
                update,
                f"✅ Додано правило ({language}): {source} -> {target}\n"
                f"📁 {dictionary_path}\n{dictionary_hint}".strip()
            )
        else:
            await self._reply_text(update, "❌ Не вдалося зберегти правило.")

    async def _handle_pronunciation_remove(self, update, text: str) -> None:
        if not self._orchestrator:
            await self._reply_text(update, "❌ Оркестратор не ініціалізовано.")
            return
        
        payload = text.strip()
        lowered = payload.lower()
        for prefix in [
            "видали правило",
            "видалити правило",
            "удали правило",
            "delete rule",
            "remove rule",
            "видали слово",
            "удали слово",
        ]:
            if lowered.startswith(prefix):
                payload = payload[len(prefix):].strip()
                break
        
        language = self._extract_language_from_command(payload)
        source = self._remove_language_token(payload).strip(" \"'")
        if not source:
            await self._reply_text(update, "❌ Формат: видали правило [ru|uk|en] source")
            return
        
        if self._orchestrator.remove_pronunciation_rule(language, source):
            await self._reply_text(update, f"✅ Видалено правило ({language}): {source}")
        else:
            await self._reply_text(update, f"⚠️ Правило не знайдено ({language}): {source}")

    async def _handle_pronunciation_list(self, update, text: str) -> None:
        if not self._orchestrator:
            await self._reply_text(update, "❌ Оркестратор не ініціалізовано.")
            return
        
        language = self._extract_language_from_command(text)
        rules = self._orchestrator.list_pronunciation_rules(language)
        if not rules:
            await self._reply_text(update, f"📚 Словник ({language}) порожній.")
            return
        
        sorted_items = sorted(rules.items(), key=lambda item: item[0].lower())
        preview = sorted_items[:80]
        lines = [f"📚 Словник вимови ({language}), правил: {len(sorted_items)}"]
        lines.extend([f"• {src} -> {dst}" for src, dst in preview])
        if len(sorted_items) > len(preview):
            lines.append(f"... і ще {len(sorted_items) - len(preview)}")
        
        await self._reply_text(update, "\n".join(lines))

    async def _reply_text(self, source, text: str, **kwargs) -> None:
        """Уніфікована відповідь для Update і CallbackQuery."""
        try:
            if hasattr(source, "message") and source.message:
                await source.message.reply_text(text, **kwargs)
                return
            if hasattr(source, "edit_message_text"):
                try:
                    await source.edit_message_text(text, **kwargs)
                    return
                except Exception:
                    pass
            if hasattr(source, "message") and source.message:
                await source.message.reply_text(text, **kwargs)
        except Exception as e:
            if self._is_transient_telegram_error(e):
                logger.warning(f"[AudioDramaBot] _reply_text transient network error: {e}")
                return
            logger.error(f"[AudioDramaBot] _reply_text failed: {e}")

    async def _reply_audio(self, source, audio_path: str, caption: str = "", **kwargs) -> Optional[Any]:
        """Уніфікована відправка аудіо для Update і CallbackQuery."""
        if not audio_path or not os.path.exists(audio_path):
            await self._reply_text(source, f"❌ Аудіофайл не знайдено: {audio_path}")
            return None
        safe_caption = (caption or "").strip()
        if len(safe_caption) > 1000:
            safe_caption = safe_caption[:997] + "..."

        async def _send_with_message(message_obj) -> Optional[Any]:
            if message_obj is None:
                return None
            try:
                with open(audio_path, "rb") as audio_file:
                    return await message_obj.reply_audio(audio=audio_file, caption=safe_caption, **kwargs)
            except Exception as err_audio:
                logger.warning(f"[AudioDramaBot] reply_audio fallback to document: {err_audio}")
                try:
                    with open(audio_path, "rb") as audio_file:
                        return await message_obj.reply_document(document=audio_file, caption=safe_caption, **kwargs)
                except Exception as err_doc:
                    logger.warning(f"[AudioDramaBot] reply_document failed: {err_doc}")
                    return None

        sent = None
        try:
            if hasattr(source, "message") and source.message:
                sent = await _send_with_message(source.message)
            elif hasattr(source, "edit_message_text") and hasattr(source, "message") and source.message:
                sent = await _send_with_message(source.message)

            if sent:
                return sent

            try:
                with open(audio_path, "rb") as audio_file:
                    sent = await self._bot.send_audio(chat_id=self.chat_id, audio=audio_file, caption=safe_caption, **kwargs)
                if sent:
                    return sent
            except Exception as err_audio:
                logger.warning(f"[AudioDramaBot] _bot.send_audio failed in _reply_audio: {err_audio}")
                try:
                    with open(audio_path, "rb") as audio_file:
                        sent = await self._bot.send_document(chat_id=self.chat_id, document=audio_file, caption=safe_caption, **kwargs)
                    if sent:
                        return sent
                except Exception as err_doc:
                    logger.warning(f"[AudioDramaBot] _bot.send_document failed in _reply_audio: {err_doc}")

            compressed_copy = await asyncio.to_thread(self._build_telegram_audio_fallback_copy, audio_path)
            try:
                if compressed_copy and os.path.exists(compressed_copy):
                    compressed_caption = (safe_caption + "\nℹ️ Стиснута копія для Telegram.").strip()
                    if len(compressed_caption) > 1000:
                        compressed_caption = compressed_caption[:997] + "..."
                    if hasattr(source, "message") and source.message:
                        try:
                            with open(compressed_copy, "rb") as audio_file:
                                sent = await source.message.reply_audio(audio=audio_file, caption=compressed_caption, **kwargs)
                            if sent:
                                return sent
                        except Exception:
                            try:
                                with open(compressed_copy, "rb") as audio_file:
                                    sent = await source.message.reply_document(document=audio_file, caption=compressed_caption, **kwargs)
                                if sent:
                                    return sent
                            except Exception:
                                pass
                    try:
                        with open(compressed_copy, "rb") as audio_file:
                            sent = await self._bot.send_audio(chat_id=self.chat_id, audio=audio_file, caption=compressed_caption, **kwargs)
                        if sent:
                            return sent
                    except Exception:
                        try:
                            with open(compressed_copy, "rb") as audio_file:
                                sent = await self._bot.send_document(chat_id=self.chat_id, document=audio_file, caption=compressed_caption, **kwargs)
                            if sent:
                                return sent
                        except Exception:
                            pass
            finally:
                try:
                    if compressed_copy and os.path.exists(compressed_copy):
                        os.remove(compressed_copy)
                except Exception:
                    pass
        except Exception as e:
            if self._is_transient_telegram_error(e):
                logger.warning(f"[AudioDramaBot] _reply_audio transient network error: {e}")
                return None
            logger.error(f"[AudioDramaBot] _reply_audio failed: {e}")

        await self._reply_text(source, "❌ Не вдалося відправити аудіофайл у Telegram (audio/document).")
        return None

    async def _edit_callback_response(self, query, text: str) -> None:
        """Безпечне оновлення повідомлення callback (audio caption або text)."""
        try:
            await query.edit_message_caption(caption=text)
            return
        except Exception:
            pass
        try:
            await query.edit_message_text(text=text)
            return
        except Exception:
            pass
        if hasattr(query, "message") and query.message:
            try:
                await query.message.reply_text(text)
            except Exception:
                pass

    async def send_or_update_progress(self, message: str, force_new: bool = False) -> bool:
        """Відправляє або оновлює одне статус-повідомлення прогресу."""
        if not self._bot:
            return False
        
        try:
            if force_new or self._progress_message_id is None or self._progress_chat_id is None:
                sent = await self._bot.send_message(chat_id=self.chat_id, text=message)
                self._progress_message_id = sent.message_id
                self._progress_chat_id = sent.chat_id
                return True
            
            await self._bot.edit_message_text(
                chat_id=self._progress_chat_id,
                message_id=self._progress_message_id,
                text=message
            )
            return True
        except Exception as e:
            # Часта штатна ситуація для Telegram: текст без змін.
            if "message is not modified" in str(e).lower():
                return True
            logger.warning(f"[AudioDramaBot] Progress update fallback to new message: {e}")
            try:
                sent = await self._bot.send_message(chat_id=self.chat_id, text=message)
                self._progress_message_id = sent.message_id
                self._progress_chat_id = sent.chat_id
                return True
            except Exception as inner:
                logger.error(f"[AudioDramaBot] send_or_update_progress failed: {inner}")
                return False

    def reset_progress_message(self) -> None:
        """Скидає прив'язку до progress-повідомлення."""
        self._progress_message_id = None
        self._progress_chat_id = None

    def _register_character_token(self, character_name: str) -> str:
        """Реєструє короткий callback token для персонажа."""
        base = hashlib.md5((character_name or "").encode("utf-8")).hexdigest()[:12]
        token = base
        idx = 0
        while token in self._callback_character_map and self._callback_character_map[token] != character_name:
            idx += 1
            token = f"{base[:10]}{idx:02d}"
        self._callback_character_map[token] = character_name
        return token

    def _resolve_character_token(self, token_or_name: str) -> str:
        """Повертає ім'я персонажа з callback token або raw-рядка."""
        if token_or_name in self._callback_character_map:
            return self._callback_character_map[token_or_name]
        
        token = (token_or_name or "").strip()
        if not token:
            return token_or_name
        
        # Fallback: якщо map втрачено (рестарт), відновлюємо за md5-токеном зі списку персонажів.
        if self._current_scenario and re.fullmatch(r"[a-f0-9]{12}", token):
            for character in self._current_scenario.characters or []:
                digest = hashlib.md5((character.name or "").encode("utf-8")).hexdigest()[:12]
                if digest == token:
                    self._callback_character_map[token] = character.name
                    return character.name
            # Підтримка токенів з суфіксом колізії (base[:10] + NN)
            for character in self._current_scenario.characters or []:
                digest = hashlib.md5((character.name or "").encode("utf-8")).hexdigest()[:12]
                if digest[:10] == token[:10]:
                    self._callback_character_map[token] = character.name
                    return character.name
        
        if self._current_scenario:
            for character in self._current_scenario.characters or []:
                if are_names_equivalent(token, character.name):
                    return character.name
        return token_or_name

    def _voice_card_key(self, character_name: str) -> str:
        return (canonical_character_name(character_name) or character_name or "").strip().lower()

    def track_voice_card_message(self, character_name: str, message_obj: Any) -> None:
        """Запам'ятовує message id картки тестового голосу."""
        if not character_name or not message_obj:
            return
        key = self._voice_card_key(character_name)
        chat_id = getattr(message_obj, "chat_id", None) or getattr(getattr(message_obj, "chat", None), "id", None)
        message_id = getattr(message_obj, "message_id", None)
        if chat_id is None or message_id is None:
            return
        self._voice_card_messages[key] = (chat_id, int(message_id), character_name)

    async def mark_voice_card_approved(self, character_name: str, status_text: str) -> None:
        """Ставить зелену галочку в картці голосу при approve."""
        if not self._bot or not character_name:
            return
        key = self._voice_card_key(character_name)
        target = self._voice_card_messages.get(key)
        if target is None:
            # Fuzzy fallback
            for stored_key, value in self._voice_card_messages.items():
                if are_names_equivalent(stored_key, key):
                    target = value
                    break
        if target is None:
            return
        chat_id, message_id, resolved_name = target
        caption = f"✅ Голос '{resolved_name}' прийнято.\n{status_text}"
        try:
            await self._bot.edit_message_caption(chat_id=chat_id, message_id=message_id, caption=caption)
        except Exception:
            try:
                await self._bot.edit_message_text(chat_id=chat_id, message_id=message_id, text=caption)
            except Exception:
                pass

    def _extract_character_name_from_text(self, text: str) -> Optional[str]:
        """Повертає ім'я персонажа з команди у довільній формі."""
        normalized = (text or "").strip().lower()
        
        if is_narrator_alias(normalized) or any(alias in normalized for alias in NARRATOR_ALIASES):
            return NARRATOR_NAME
        
        if self._current_scenario and self._current_scenario.characters:
            for character in self._current_scenario.characters:
                if character.name.lower() in normalized:
                    return character.name
                # Відмінки типу "диктора" теж мають матчитися на "Диктор"
                if is_narrator_alias(character.name) and any(alias in normalized for alias in NARRATOR_ALIASES):
                    return NARRATOR_NAME
            # Fuzzy match для друкарських/трансліт-помилок
            words = re.findall(r"[a-zа-яіїєґё']{3,}", normalized)
            for word in words:
                for character in self._current_scenario.characters:
                    if are_names_equivalent(word, character.name):
                        return character.name
        
        for marker in ["зміни голос", "перегенеруй голос", "перегенеруй", "голос"]:
            if marker in normalized:
                candidate = normalized.split(marker, 1)[1].strip()
                if candidate:
                    if is_narrator_alias(candidate.split()[0] if candidate.split() else candidate):
                        return NARRATOR_NAME
                    if self._current_scenario and self._current_scenario.characters:
                        for character in self._current_scenario.characters:
                            if character.name.lower() in candidate:
                                return character.name
                            if are_names_equivalent(candidate, character.name):
                                return character.name
                    return candidate.split()[0]
        
        return None

    def _extract_voice_modification(self, text: str, character_name: Optional[str]) -> str:
        """Витягує модифікацію голосу з фрази."""
        normalized = (text or "").strip().lower()
        if not normalized:
            return ""
        
        if character_name:
            lowered_name = character_name.lower()
            normalized = re.sub(
                rf"\b{re.escape(lowered_name)}[a-zа-яіїєґ']*\b",
                " ",
                normalized
            )
            if character_name == NARRATOR_NAME:
                for alias in NARRATOR_ALIASES:
                    normalized = re.sub(rf"\b{re.escape(alias)}[a-zа-яіїєґ']*\b", " ", normalized)
        
        markers = [
            "зміни голос",
            "перегенеруй голос",
            "перегенеруй",
            "додай",
            "добав",
            "голос",
            "голосу",
            "будь ласка",
            "пожалуйста",
        ]
        for marker in markers:
            normalized = normalized.replace(marker, " ")
        
        for filler in ["на", "для", "і", "та"]:
            normalized = re.sub(rf"\b{re.escape(filler)}\b", " ", normalized)
        normalized = re.sub(r"\b(жіноч\w*|чоловіч\w*|female|male|женск\w*|мужск\w*|нейтрал\w*|neutral)\b", " ", normalized)
        normalized = re.sub(r"\s+", " ", normalized).strip(" .,:;!-")
        if len(normalized) < 3:
            return ""
        return normalized
    
    async def _handle_voice_regeneration(self, update, text: str) -> None:
        """Обробляє команду перегенерації голосу."""
        character_name = self._extract_character_name_from_text(text)
        modification = self._extract_voice_modification(text, character_name)
        gender_override = extract_gender_from_text(text)
        
        if not character_name and self._current_scenario:
            # Показуємо список персонажів для вибору
            keyboard = self._get_character_keyboard()
            await self._reply_text(
                update,
                "Оберіть персонажа для перегенерації голосу:",
                reply_markup=keyboard
            )
            return
        
        if self._orchestrator and character_name:
            self._orchestrator.update_voice_approval(character_name, approved=False)
            await self._reply_text(update, f"🔄 Перегенерую голос для '{character_name}'...")
            
            try:
                preview_mode = False
                resolved = self._orchestrator._resolve_voice_approval_name(character_name)
                if resolved:
                    preview_mode = not self._orchestrator._voice_approvals.get(resolved, True)
                new_voice_path = await self._orchestrator.regenerate_character_voice(
                    character_name,
                    modification,
                    gender_override,
                    preview=preview_mode
                )
                
                if new_voice_path:
                    drive_hint = self._orchestrator.format_drive_hint(new_voice_path)
                    applied_gender = ""
                    if gender_override:
                        applied_gender = f"\n👤 Стать голосу: {gender_override}"
                    approved_count, total_count, _ = self._orchestrator.get_voice_approval_summary()
                    progress_text = self._orchestrator.get_voice_approval_progress_text()
                    md_path = os.path.splitext(new_voice_path)[0] + ".md"
                    md_info = f"\n📝 {md_path}" if os.path.exists(md_path) else ""
                    mode_info = "3с прев'ю" if preview_mode else "повний референс"
                    await self._reply_text(
                        update,
                        (
                            f"✅ Новий голос створено ({mode_info})!\n"
                            f"📁 {new_voice_path}{applied_gender}{md_info}\n{drive_hint}\n"
                            f"⌛ Очікує затвердження: {approved_count}/{total_count}\n"
                            f"{progress_text}"
                        ).strip(),
                        reply_markup=self._get_voice_keyboard(character_name)
                    )
                    sent_audio = await self._reply_audio(
                        update,
                        new_voice_path,
                        caption=f"🎤 Новий голос: {character_name}{applied_gender}"
                    )
                    if sent_audio:
                        self.track_voice_card_message(character_name, sent_audio)
                else:
                    await self._reply_text(
                        update,
                        (
                            f"❌ Не вдалося перегенерувати голос для '{character_name}'.\n"
                            "Спробуй ще раз або завантаж свій референс через кнопку '📥 Свій голос'.\n"
                            f"{self._build_context_help_text()}"
                        )
                    )
            except Exception as e:
                await self._reply_text(
                    update,
                    (
                        f"❌ Помилка перегенерації: {str(e)}\n"
                        "Можеш натиснути '📥 Свій голос' і завантажити готовий референс."
                    )
                )
    
    async def _handle_voice_modification(self, update, text: str) -> None:
        """Обробляє команду модифікації голосу."""
        character_name = self._extract_character_name_from_text(text)
        
        if not character_name:
            await self._reply_text(update, "Не знайдено ім'я персонажа. Спробуй: 'додай голосу Марка кавказський акцент'.")
            return
        
        modification = self._extract_voice_modification(text, character_name)
        
        if not modification:
            modification = "оновлений стиль мовлення"
        
        if self._orchestrator:
            self._orchestrator.update_voice_approval(character_name, approved=False)
            await self._reply_text(
                update,
                f"🔄 Модифікую голос '{character_name}': {modification}"
            )
            gender_override = extract_gender_from_text(text)
            preview_mode = False
            resolved = self._orchestrator._resolve_voice_approval_name(character_name)
            if resolved:
                preview_mode = not self._orchestrator._voice_approvals.get(resolved, True)
            
            new_voice = await self._orchestrator.regenerate_character_voice(
                character_name,
                modification,
                gender_override,
                preview=preview_mode
            )
            
            if new_voice:
                drive_hint = self._orchestrator.format_drive_hint(new_voice)
                approved_count, total_count, _ = self._orchestrator.get_voice_approval_summary()
                progress_text = self._orchestrator.get_voice_approval_progress_text()
                md_path = os.path.splitext(new_voice)[0] + ".md"
                md_info = f"\n📝 {md_path}" if os.path.exists(md_path) else ""
                mode_info = "3с прев'ю" if preview_mode else "повний референс"
                await self._reply_text(
                    update,
                    (
                        f"✅ Голос модифіковано ({mode_info})!\n📁 {new_voice}{md_info}\n{drive_hint}\n"
                        f"⌛ Очікує затвердження: {approved_count}/{total_count}\n"
                        f"{progress_text}"
                    ).strip()
                )
                sent_audio = await self._reply_audio(
                    update,
                    new_voice,
                    caption=f"🎤 Модифікований голос: {character_name}"
                )
                if sent_audio:
                    self.track_voice_card_message(character_name, sent_audio)
            else:
                await self._reply_text(
                    update,
                    (
                        "❌ Не вдалося модифікувати голос.\n"
                        "Спробуй іншу характеристику або завантаж власний референс голосу."
                    )
                )

    async def _handle_voice_approval(self, update, text: str) -> None:
        """Ручне затвердження голосу командою."""
        character_name = self._extract_character_name_from_text(text)
        if not character_name:
            await self._reply_text(update, "❌ Не вдалося визначити персонажа для затвердження голосу.")
            return
        if not self._orchestrator:
            await self._reply_text(update, "❌ Оркестратор не ініціалізовано.")
            return
        
        resolved, approved_count, total_count, all_approved = self._orchestrator.update_voice_approval(
            character_name,
            approved=True
        )
        if not resolved:
            await self._reply_text(update, f"⚠️ Персонажа '{character_name}' немає у списку затвердження.")
            return
        
        message = f"✅ Голос '{resolved}' затверджено ({approved_count}/{total_count})"
        message += f"\n{self._orchestrator.get_voice_approval_progress_text()}"
        if all_approved:
            message += "\n🎬 Усі голоси затверджені, можна генерувати фінальне аудіо."
        await self.mark_voice_card_approved(resolved, self._orchestrator.get_voice_approval_progress_text())
        await self._reply_text(update, message)
    
    async def handle_callback(self, update, context) -> None:
        """Обробляє натискання inline кнопок."""
        query = update.callback_query
        try:
            try:
                await query.answer()
            except Exception as e:
                logger.warning(f"[AudioDramaBot] callback answer warning: {e}")
            
            data = query.data
            
            if data in {"prevoice_continue", "prevoice_skip", "prevoice_stop"}:
                self._clear_pending_reference_selection(clear_md=False)
                if self._orchestrator:
                    await self._orchestrator.handle_pre_voice_reference_decision(True)
                if data == "prevoice_skip":
                    await self._edit_callback_response(query, "⏭ Пропускаю завантаження своїх референсів. Запускаю тестові голоси...")
                elif data == "prevoice_stop":
                    await self._edit_callback_response(query, "ℹ️ Стара кнопка 'Зупинити' оброблена як 'Пропустити'. Запускаю тестові голоси...")
                else:
                    await self._edit_callback_response(query, "▶️ Продовжую. Запускаю генерацію тестових голосів...")

            elif data == "prevoice_pause":
                self._is_processing = False
                self._clear_pending_reference_selection(clear_md=False)
                if self._orchestrator:
                    await self._orchestrator.handle_pre_voice_reference_decision(False)
                await self._edit_callback_response(query, "⏸ Етап pre-voice поставлено на паузу. Для продовження надішли /resume.")

            elif data == "upl_cancel":
                cleared = self._clear_pending_reference_selection(clear_md=False)
                if cleared:
                    await self._edit_callback_response(query, "↩️ Режим завантаження свого голосу скасовано.")
                else:
                    await self._edit_callback_response(query, "ℹ️ Режим завантаження голосу не був активним.")

            elif data == "stop":
                self._is_processing = False
                self._clear_pending_reference_selection(clear_md=False)
                if self._orchestrator:
                    if self._orchestrator.is_waiting_for_pre_voice_reference():
                        await self._orchestrator.handle_pre_voice_reference_decision(False)
                        await self._edit_callback_response(query, "⏹ Зупинено на етапі завантаження референсів.")
                    else:
                        await self._orchestrator.handle_checkpoint_decision(False)
                        await self._edit_callback_response(query, "⏹ Генерацію зупинено.")
                else:
                    await self._edit_callback_response(query, "⏹ Генерацію зупинено.")
            
            elif data in {"continue", "generate"}:
                if self._orchestrator:
                    if self._orchestrator.is_waiting_for_pre_voice_reference():
                        await self._orchestrator.handle_pre_voice_reference_decision(True)
                        await self._edit_callback_response(query, "▶️ Продовжую. Запускаю генерацію тестових голосів...")
                    else:
                        await self._orchestrator.handle_checkpoint_decision(True)
                        await self._edit_callback_response(query, "▶️ Продовжую генерацію...")
                else:
                    await self._edit_callback_response(query, "▶️ Продовжую генерацію...")
            
            elif data.startswith("reg:"):
                token = data.replace("reg:", "", 1)
                character_name = self._resolve_character_token(token)
                await self._handle_voice_regeneration(query, f"зміни голос {character_name}")
            
            elif data.startswith("regenerate_"):
                character_name = data.replace("regenerate_", "", 1)
                character_name = self._resolve_character_token(character_name)
                await self._handle_voice_regeneration(query, f"зміни голос {character_name}")
            
            elif data.startswith("upl:"):
                token = data.replace("upl:", "", 1)
                character_name = self._resolve_character_token(token)
                self._pending_reference_character = character_name
                await self._edit_callback_response(
                    query,
                    (
                        f"📥 Режим завантаження свого голосу для '{character_name}' активовано.\n"
                        "Надішли аудіо/voice у чат.\n"
                        "Опційно: спочатку надішли .md файл з описом голосу."
                    )
                )
                if hasattr(query, "message") and query.message:
                    try:
                        await query.message.reply_text(
                            "Керування режимом upload:",
                            reply_markup=self._get_reference_upload_mode_keyboard()
                        )
                    except Exception:
                        pass
            
            elif data.startswith("acc:") or data.startswith("accept_"):
                if data.startswith("acc:"):
                    token = data.replace("acc:", "", 1)
                else:
                    token = data.replace("accept_", "", 1)
                character_name = self._resolve_character_token(token)
                if self._orchestrator:
                    resolved, approved_count, total_count, all_approved = self._orchestrator.update_voice_approval(
                        character_name,
                        approved=True
                    )
                    if resolved:
                        progress_text = self._orchestrator.get_voice_approval_progress_text()
                        message = f"✅ Голос '{resolved}' прийнято ({approved_count}/{total_count}).\n{progress_text}"
                        if all_approved:
                            message += "\n🎬 Усі голоси затверджені."
                        await self.mark_voice_card_approved(resolved, progress_text)
                        await self._edit_callback_response(query, message)
                    else:
                        await self._edit_callback_response(query, f"⚠️ Персонажа '{character_name}' не знайдено у списку затвердження.")
                else:
                    await self._edit_callback_response(query, f"✅ Голос '{character_name}' прийнято.")
            
            elif data == "final_generate":
                if self._orchestrator and self._current_scenario:
                    await self._edit_callback_response(query, "🎬 Генерую фінальну аудіоп'єсу...")
                    # Запускаємо фінальну генерацію
            else:
                await self._edit_callback_response(query, f"ℹ️ Невідома команда: {data}")
        except Exception as e:
            logger.error(f"[AudioDramaBot] handle_callback failed: {e}")
            try:
                await self._edit_callback_response(query, f"❌ Помилка callback: {str(e)}")
            except Exception:
                pass
    
    async def handle_voice_command(self, update, context) -> None:
        """Обробляє голосові команди."""
        # Для майбутнього розширення - обробка голосових повідомлень
        await update.message.reply_text("🎤 Голосові команди поки не підтримуються.")
    
    async def handle_regenerate_request(self, update, context) -> None:
        """Обробляє запит на перегенерацію."""
        await self._handle_voice_regeneration(update, update.message.text)
    
    async def send_checkpoint(
        self,
        audio_path: str,
        duration: float,
        message: str = "",
        drive_hint: str = ""
    ) -> bool:
        """Відправляє контрольну точку в Telegram."""
        if not self._bot:
            return False
        
        try:
            caption = f"""📍 Контрольна точка

⏱ Тривалість: {duration:.1f} сек
📁 Файл: {os.path.basename(audio_path)}
{message}
{drive_hint}"""
            
            with open(audio_path, 'rb') as audio_file:
                await self._bot.send_audio(
                    chat_id=self.chat_id,
                    audio=audio_file,
                    caption=caption,
                    reply_markup=self._get_checkpoint_keyboard()
                )
            
            logger.info(f"[AudioDramaBot] Checkpoint sent: {duration:.1f}s")
            return True
            
        except Exception as e:
            logger.error(f"[AudioDramaBot] Failed to send checkpoint: {e}")
            return False
    
    async def send_inline_buttons(
        self,
        message: str,
        buttons: List[Dict],
        chat_id: Optional[str] = None
    ) -> bool:
        """Відправляє повідомлення з inline кнопками."""
        if not self._bot:
            return False
        
        try:
            from telegram import InlineKeyboardButton, InlineKeyboardMarkup
            
            keyboard = []
            for button_row in buttons:
                row = []
                for btn in button_row:
                    row.append(InlineKeyboardButton(
                        text=btn.get("text", ""),
                        callback_data=btn.get("callback_data", "")
                    ))
                keyboard.append(row)
            
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            await self._bot.send_message(
                chat_id=chat_id or self.chat_id,
                text=message,
                reply_markup=reply_markup
            )
            
            return True
            
        except Exception as e:
            logger.error(f"[AudioDramaBot] Failed to send inline buttons: {e}")
            return False
    
    async def send_message(self, message: str, chat_id: Optional[str] = None) -> bool:
        """Відправляє текстове повідомлення."""
        if not self._bot:
            return False
        
        try:
            await self._bot.send_message(
                chat_id=chat_id or self.chat_id,
                text=message
            )
            return True
        except Exception as e:
            logger.error(f"[AudioDramaBot] Failed to send message: {e}")
            return False

    def _build_telegram_audio_fallback_copy(self, audio_path: str) -> Optional[str]:
        """Створює стиснуту mp3-копію для Telegram, якщо оригінал не відправляється."""
        if not audio_path or not os.path.exists(audio_path):
            return None
        tmp_path = ""
        try:
            from pydub import AudioSegment
            with tempfile.NamedTemporaryFile(delete=False, suffix=".mp3") as tmp:
                tmp_path = tmp.name
            audio = AudioSegment.from_file(audio_path)
            if len(audio) <= 0:
                return None
            # Mono + 24kHz + 96kbit дає стабільно менший розмір для Telegram.
            audio = audio.set_channels(1).set_frame_rate(24000)
            audio.export(tmp_path, format="mp3", bitrate="96k")
            if os.path.exists(tmp_path) and os.path.getsize(tmp_path) > 0:
                return tmp_path
        except Exception as e:
            logger.warning(f"[AudioDramaBot] Failed to build compressed Telegram copy: {e}")
        try:
            if tmp_path and os.path.exists(tmp_path):
                os.remove(tmp_path)
        except Exception:
            pass
        return None
    
    async def send_audio(self, audio_path: str, caption: str = "", reply_markup: Any = None) -> Optional[Any]:
        """Відправляє аудіо файл."""
        if not self._bot:
            return None
        if not audio_path or not os.path.exists(audio_path):
            logger.error(f"[AudioDramaBot] Audio file not found: {audio_path}")
            return None

        safe_caption = (caption or "").strip()
        if len(safe_caption) > 1000:
            safe_caption = safe_caption[:997] + "..."

        send_errors: List[str] = []

        # 1) Preferred: send as audio.
        try:
            with open(audio_path, "rb") as audio_file:
                sent = await self._bot.send_audio(
                    chat_id=self.chat_id,
                    audio=audio_file,
                    caption=safe_caption,
                    reply_markup=reply_markup
                )
            return sent
        except Exception as e:
            send_errors.append(f"send_audio(original): {e}")
            logger.warning(f"[AudioDramaBot] send_audio failed, trying send_document: {e}")

        # 2) Fallback: send as generic document.
        try:
            with open(audio_path, "rb") as audio_file:
                sent = await self._bot.send_document(
                    chat_id=self.chat_id,
                    document=audio_file,
                    caption=safe_caption,
                    reply_markup=reply_markup
                )
            return sent
        except Exception as e:
            send_errors.append(f"send_document(original): {e}")
            logger.warning(f"[AudioDramaBot] send_document failed, trying compressed copy: {e}")

        # 3) Last fallback: compressed mp3 copy.
        compressed_copy = await asyncio.to_thread(self._build_telegram_audio_fallback_copy, audio_path)
        try:
            if compressed_copy and os.path.exists(compressed_copy):
                compressed_caption = (safe_caption + "\nℹ️ Стиснута копія для Telegram.").strip()
                if len(compressed_caption) > 1000:
                    compressed_caption = compressed_caption[:997] + "..."
                try:
                    with open(compressed_copy, "rb") as audio_file:
                        sent = await self._bot.send_audio(
                            chat_id=self.chat_id,
                            audio=audio_file,
                            caption=compressed_caption,
                            reply_markup=reply_markup
                        )
                    return sent
                except Exception as e:
                    send_errors.append(f"send_audio(compressed): {e}")

                try:
                    with open(compressed_copy, "rb") as audio_file:
                        sent = await self._bot.send_document(
                            chat_id=self.chat_id,
                            document=audio_file,
                            caption=compressed_caption,
                            reply_markup=reply_markup
                        )
                    return sent
                except Exception as e:
                    send_errors.append(f"send_document(compressed): {e}")
        finally:
            try:
                if compressed_copy and os.path.exists(compressed_copy):
                    os.remove(compressed_copy)
            except Exception:
                pass

        logger.error(
            "[AudioDramaBot] Failed to send audio after all fallbacks: "
            + " | ".join(send_errors[:6])
        )
        return None
    
    def _get_main_keyboard(self) -> Any:
        """Повертає головну клавіатуру."""
        from telegram import InlineKeyboardButton, InlineKeyboardMarkup
        
        keyboard = [
            [
                InlineKeyboardButton("▶️ Генерація", callback_data="generate"),
                InlineKeyboardButton("⏹ Зупинити", callback_data="stop")
            ]
        ]
        return InlineKeyboardMarkup(keyboard)
    
    def _get_checkpoint_keyboard(self) -> Any:
        """Повертає клавіатуру для контрольної точки."""
        from telegram import InlineKeyboardButton, InlineKeyboardMarkup
        
        keyboard = [
            [
                InlineKeyboardButton("▶️ Генерація", callback_data="generate"),
                InlineKeyboardButton("⏹ Зупинити", callback_data="stop")
            ]
        ]
        return InlineKeyboardMarkup(keyboard)
    
    def _get_character_keyboard(self) -> Any:
        """Повертає клавіатуру з персонажами."""
        from telegram import InlineKeyboardButton, InlineKeyboardMarkup
        
        keyboard = []
        if self._current_scenario:
            for char in self._current_scenario.characters:
                token = self._register_character_token(char.name)
                keyboard.append([
                    InlineKeyboardButton(
                        f"🎭 {char.name}",
                        callback_data=f"reg:{token}"
                    )
                ])
        
        return InlineKeyboardMarkup(keyboard)
    
    def _get_voice_keyboard(self, character_name: str) -> Any:
        """Повертає клавіатуру для голосу персонажа."""
        from telegram import InlineKeyboardButton, InlineKeyboardMarkup
        token = self._register_character_token(character_name)
        
        keyboard = [
            [
                InlineKeyboardButton("✅ Прийняти", callback_data=f"acc:{token}"),
                InlineKeyboardButton("🔄 Перегенерувати", callback_data=f"reg:{token}")
            ],
            [
                InlineKeyboardButton("📥 Свій голос", callback_data=f"upl:{token}")
            ]
        ]
        return InlineKeyboardMarkup(keyboard)

    def _get_reference_upload_mode_keyboard(self) -> Any:
        """Повертає клавіатуру керування для режиму upload референсу."""
        from telegram import InlineKeyboardButton, InlineKeyboardMarkup

        keyboard = [
            [InlineKeyboardButton("↩️ Назад", callback_data="upl_cancel")],
            [
                InlineKeyboardButton("▶️ Продовжити", callback_data="prevoice_continue"),
                InlineKeyboardButton("⏭ Пропустити", callback_data="prevoice_skip"),
            ],
            [InlineKeyboardButton("⏸ Пауза", callback_data="prevoice_pause")],
        ]
        return InlineKeyboardMarkup(keyboard)
    
    async def run(self) -> None:
        """Запускає бота."""
        if self._application:
            await self._application.initialize()
            await self._application.start()
            await self._application.updater.start_polling()
            
            logger.info("[AudioDramaBot] Bot started polling...")
            
            # Тримаємо бота активним
            while True:
                await asyncio.sleep(1)
    
    async def stop(self) -> None:
        """Зупиняє бота."""
        if self._application:
            await self._application.updater.stop()
            await self._application.stop()
            await self._application.shutdown()
            
            logger.info("[AudioDramaBot] Bot stopped")


# =============================================================================
# SECTION 8: MAIN ORCHESTRATOR
# =============================================================================

class AudioDramaOrchestrator:
    """Головний координатор процесу створення аудіоп'єси."""
    
    def __init__(self, config: OrchestratorConfig):
        self.config = config
        
        # API клієнти
        self._gemini_client: Optional[GeminiClient] = None
        self._pexels_client: Optional[PexelsClient] = None
        
        # TTS компоненти (з voice_library_all_in_one)
        self._design_engine = None
        self._clone_engine = None
        self._library_manager = None
        self._cache = None
        self._create_voice_from_reference_fn = None
        self._voice_library_entry_cls = None
        
        # Audio компоненти
        self._voice_manager: Optional[VoiceManager] = None
        self._audio_mixer: Optional[AudioMixer] = None
        self._audio_assembler: Optional[AudioAssembler] = None
        self._pronunciation_manager: Optional[PronunciationManager] = None
        self._quality_guard: Optional[AudioQualityGuard] = None
        
        # Telegram бот
        self._bot: Optional[AudioDramaBot] = None
        
        # Поточний стан
        self._current_scenario: Optional[Scenario] = None
        self._current_work_dir: Optional[str] = None
        self._generated_duration: float = 0.0
        self._checkpoints: List[Checkpoint] = []
        self._checkpoint_targets_sent: set[int] = set()
        self._stop_requested: bool = False
        self._checkpoint_resume_event: asyncio.Event = asyncio.Event()
        self._checkpoint_resume_event.set()
        self._waiting_for_checkpoint_resume: bool = False
        self._pre_voice_reference_event: asyncio.Event = asyncio.Event()
        self._pre_voice_reference_event.set()
        self._waiting_for_pre_voice_reference: bool = False
        self._voice_approvals: Dict[str, bool] = {}
        self._voice_approval_event: asyncio.Event = asyncio.Event()
        self._voice_approval_event.clear()
        self._voice_update_event: asyncio.Event = asyncio.Event()
        self._voice_update_event.clear()
        self._current_approval_character: Optional[str] = None
        self._reference_render_tasks: Dict[str, asyncio.Task] = {}
        self._reference_render_errors: Dict[str, str] = {}
        self._reference_render_lock: asyncio.Lock = asyncio.Lock()
        self._preview_generation_active: bool = False
        self._last_progress_percent: int = -1
        self._last_progress_min_bucket: int = -1
        self._current_language: str = normalize_language_code(self.config.language, "russian")
        self._resume_pointer_file: str = os.path.join(self.config.drive_base_path, "last_paused_session.json")
        self._state_file_name: str = "generation_state.json"
        self._quality_reports_dir_name: str = "quality_reports"
        self._is_initialized = False
        self._drive_service = None
        self._drive_service_failed: bool = False
        self._drive_service_notice_sent: bool = False
        self._drive_item_id_cache: Dict[Tuple[str, bool], Optional[str]] = {}
        self._drive_url_cache: Dict[str, Optional[str]] = {}
        self._dirty_characters: set[str] = set()
        self._last_generation_stats: Dict[str, int] = {
            "total_dialogues": 0,
            "reused_dialogues": 0,
            "regenerated_dialogues": 0,
        }
        self._qwen_runtime_install_attempted: bool = False
        self._qwen_runtime_install_error: str = ""
        self._telegram_token: str = ""
        self._telegram_chat_id: str = ""
        self._last_final_audio_sent_ok: bool = False
    
    async def initialize(
        self,
        telegram_token: Optional[str] = None,
        telegram_chat_id: Optional[str] = None,
        gemini_api_key: Optional[str] = None,
        pexels_api_key: Optional[str] = None
    ) -> bool:
        """Ініціалізує всі компоненти."""
        logger.info("[Orchestrator] Initializing...")
        
        try:
            self._telegram_token = (telegram_token or "").strip()
            self._telegram_chat_id = (telegram_chat_id or "").strip()

            # Створюємо директорії
            os.makedirs(self.config.drive_base_path, exist_ok=True)
            os.makedirs(self.config.cache_dir, exist_ok=True)
            pronunciation_path = os.path.join(self.config.drive_base_path, "pronunciation_dictionaries.json")
            self._pronunciation_manager = PronunciationManager(pronunciation_path)
            self._quality_guard = AudioQualityGuard(
                enabled=self.config.quality_guard_enabled,
                max_attempts=self.config.quality_guard_max_attempts,
                similarity_threshold=self.config.quality_guard_similarity_threshold,
                min_tail_silence_ms=self.config.quality_guard_min_tail_silence_ms,
                use_openai_asr=self.config.quality_guard_use_openai_asr,
                openai_asr_model=self.config.quality_guard_asr_model,
                openai_api_key=os.getenv("OPENAI_API_KEY"),
                use_local_asr=self.config.quality_guard_use_local_asr,
                local_asr_model=self.config.quality_guard_local_asr_model,
                local_asr_device=self.config.quality_guard_local_asr_device,
                local_asr_compute_type=self.config.quality_guard_local_asr_compute_type
            )
            
            # Ініціалізуємо Gemini
            if gemini_api_key:
                self._gemini_client = GeminiClient(gemini_api_key)
                logger.info("[Orchestrator] Gemini client initialized")
            
            # Ініціалізуємо Pexels
            if pexels_api_key:
                self._pexels_client = PexelsClient(pexels_api_key)
                logger.info("[Orchestrator] Pexels client initialized")
            
            # Ініціалізуємо TTS двигуни
            await self._initialize_tts_engines()
            
            # Ініціалізуємо Voice Manager
            self._voice_manager = VoiceManager(
                design_engine=self._design_engine,
                clone_engine=self._clone_engine,
                library_manager=self._library_manager,
                cache=self._cache,
                create_from_reference_fn=self._create_voice_from_reference_fn,
                text_preprocessor=self._prepare_tts_text
            )
            
            # Ініціалізуємо Audio компоненти
            self._audio_mixer = AudioMixer()
            self._audio_assembler = AudioAssembler()
            
            # Ініціалізуємо Telegram бот
            if self._telegram_token and self._telegram_chat_id:
                self._bot = AudioDramaBot(self._telegram_token, self._telegram_chat_id)
                self._bot.set_orchestrator(self)
                await self._bot.initialize()
                logger.info("[Orchestrator] Telegram bot initialized")
            
            self._is_initialized = True
            logger.info("[Orchestrator] Initialization complete!")
            return True
            
        except Exception as e:
            logger.error(f"[Orchestrator] Initialization failed: {e}")
            return False
    
    async def _initialize_tts_engines(self) -> None:
        """Ініціалізує TTS двигуни з багаторівневим fallback."""
        errors: List[str] = []

        # One-script режим: тільки внутрішні Qwen3-TTS двигуни цього файлу.
        if self.config.one_script_mode:
            if self.config.qwen_precheck_on_init:
                try:
                    self._ensure_qwen_tts_runtime(allow_install=True)
                except Exception as runtime_error:
                    errors.append(f"qwen runtime precheck failed: {runtime_error}")
            else:
                logger.warning(
                    "[Orchestrator] Qwen runtime precheck deferred on init "
                    "(AUDIO_DRAMA_QWEN_PRECHECK_ON_INIT=0). "
                    "Bot will start faster; runtime checks run on first synthesis."
                )
            if await self._try_init_tts_internal_engines(errors):
                logger.info("[Orchestrator] One-script mode enabled (internal Qwen engines only)")
                return
            details = " | ".join(errors[:6]) if errors else "unknown error"
            if self.config.require_qwen_dual_models:
                raise RuntimeError(
                    "One-script Qwen mode required, but internal Qwen3-TTS init failed: "
                    f"{details}"
                )
            logger.warning(
                "[Orchestrator] One-script internal Qwen init failed, switching to no-TTS fallback: "
                f"{details}"
            )
            self._create_fallback_engines()
            return

        # Strict режим: лише дві Qwen3-TTS моделі (VoiceDesign + VoiceClone).
        if self.config.require_qwen_dual_models:
            if self.config.qwen_precheck_on_init:
                try:
                    self._ensure_qwen_tts_runtime(allow_install=True)
                except Exception as runtime_error:
                    errors.append(f"qwen runtime precheck failed: {runtime_error}")
            else:
                logger.warning(
                    "[Orchestrator] Strict Qwen mode enabled, but runtime precheck is deferred "
                    "(AUDIO_DRAMA_QWEN_PRECHECK_ON_INIT=0)."
                )
            if await self._try_init_tts_internal_engines(errors):
                logger.info("[Orchestrator] Qwen dual-model mode enabled (VoiceDesign + VoiceClone)")
                return
            details = " | ".join(errors[:6]) if errors else "unknown error"
            raise RuntimeError(
                "Qwen dual-model mode required, but internal Qwen3-TTS init failed: "
                f"{details}"
            )
        
        if self.config.prefer_internal_tts:
            # 1) Internal fallback: вбудовані qwen-tts двигуни без зовнішніх py-файлів
            if await self._try_init_tts_internal_engines(errors):
                return
            
            if self.config.allow_external_tts_fallback:
                # 2) Fallback: voice_library_all_in_one (в sys.path або через прямий файл)
                if await self._try_init_tts_from_all_in_one(errors):
                    return
                
                # 3) Fallback: модульний стек (voicebox_adapter_colab + voice_library_manager)
                if await self._try_init_tts_from_modular_stack(errors):
                    return
            else:
                errors.append("external TTS fallback disabled by config (allow_external_tts_fallback=False)")
        else:
            # 1) Основний шлях: voice_library_all_in_one (в sys.path або через прямий файл)
            if await self._try_init_tts_from_all_in_one(errors):
                return
            
            # 2) Fallback: модульний стек (voicebox_adapter_colab + voice_library_manager)
            if await self._try_init_tts_from_modular_stack(errors):
                return
            
            # 3) Internal fallback
            if await self._try_init_tts_internal_engines(errors):
                return
        
        # 4) Фінальний fallback без TTS
        logger.warning("[Orchestrator] Failed to initialize TTS stack. Falling back to no-TTS mode.")
        for i, err in enumerate(errors[:12], start=1):
            logger.warning(f"[Orchestrator] TTS init error {i}: {err}")
        self._create_fallback_engines()

    def _engine_is_loaded(self, engine: Any) -> bool:
        """Уніфікована перевірка завантаження моделі."""
        if engine is None:
            return False
        if hasattr(engine, "is_loaded"):
            try:
                return bool(engine.is_loaded())
            except Exception:
                pass
        return getattr(engine, "_model", None) is not None

    async def _load_engine(self, engine: Any, engine_name: str) -> bool:
        """Безпечне завантаження моделі двигуна."""
        if engine is None:
            return False
        if self._engine_is_loaded(engine):
            return True
        if not hasattr(engine, "load_model"):
            logger.warning(f"[Orchestrator] {engine_name} has no load_model()")
            return False
        try:
            loaded = await asyncio.to_thread(engine.load_model)
            if loaded:
                logger.info(f"[Orchestrator] {engine_name} loaded")
            else:
                logger.warning(f"[Orchestrator] {engine_name} failed to load")
                # Для Qwen-двигунів робимо одну спробу автоінсталяції runtime і retry
                if "Voice" in engine_name:
                    if not self.config.qwen_auto_install_on_load:
                        logger.warning(
                            f"[Orchestrator] Runtime auto-install skipped for {engine_name} "
                            "(AUDIO_DRAMA_QWEN_AUTO_INSTALL_ON_LOAD=0)."
                        )
                    else:
                        if self._bot and not self._qwen_runtime_install_attempted:
                            await self._send_progress(
                                "⚙️ Qwen TTS runtime відсутній. Запускаю одноразове встановлення "
                                "(це може тривати 2-8 хв)."
                            )
                        try:
                            self._ensure_qwen_tts_runtime(allow_install=True)
                            loaded = await asyncio.to_thread(engine.load_model)
                            if loaded:
                                if self._bot:
                                    await self._send_progress("✅ Qwen TTS runtime встановлено, продовжую генерацію.")
                                logger.info(f"[Orchestrator] {engine_name} loaded after runtime install")
                                return True
                        except Exception as runtime_error:
                            logger.warning(f"[Orchestrator] Runtime install/retry failed for {engine_name}: {runtime_error}")
                            if self._bot:
                                await self._send_progress(
                                    "❌ Не вдалося встановити/завантажити Qwen TTS runtime.\n"
                                    "Можна завантажити свої референси голосів і продовжити через /resume."
                                )
            return bool(loaded)
        except Exception as e:
            logger.warning(f"[Orchestrator] {engine_name} load error: {e}")
            return False

    def _unload_engine(self, engine: Any, engine_name: str) -> None:
        """Вивантажує модель з пам'яті/VRAM."""
        if engine is None:
            return
        try:
            if hasattr(engine, "set_model"):
                try:
                    engine.set_model(None)
                except Exception:
                    pass
            if hasattr(engine, "_model"):
                try:
                    setattr(engine, "_model", None)
                except Exception:
                    pass
            logger.info(f"[Orchestrator] {engine_name} unloaded")
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to unload {engine_name}: {e}")
        
        try:
            import gc
            gc.collect()
        except Exception:
            pass
        
        try:
            import torch
            if torch.cuda.is_available():
                torch.cuda.empty_cache()
        except Exception:
            pass

    async def _ensure_design_engine_loaded(self) -> bool:
        """Готує VoiceDesign engine (lazy-load, з unload clone при нестачі пам'яті)."""
        if self._design_engine is None:
            return False
        if self._engine_is_loaded(self._design_engine):
            return True
        
        if self.config.unload_design_before_clone and self._engine_is_loaded(self._clone_engine):
            self._unload_engine(self._clone_engine, "VoiceClone")
        
        loaded = await self._load_engine(self._design_engine, "VoiceDesign")
        if not loaded and self.config.unload_design_before_clone:
            # Спроба повернути clone, якщо design не піднявся
            await self._load_engine(self._clone_engine, "VoiceClone")
        return loaded

    async def _ensure_clone_engine_loaded(self) -> bool:
        """Готує VoiceClone engine перед генерацією реплік."""
        if self._clone_engine is None:
            return False
        if self._engine_is_loaded(self._clone_engine):
            return True
        
        # Щоб уникнути OOM у Colab, перед clone вивантажуємо design
        if self.config.unload_design_before_clone and self._engine_is_loaded(self._design_engine):
            self._unload_engine(self._design_engine, "VoiceDesign")
        
        loaded = await self._load_engine(self._clone_engine, "VoiceClone")
        if not loaded and self.config.unload_design_before_clone:
            # Повертаємо design назад для fallback, якщо clone не піднявся
            await self._load_engine(self._design_engine, "VoiceDesign")
        return loaded

    async def _try_init_tts_from_all_in_one(self, errors: List[str]) -> bool:
        """Пробує ініціалізувати TTS через voice_library_all_in_one."""
        module = self._load_module_with_candidates(
            module_name="voice_library_all_in_one",
            filename="voice_library_all_in_one.py",
            errors=errors
        )
        if module is None:
            return False
        
        required = [
            "VoiceDesignEngine",
            "VoiceCloneEngine",
            "VoicePromptCache",
            "init_voice_library_colab"
        ]
        missing = [name for name in required if not hasattr(module, name)]
        if missing:
            errors.append(f"voice_library_all_in_one missing required symbols: {', '.join(missing)}")
            return False
        
        try:
            VoiceDesignEngine = getattr(module, "VoiceDesignEngine")
            VoiceCloneEngine = getattr(module, "VoiceCloneEngine")
            VoicePromptCache = getattr(module, "VoicePromptCache")
            init_voice_library_colab = getattr(module, "init_voice_library_colab")
            create_voice_from_reference = getattr(module, "create_voice_from_reference", None)
            
            voices_drive_path = os.path.join(self.config.drive_base_path, "voices")
            os.makedirs(voices_drive_path, exist_ok=True)
            
            library = init_voice_library_colab(
                drive_path=voices_drive_path,
                cache_dir=os.path.join(self.config.cache_dir, "voice_cache"),
                load_models=False
            )
            
            self._design_engine = VoiceDesignEngine()
            self._clone_engine = VoiceCloneEngine(
                cache_dir=os.path.join(self.config.cache_dir, "voice_cache")
            )
            self._library_manager = library
            self._cache = VoicePromptCache(
                cache_dir=os.path.join(self.config.cache_dir, "prompt_cache")
            )
            self._create_voice_from_reference_fn = create_voice_from_reference
            self._voice_library_entry_cls = None
            
            design_loaded = False
            if not self.config.lazy_load_design:
                logger.info("[Orchestrator] Loading VoiceDesign model (all_in_one)...")
                design_loaded = await self._load_engine(self._design_engine, "VoiceDesign")
            else:
                logger.info("[Orchestrator] VoiceDesign will be loaded lazily on first test-voice generation")
            clone_loaded = False
            if not self.config.lazy_load_clone:
                clone_loaded = await self._load_engine(self._clone_engine, "VoiceClone")
            else:
                logger.info("[Orchestrator] VoiceClone will be loaded lazily on first dialogue generation")
            
            if (not self.config.lazy_load_design and not self.config.lazy_load_clone
                and not design_loaded and not clone_loaded):
                errors.append("all_in_one imported, but both TTS models failed to load")
                return False
            
            if self._create_voice_from_reference_fn is None:
                self._create_voice_from_reference_fn = self._create_voice_from_reference_fallback
            
            logger.info("[Orchestrator] TTS initialized via voice_library_all_in_one")
            return True
        except Exception as e:
            errors.append(f"all_in_one init failed: {e}")
            logger.exception("[Orchestrator] all_in_one initialization exception")
            return False

    async def _try_init_tts_from_modular_stack(self, errors: List[str]) -> bool:
        """Пробує ініціалізувати TTS через voicebox_adapter_colab + voice_library_manager."""
        adapter_module = self._load_module_with_candidates(
            module_name="voicebox_adapter_colab",
            filename="voicebox_adapter_colab.py",
            errors=errors
        )
        manager_module = self._load_module_with_candidates(
            module_name="voice_library_manager",
            filename="voice_library_manager.py",
            errors=errors
        )
        if adapter_module is None or manager_module is None:
            return False
        
        required_adapter = ["VoiceDesignEngine", "VoiceCloneEngine", "VoicePromptCache"]
        missing_adapter = [name for name in required_adapter if not hasattr(adapter_module, name)]
        if missing_adapter:
            errors.append(f"voicebox_adapter_colab missing symbols: {', '.join(missing_adapter)}")
            return False
        
        if not hasattr(manager_module, "VoiceLibraryManager"):
            errors.append("voice_library_manager missing symbol: VoiceLibraryManager")
            return False
        
        try:
            VoiceDesignEngine = getattr(adapter_module, "VoiceDesignEngine")
            VoiceCloneEngine = getattr(adapter_module, "VoiceCloneEngine")
            VoicePromptCache = getattr(adapter_module, "VoicePromptCache")
            VoiceLibraryManager = getattr(manager_module, "VoiceLibraryManager")
            VoiceLibraryEntry = getattr(manager_module, "VoiceLibraryEntry", None)
            
            voices_drive_path = os.path.join(self.config.drive_base_path, "voices")
            os.makedirs(voices_drive_path, exist_ok=True)
            
            self._design_engine = VoiceDesignEngine()
            self._clone_engine = VoiceCloneEngine(
                cache_dir=os.path.join(self.config.cache_dir, "voice_cache")
            )
            self._cache = VoicePromptCache(
                cache_dir=os.path.join(self.config.cache_dir, "prompt_cache")
            )
            self._library_manager = VoiceLibraryManager(
                voice_clone_engine=self._clone_engine,
                voice_design_engine=self._design_engine,
                cache=self._cache,
                drive_path=voices_drive_path
            )
            self._voice_library_entry_cls = VoiceLibraryEntry
            self._create_voice_from_reference_fn = self._create_voice_from_reference_fallback
            
            if hasattr(self._library_manager, "load_library"):
                await asyncio.to_thread(self._library_manager.load_library)
            
            design_loaded = False
            if not self.config.lazy_load_design:
                logger.info("[Orchestrator] Loading VoiceDesign model (modular stack)...")
                design_loaded = await self._load_engine(self._design_engine, "VoiceDesign")
            else:
                logger.info("[Orchestrator] VoiceDesign will be loaded lazily on first test-voice generation")
            clone_loaded = False
            if not self.config.lazy_load_clone:
                clone_loaded = await self._load_engine(self._clone_engine, "VoiceClone")
            else:
                logger.info("[Orchestrator] VoiceClone will be loaded lazily on first dialogue generation")
            
            if (not self.config.lazy_load_design and not self.config.lazy_load_clone
                and not design_loaded and not clone_loaded):
                errors.append("modular stack imported, but both TTS models failed to load")
                return False
            
            logger.info("[Orchestrator] TTS initialized via modular stack")
            return True
        except Exception as e:
            errors.append(f"modular stack init failed: {e}")
            logger.exception("[Orchestrator] modular stack initialization exception")
            return False

    async def _try_init_tts_internal_engines(self, errors: List[str]) -> bool:
        """Пробує ініціалізувати TTS через внутрішні engine-класи цього файлу."""
        try:
            voices_drive_path = os.path.join(self.config.drive_base_path, "voices")
            os.makedirs(voices_drive_path, exist_ok=True)
            
            design_model = os.getenv("QWEN_VOICEDESIGN_MODEL", "Qwen/Qwen3-TTS-12Hz-1.7B-VoiceDesign")
            clone_model = os.getenv("QWEN_VOICECLONE_MODEL", "Qwen/Qwen3-TTS-12Hz-1.7B-Base")
            
            self._design_engine = InternalVoiceDesignEngine(model_name=design_model)
            self._clone_engine = InternalVoiceCloneEngine(
                model_name=clone_model,
                cache_dir=os.path.join(self.config.cache_dir, "voice_cache")
            )
            self._cache = None
            self._library_manager = SimpleVoiceLibraryManager(
                clone_engine=self._clone_engine,
                drive_path=voices_drive_path
            )
            self._voice_library_entry_cls = None
            self._create_voice_from_reference_fn = self._create_voice_from_reference_fallback
            
            self._library_manager.load_library()
            
            design_loaded = False
            if not self.config.lazy_load_design:
                logger.info("[Orchestrator] Loading VoiceDesign model (internal engines)...")
                design_loaded = await self._load_engine(self._design_engine, "VoiceDesign")
            else:
                logger.info("[Orchestrator] VoiceDesign will be loaded lazily on first test-voice generation")
            clone_loaded = False
            if not self.config.lazy_load_clone:
                clone_loaded = await self._load_engine(self._clone_engine, "VoiceClone")
            else:
                logger.info("[Orchestrator] VoiceClone will be loaded lazily on first dialogue generation")
            
            if (not self.config.lazy_load_design and not self.config.lazy_load_clone
                and not design_loaded and not clone_loaded):
                errors.append("internal engines initialized, but both models failed to load")
                return False
            
            logger.info("[Orchestrator] TTS initialized via internal engines")
            return True
        except Exception as e:
            errors.append(f"internal engines init failed: {e}")
            logger.exception("[Orchestrator] internal engines initialization exception")
            return False

    def _load_module_with_candidates(
        self,
        module_name: str,
        filename: str,
        errors: List[str]
    ):
        """Завантажує модуль з sys.path, а якщо не вдалось - з локальних файлів."""
        import importlib
        import importlib.util
        
        try:
            return importlib.import_module(module_name)
        except Exception as e:
            errors.append(f"import {module_name} failed: {e}")
        
        candidates = self._find_candidate_module_paths(filename)
        for file_path in candidates:
            try:
                spec = importlib.util.spec_from_file_location(module_name, file_path)
                if spec is None or spec.loader is None:
                    continue
                module = importlib.util.module_from_spec(spec)
                sys.modules[module_name] = module
                spec.loader.exec_module(module)
                logger.info(f"[Orchestrator] Loaded {module_name} from {file_path}")
                return module
            except Exception as e:
                errors.append(f"load {module_name} from {file_path} failed: {e}")
        
        return None

    def _find_candidate_module_paths(self, filename: str) -> List[str]:
        """Повертає список можливих шляхів до модуля у Colab/Drive."""
        paths: List[str] = []
        
        base_dirs = [
            os.getcwd(),
            os.path.dirname(os.path.abspath(__file__)) if "__file__" in globals() else "",
            self.config.drive_base_path,
            os.path.dirname(self.config.drive_base_path),
            "/content",
            "/content/drive/MyDrive",
        ]
        
        for base_dir in base_dirs:
            if not base_dir:
                continue
            candidate = os.path.join(base_dir, filename)
            if os.path.exists(candidate):
                paths.append(candidate)
        
        # Глибокий пошук у MyDrive (обмежений першими знахідками)
        drive_pattern = f"/content/drive/MyDrive/**/{filename}"
        for found in glob.glob(drive_pattern, recursive=True)[:20]:
            if os.path.exists(found):
                paths.append(found)
        
        # Унікалізація зі збереженням порядку
        unique_paths: List[str] = []
        seen = set()
        for path in paths:
            norm = os.path.normpath(path)
            if norm in seen:
                continue
            seen.add(norm)
            unique_paths.append(path)
        
        return unique_paths

    def _extract_offending_backend(self, error_message: str) -> Optional[str]:
        """Витягує назву backend з помилки transformers import_utils."""
        match = re.search(r"Offending backend:\s*([A-Za-z0-9_]+)", str(error_message or ""))
        if not match:
            return None
        token = (match.group(1) or "").strip().lower()
        return token or None

    def _patch_transformers_backend_mapping(
        self,
        backend_name: str = "keras_nlp",
        force_optional_true: bool = False
    ) -> bool:
        """Патчить BACKENDS_MAPPING transformers для сумісності qwen_tts."""
        backend = (backend_name or "").strip().lower()
        if not backend:
            return False

        module_name = {
            "tf": "tensorflow",
            "tensorflow": "tensorflow",
            "tf_text": "tensorflow_text",
            "tensorflowtext": "tensorflow_text",
        }.get(backend, backend)

        try:
            from transformers.utils import import_utils as tr_import_utils
        except Exception:
            return False

        mapping = getattr(tr_import_utils, "BACKENDS_MAPPING", None)
        if mapping is None:
            return False

        default_message = (
            f"{{0}} requires the {backend} library but it was not found in your environment. "
            f"Install it with `pip install {module_name}`."
        )

        if backend in mapping:
            if not force_optional_true:
                return False
            try:
                existing = mapping.get(backend)
                message = (
                    existing[1]
                    if isinstance(existing, (tuple, list)) and len(existing) > 1
                    else default_message
                )
                mapping[backend] = (lambda: True, message)
                logger.warning(
                    f"[Orchestrator] Forced transformers backend {backend} as optional=True"
                )
                return True
            except Exception as e:
                logger.warning(f"[Orchestrator] Failed to force backend {backend}: {e}")
                return False

        def _backend_available() -> bool:
            if force_optional_true:
                return True
            try:
                return importlib.util.find_spec(module_name) is not None
            except Exception:
                return False

        try:
            mapping[backend] = (_backend_available, default_message)
            logger.warning(
                f"[Orchestrator] Patched transformers BACKENDS_MAPPING with {backend} "
                f"(force_optional_true={force_optional_true})"
            )
            return True
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to patch backend mapping for {backend}: {e}")
            return False

    def _patch_transformers_offline_mode_compat(self) -> bool:
        """Додає сумісність для transformers.utils.is_offline_mode (прибрано в нових релізах)."""
        try:
            import transformers
            from transformers import utils as tr_utils
        except Exception:
            return False

        if hasattr(tr_utils, "is_offline_mode"):
            return False

        resolver = None
        try:
            from transformers.utils.hub import is_offline_mode as hub_is_offline_mode
            resolver = hub_is_offline_mode
        except Exception:
            resolver = None

        if resolver is None:
            def _env_offline_mode() -> bool:
                markers = {"1", "true", "yes", "y", "on"}
                hf = str(os.getenv("HF_HUB_OFFLINE", "0")).strip().lower() in markers
                tr = str(os.getenv("TRANSFORMERS_OFFLINE", "0")).strip().lower() in markers
                return hf or tr
            resolver = _env_offline_mode

        try:
            setattr(tr_utils, "is_offline_mode", resolver)
            try:
                if hasattr(transformers, "utils"):
                    setattr(transformers.utils, "is_offline_mode", resolver)
            except Exception:
                pass
            logger.warning("[Orchestrator] Patched transformers.utils.is_offline_mode compatibility symbol")
            return True
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to patch transformers.utils.is_offline_mode: {e}")
            return False

    def _patch_transformers_download_url_compat(self) -> bool:
        """Додає сумісність для transformers.utils.download_url (прибрано в нових релізах)."""
        try:
            import transformers
            from transformers import utils as tr_utils
        except Exception:
            return False

        if hasattr(tr_utils, "download_url"):
            return False

        def _compat_download_url(url: str, proxies: Optional[Dict[str, str]] = None) -> str:
            if not url:
                raise ValueError("download_url expects a non-empty url")

            from urllib.parse import urlparse
            from urllib.request import Request, ProxyHandler, build_opener, urlopen

            parsed = urlparse(url)
            file_name = os.path.basename(parsed.path or "")
            if not file_name:
                file_name = f"download_{hashlib.md5(url.encode('utf-8')).hexdigest()[:16]}.bin"
            cache_dir = os.path.join(tempfile.gettempdir(), "transformers_download_url_cache")
            os.makedirs(cache_dir, exist_ok=True)
            out_path = os.path.join(cache_dir, file_name)
            if os.path.exists(out_path) and os.path.getsize(out_path) > 0:
                return out_path

            request = Request(url, headers={"User-Agent": "transformers-compat"})
            if proxies:
                opener = build_opener(ProxyHandler(proxies))
                with opener.open(request, timeout=180) as response, open(out_path, "wb") as output:
                    shutil.copyfileobj(response, output)
            else:
                with urlopen(request, timeout=180) as response, open(out_path, "wb") as output:
                    shutil.copyfileobj(response, output)
            return out_path

        try:
            setattr(tr_utils, "download_url", _compat_download_url)
            try:
                if hasattr(transformers, "utils"):
                    setattr(transformers.utils, "download_url", _compat_download_url)
            except Exception:
                pass
            logger.warning("[Orchestrator] Patched transformers.utils.download_url compatibility symbol")
            return True
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to patch transformers.utils.download_url: {e}")
            return False

    def _patch_transformers_is_remote_url_compat(self) -> bool:
        """Додає сумісність для transformers.utils.is_remote_url."""
        try:
            import transformers
            from transformers import utils as tr_utils
        except Exception:
            return False

        if hasattr(tr_utils, "is_remote_url"):
            return False

        resolver = None
        try:
            from transformers.utils.hub import is_remote_url as hub_is_remote_url
            resolver = hub_is_remote_url
        except Exception:
            resolver = None

        if resolver is None:
            def _compat_is_remote_url(url_or_filename: Any) -> bool:
                if url_or_filename is None:
                    return False
                try:
                    from urllib.parse import urlparse
                    parsed = urlparse(str(url_or_filename))
                    return parsed.scheme in {"http", "https", "ftp", "ftps", "s3", "gs"}
                except Exception:
                    return False
            resolver = _compat_is_remote_url

        try:
            setattr(tr_utils, "is_remote_url", resolver)
            try:
                if hasattr(transformers, "utils"):
                    setattr(transformers.utils, "is_remote_url", resolver)
            except Exception:
                pass
            logger.warning("[Orchestrator] Patched transformers.utils.is_remote_url compatibility symbol")
            return True
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to patch transformers.utils.is_remote_url: {e}")
            return False

    def _runtime_debug_enabled(self) -> bool:
        """Повертає True, якщо увімкнено розширений runtime debug."""
        value = str(os.getenv("AUDIO_DRAMA_RUNTIME_DEBUG", "0")).strip().lower()
        return value in {"1", "true", "yes", "on"}

    def _purge_modules_by_prefix(self, prefixes: Tuple[str, ...]) -> int:
        """Видаляє модулі з sys.modules за префіксами."""
        removed = 0
        normalized = tuple(p for p in prefixes if p)
        if not normalized:
            return 0
        for module_name in list(sys.modules.keys()):
            for prefix in normalized:
                if module_name == prefix or module_name.startswith(prefix + "."):
                    sys.modules.pop(module_name, None)
                    removed += 1
                    break
        try:
            import importlib
            importlib.invalidate_caches()
        except Exception:
            pass
        return removed

    def _cleanup_transformers_ipynb_checkpoints(self) -> int:
        """Прибирає .ipynb_checkpoints з transformers пакета (відомий баг імпортів)."""
        try:
            import transformers
            base_dir = os.path.dirname(str(getattr(transformers, "__file__", "") or ""))
        except Exception:
            return 0
        if not base_dir or not os.path.isdir(base_dir):
            return 0

        removed = 0
        try:
            for root, dirs, _ in os.walk(base_dir):
                for d in list(dirs):
                    if d == ".ipynb_checkpoints":
                        target = os.path.join(root, d)
                        try:
                            shutil.rmtree(target, ignore_errors=True)
                            removed += 1
                        except Exception:
                            pass
        except Exception:
            return removed

        if removed > 0:
            logger.warning(
                f"[Orchestrator] Removed {removed} '.ipynb_checkpoints' dirs from transformers package"
            )
        return removed

    def _diagnose_transformers_automodel_import(self) -> str:
        """Повертає детальний traceback для проблем імпорту AutoModel."""
        snippet = (
            "import traceback\n"
            "print('[diag] python ok')\n"
            "try:\n"
            "    import transformers\n"
            "    print('[diag] transformers', getattr(transformers, '__version__', '?'))\n"
            "except Exception:\n"
            "    traceback.print_exc()\n"
            "try:\n"
            "    from transformers import AutoModel\n"
            "    print('[diag] AutoModel ok', AutoModel)\n"
            "except Exception:\n"
            "    traceback.print_exc()\n"
            "try:\n"
            "    from transformers.modeling_utils import PreTrainedModel\n"
            "    print('[diag] PreTrainedModel ok', PreTrainedModel)\n"
            "except Exception:\n"
            "    traceback.print_exc()\n"
            "try:\n"
            "    import torchvision\n"
            "    print('[diag] torchvision', getattr(torchvision, '__version__', '?'))\n"
            "except Exception:\n"
            "    traceback.print_exc()\n"
        )
        try:
            result = subprocess.run(
                [sys.executable, "-c", snippet],
                check=False,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                timeout=120
            )
            output = (result.stdout or "").strip()
            return output[-5000:] if output else f"diagnostic empty (rc={result.returncode})"
        except Exception as e:
            return f"diagnostic failed: {e}"

    def _looks_like_torchvision_mismatch(self, message: str) -> bool:
        value = str(message or "").lower()
        markers = [
            "torchvision::nms",
            "operator torchvision",
            "interpolationmode.nearest_exact",
            "attributeerror: nearest_exact",
            "detected that pytorch and torchvision were compiled with different cuda",
            "could not import module 'pretrainedmodel'",
            "could not import module 'automodel'",
        ]
        if any(marker in value for marker in markers):
            return True
        if "torchvision" in value and ("runtimeerror" in value or "importerror" in value):
            return True
        return False

    def _attempt_torchvision_text_only_workaround(self) -> bool:
        """
        Workaround з форумів HF: для text-only задач можна прибрати torchvision,
        якщо він ламає імпорт transformers через nms/InterpolationMode.
        """
        current_tv = self._get_installed_package_version("torchvision")
        if not current_tv:
            return False
        logger.warning(
            f"[Orchestrator] Trying text-only workaround: uninstall torchvision ({current_tv}) "
            "to unblock transformers AutoModel import"
        )
        ok, err = self._run_install_step(
            [sys.executable, "-m", "pip", "uninstall", "-y", "torchvision"],
            "uninstall torchvision workaround",
            timeout_sec=300,
            stream_output=self._runtime_debug_enabled()
        )
        if not ok:
            logger.warning(f"[Orchestrator] torchvision workaround failed: {err}")
            return False
        removed = self._purge_modules_by_prefix(("torchvision", "transformers"))
        logger.warning(f"[Orchestrator] torchvision workaround applied (purged modules: {removed})")
        return True

    def _validate_transformers_automodel_runtime(self) -> Tuple[bool, str]:
        """Перевіряє базову працездатність transformers AutoModel."""
        try:
            import transformers  # noqa: F401
            from transformers import AutoModel  # noqa: F401
            return True, ""
        except Exception as e:
            detail = f"{type(e).__name__}: {e}"
            if self._runtime_debug_enabled():
                diagnostic = self._diagnose_transformers_automodel_import()
                if diagnostic:
                    detail = f"{detail} | {diagnostic}"
            return False, detail

    def _log_qwen_runtime_diagnostics(self, stage: str, extra: str = "") -> None:
        """Логує діагностику runtime для швидкого дебагу у Colab."""
        if not self._runtime_debug_enabled():
            return
        try:
            package_versions: Dict[str, str] = {}
            for pkg in ("transformers", "accelerate", "qwen-tts", "torch", "torchvision", "huggingface-hub"):
                package_versions[pkg] = self._get_installed_package_version(pkg) or "not-installed"

            utils_flags: Dict[str, str] = {}
            transformers_path = "unknown"
            try:
                import transformers
                transformers_path = str(getattr(transformers, "__file__", "unknown"))
                from transformers import utils as tr_utils
                for symbol in ("download_url", "is_offline_mode", "is_jax_tensor", "is_remote_url"):
                    utils_flags[symbol] = "1" if hasattr(tr_utils, symbol) else "0"
            except Exception as e:
                utils_flags["import_error"] = str(e)

            logger.warning(
                "[Orchestrator][RuntimeDebug] "
                f"stage={stage}; py={sys.version.split()[0]}; "
                f"versions={package_versions}; "
                f"transformers_path={transformers_path}; "
                f"utils={utils_flags}"
                + (f"; extra={extra}" if extra else "")
            )
        except Exception as e:
            logger.warning(f"[Orchestrator][RuntimeDebug] stage={stage}; diagnostic failed: {e}")

    def _extract_missing_transformers_utils_symbol(self, error_message: str) -> Optional[str]:
        """Витягує назву відсутнього символу з ImportError по transformers.utils."""
        match = re.search(
            (
                r"cannot import name ['\"]([A-Za-z0-9_]+)['\"] from "
                r"['\"]transformers\.utils(?:\.import_utils)?['\"]"
            ),
            str(error_message or "")
        )
        if not match:
            return None
        token = (match.group(1) or "").strip()
        return token or None

    def _patch_transformers_utils_symbol_compat(self, symbol_name: str) -> bool:
        """Патчить відсутній символ у transformers.utils для старого qwen_tts коду."""
        symbol = (symbol_name or "").strip()
        if not symbol:
            return False

        if symbol == "is_offline_mode":
            return self._patch_transformers_offline_mode_compat()
        if symbol == "download_url":
            return self._patch_transformers_download_url_compat()
        if symbol == "is_remote_url":
            return self._patch_transformers_is_remote_url_compat()

        try:
            import transformers
            from transformers import utils as tr_utils
            from transformers.utils import import_utils as tr_import_utils
        except Exception:
            return False

        if hasattr(tr_utils, symbol) and hasattr(tr_import_utils, symbol):
            return False

        resolver = None

        # 1) Найкращий варіант: символ ще є в import_utils.
        candidate = getattr(tr_import_utils, symbol, None)
        if candidate is not None:
            resolver = candidate

        # 2) Fallback для приватних *_version (напр. _torch_version).
        if resolver is None and symbol.startswith("_") and symbol.endswith("_version"):
            token = symbol[1:-len("_version")]
            package_aliases = {
                "torch": "torch",
                "tf": "tensorflow",
                "tensorflow": "tensorflow",
                "flax": "flax",
                "jax": "jax",
                "numpy": "numpy",
                "np": "numpy",
                "vision": "Pillow",
                "pil": "Pillow",
            }
            package_name = package_aliases.get(token, token.replace("_", "-"))
            version_value = self._get_installed_package_version(package_name)
            if not version_value and token == "torch":
                try:
                    import torch
                    version_value = str(getattr(torch, "__version__", "") or "")
                except Exception:
                    version_value = ""
            if not version_value and token in {"tf", "tensorflow"}:
                try:
                    import tensorflow as tf
                    version_value = str(getattr(tf, "__version__", "") or "")
                except Exception:
                    version_value = ""
            resolver = version_value or "0.0.0"

        # 3) Fallback для приватних *_available (напр. _torch_available).
        if resolver is None and symbol.startswith("_") and symbol.endswith("_available"):
            token = symbol[1:-len("_available")]
            module_aliases = {
                "tf": "tensorflow",
                "tensorflow": "tensorflow",
                "torch": "torch",
                "flax": "flax",
                "jax": "jax",
                "vision": "PIL",
                "pil": "PIL",
                "numpy": "numpy",
            }
            module_name = module_aliases.get(token, token)

            def _private_available(mod: str = module_name) -> bool:
                try:
                    return importlib.util.find_spec(mod) is not None
                except Exception:
                    return False

            resolver = _private_available

        # 4) Евристичний fallback для is_*_available.
        if resolver is None and symbol.startswith("is_") and symbol.endswith("_available"):
            module_token = symbol[len("is_"):-len("_available")]
            module_aliases = {
                "tf": "tensorflow",
                "flax": "flax",
                "jax": "jax",
                "torch": "torch",
                "vision": "PIL",
                "sentencepiece": "sentencepiece",
                "tokenizers": "tokenizers",
                "safetensors": "safetensors",
                "datasets": "datasets",
            }
            module_name = module_aliases.get(module_token, module_token)

            if symbol == "is_flax_available":
                def _is_flax_available() -> bool:
                    try:
                        return (
                            importlib.util.find_spec("jax") is not None
                            and importlib.util.find_spec("flax") is not None
                        )
                    except Exception:
                        return False
                resolver = _is_flax_available
            else:
                def _generic_is_available(mod: str = module_name) -> bool:
                    try:
                        return importlib.util.find_spec(mod) is not None
                    except Exception:
                        return False
                resolver = _generic_is_available

        # 5) Евристичний fallback для is_*_tensor (напр. is_jax_tensor).
        if resolver is None and symbol.startswith("is_") and symbol.endswith("_tensor"):
            tensor_token = symbol[len("is_"):-len("_tensor")]
            tensor_aliases = {
                "tf": "tensorflow",
                "tensorflow": "tensorflow",
                "torch": "torch",
                "jax": "jax",
                "numpy": "numpy",
                "np": "numpy",
            }
            module_prefix = tensor_aliases.get(tensor_token, tensor_token)

            if symbol == "is_jax_tensor":
                def _is_jax_tensor(value: Any = None) -> bool:
                    if value is None:
                        return False
                    # Швидка перевірка по module path (без імпорту jax).
                    try:
                        cls_module = str(getattr(getattr(value, "__class__", None), "__module__", "") or "")
                        if cls_module.startswith("jax.") or cls_module.startswith("jaxlib."):
                            return True
                    except Exception:
                        pass
                    # Додаткова перевірка якщо jax встановлено.
                    try:
                        import jax
                        import jax.numpy as jnp
                        if hasattr(jax, "Array") and isinstance(value, jax.Array):
                            return True
                        return isinstance(value, jnp.ndarray)
                    except Exception:
                        return False
                resolver = _is_jax_tensor
            else:
                def _generic_is_tensor(value: Any = None, prefix: str = module_prefix) -> bool:
                    if value is None:
                        return False
                    try:
                        cls_module = str(getattr(getattr(value, "__class__", None), "__module__", "") or "")
                    except Exception:
                        return False
                    if not cls_module:
                        return False
                    return cls_module == prefix or cls_module.startswith(prefix + ".")
                resolver = _generic_is_tensor

        if resolver is None:
            return False

        try:
            setattr(tr_utils, symbol, resolver)
            try:
                if hasattr(transformers, "utils"):
                    setattr(transformers.utils, symbol, resolver)
            except Exception:
                pass
            try:
                if not hasattr(tr_import_utils, symbol):
                    setattr(tr_import_utils, symbol, resolver)
                    logger.warning(
                        f"[Orchestrator] Patched transformers.utils.import_utils.{symbol} compatibility symbol"
                    )
            except Exception:
                pass
            logger.warning(f"[Orchestrator] Patched transformers.utils.{symbol} compatibility symbol")
            return True
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to patch transformers.utils.{symbol}: {e}")
            return False

    def _patch_transformers_utils_exports_compat(self) -> int:
        """Підкладає відсутні export-и з import_utils у transformers.utils."""
        try:
            import transformers
            from transformers import utils as tr_utils
            from transformers.utils import import_utils as tr_import_utils
        except Exception:
            return 0

        patched = 0
        try:
            candidates = []
            for name in dir(tr_import_utils):
                if name.startswith("is_") and (name.endswith("_available") or name.endswith("_tensor")):
                    candidates.append(name)
            # Базові aliases, які часто чекає старий код.
            candidates.extend(["is_offline_mode", "download_url", "is_jax_tensor", "is_remote_url"])
            for symbol in sorted(set(candidates)):
                if hasattr(tr_utils, symbol):
                    continue
                if self._patch_transformers_utils_symbol_compat(symbol):
                    patched += 1
        except Exception:
            return patched
        return patched

    def _try_import_qwen_tts_runtime(self) -> Tuple[bool, Optional[str]]:
        """Пробує імпорт qwen_tts з патчем сумісності transformers."""
        healthy, health_detail = self._validate_transformers_automodel_runtime()
        if not healthy:
            lowered_health = str(health_detail or "").lower()
            if "ipynb_checkpoints" in lowered_health:
                removed = self._cleanup_transformers_ipynb_checkpoints()
                if removed > 0:
                    self._purge_modules_by_prefix(("transformers",))
                    healthy, health_detail = self._validate_transformers_automodel_runtime()
            if (not healthy) and self._looks_like_torchvision_mismatch(health_detail):
                if self._attempt_torchvision_text_only_workaround():
                    healthy, health_detail = self._validate_transformers_automodel_runtime()
            if not healthy:
                self._log_qwen_runtime_diagnostics("automodel_precheck_failed", extra=health_detail)
                return False, f"transformers AutoModel precheck failed: {health_detail}"

        # Підкладаємо найтиповіші backends ще до першої спроби імпорту.
        self._patch_transformers_utils_exports_compat()
        self._patch_transformers_utils_symbol_compat("is_offline_mode")
        self._patch_transformers_utils_symbol_compat("download_url")
        self._patch_transformers_utils_symbol_compat("is_remote_url")
        self._patch_transformers_utils_symbol_compat("is_flax_available")
        self._patch_transformers_utils_symbol_compat("is_jax_tensor")
        self._patch_transformers_backend_mapping("keras_nlp", force_optional_true=False)
        self._patch_transformers_backend_mapping("tensorflow_text", force_optional_true=False)

        forced_backends = set()
        last_error: Optional[str] = None

        for attempt_idx in range(6):
            try:
                # Після невдалих імпортів очищаємо partial-модулі для повторної спроби.
                for module_name in list(sys.modules.keys()):
                    if module_name == "qwen_tts" or module_name.startswith("qwen_tts."):
                        sys.modules.pop(module_name, None)
                import qwen_tts  # noqa: F401
                return True, None
            except Exception as error:
                message = str(error or "")
                last_error = message
                if self._runtime_debug_enabled():
                    logger.warning(
                        "[Orchestrator][RuntimeDebug] "
                        f"qwen_tts import attempt {attempt_idx + 1}/6 failed: "
                        f"{type(error).__name__}: {message}"
                    )
                missing_symbol = self._extract_missing_transformers_utils_symbol(message)
                if missing_symbol:
                    guard_key = f"utils_symbol:{missing_symbol}"
                    if guard_key in forced_backends:
                        return False, message
                    if self._patch_transformers_utils_symbol_compat(missing_symbol):
                        forced_backends.add(guard_key)
                        continue
                    return False, message
                offending_backend = self._extract_offending_backend(message)
                if not offending_backend:
                    return False, message
                if offending_backend in forced_backends:
                    return False, message
                # Менш агресивний підхід: не "підмальовуємо" core backends як installed.
                if offending_backend in {"tf", "tensorflow", "torch", "torchvision"}:
                    return False, message
                patched = self._patch_transformers_backend_mapping(
                    offending_backend,
                    force_optional_true=False
                )
                if not patched:
                    return False, message
                forced_backends.add(offending_backend)

        return False, last_error

    def _get_installed_package_version(self, package_name: str) -> Optional[str]:
        """Повертає встановлену версію пакета або None."""
        try:
            try:
                from importlib import metadata as importlib_metadata
            except Exception:
                import importlib_metadata  # type: ignore
            return str(importlib_metadata.version(package_name))
        except Exception:
            return None

    def _run_install_step(
        self,
        cmd: List[str],
        step_name: str,
        timeout_sec: int,
        stream_output: bool = True
    ) -> Tuple[bool, str]:
        """Запускає install-команду з таймаутом і базовим логуванням."""
        safe_timeout = max(60, int(timeout_sec or 0))
        started_at = datetime.now()
        logger.warning(
            f"[Orchestrator] [{step_name}] Running: {' '.join(cmd)} "
            f"(timeout={safe_timeout}s, stream_output={stream_output})"
        )

        run_kwargs: Dict[str, Any] = {
            "check": False,
            "timeout": safe_timeout,
        }
        if stream_output:
            run_kwargs["stdout"] = None
            run_kwargs["stderr"] = None
        else:
            run_kwargs["stdout"] = subprocess.PIPE
            run_kwargs["stderr"] = subprocess.STDOUT
            run_kwargs["text"] = True

        try:
            result = subprocess.run(cmd, **run_kwargs)
        except subprocess.TimeoutExpired:
            elapsed_sec = (datetime.now() - started_at).total_seconds()
            return False, f"timeout after {elapsed_sec:.1f}s"
        except Exception as e:
            return False, str(e)

        elapsed_sec = (datetime.now() - started_at).total_seconds()
        if result.returncode != 0:
            error_tail = ""
            if not stream_output:
                captured = getattr(result, "stdout", "") or ""
                if captured:
                    lines = [line.strip() for line in captured.splitlines() if line.strip()]
                    if lines:
                        error_tail = f"; tail: {lines[-1][:220]}"
            return False, f"return code {result.returncode} after {elapsed_sec:.1f}s{error_tail}"

        logger.info(f"[Orchestrator] [{step_name}] Completed in {elapsed_sec:.1f}s")
        return True, ""

    def _ensure_qwen_tts_runtime(self, allow_install: bool = True) -> bool:
        """Гарантує наявність qwen-tts рантайму."""
        self._log_qwen_runtime_diagnostics("ensure_runtime_start")
        imported, import_err = self._try_import_qwen_tts_runtime()
        if imported:
            self._qwen_runtime_install_error = ""
            self._log_qwen_runtime_diagnostics("ensure_runtime_already_ready")
            return True
        if not allow_install:
            raise RuntimeError(f"qwen_tts is not installed: {import_err}")
        if self._qwen_runtime_install_attempted:
            raise RuntimeError(
                "qwen-tts runtime install already attempted in this session: "
                f"{self._qwen_runtime_install_error or import_err}"
            )
        self._qwen_runtime_install_attempted = True
        logger.warning("[Orchestrator] qwen_tts not found, installing runtime (one-time)...")
        self._log_qwen_runtime_diagnostics("before_install", extra=str(import_err or ""))

        raw_timeout = os.getenv("AUDIO_DRAMA_QWEN_INSTALL_TIMEOUT_SEC", "240")
        try:
            install_timeout_sec = int(raw_timeout)
        except Exception:
            install_timeout_sec = 900
        install_timeout_sec = max(120, install_timeout_sec)

        stream_output = os.getenv("AUDIO_DRAMA_QWEN_SHOW_PIP_OUTPUT", "1").strip().lower() in {
            "1", "true", "yes", "on"
        }

        install_errors: List[str] = []
        target_transformers = "4.57.3"
        target_accelerate = "1.12.0"

        def _import_check(step_name: str) -> bool:
            ok, retry_err = self._try_import_qwen_tts_runtime()
            if ok:
                logger.info("[Orchestrator] qwen_tts runtime ready")
                self._qwen_runtime_install_error = ""
                return True
            install_errors.append(f"{step_name}: import check failed ({retry_err})")
            return False

        # 1) Спочатку ставимо сам qwen-tts (найшвидший крок).
        qwen_cmd = [
            sys.executable, "-m", "pip", "install",
            "--no-cache-dir", "--upgrade",
            "qwen-tts==0.1.1"
        ]
        ok, error = self._run_install_step(
            qwen_cmd,
            "qwen-tts==0.1.1",
            install_timeout_sec,
            stream_output=stream_output
        )
        self._log_qwen_runtime_diagnostics("after_qwen_pip", extra=f"ok={ok}; err={error}")
        if ok and _import_check("qwen-tts==0.1.1"):
            return True
        if not ok:
            install_errors.append(f"qwen-tts==0.1.1: {error}")

        # 2) Синхронізуємо transformers/accelerate тільки коли треба.
        installed_transformers = self._get_installed_package_version("transformers")
        installed_accelerate = self._get_installed_package_version("accelerate")
        needs_base_sync = (
            installed_transformers != target_transformers or
            installed_accelerate != target_accelerate
        )
        if needs_base_sync:
            logger.warning(
                "[Orchestrator] Syncing runtime deps: "
                f"transformers={installed_transformers or 'none'} -> {target_transformers}, "
                f"accelerate={installed_accelerate or 'none'} -> {target_accelerate}"
            )
            base_cmd = [
                sys.executable, "-m", "pip", "install",
                "--no-cache-dir", "--upgrade",
                f"transformers=={target_transformers}",
                f"accelerate=={target_accelerate}"
            ]
        else:
            logger.warning(
                "[Orchestrator] transformers/accelerate already pinned, "
                "forcing reinstall because qwen_tts import still fails"
            )
            base_cmd = [
                sys.executable, "-m", "pip", "install",
                "--no-cache-dir", "--upgrade", "--force-reinstall",
                f"transformers=={target_transformers}",
                f"accelerate=={target_accelerate}"
            ]

        ok, error = self._run_install_step(
            base_cmd,
            "sync transformers+accelerate",
            install_timeout_sec,
            stream_output=stream_output
        )
        self._log_qwen_runtime_diagnostics("after_transformers_sync", extra=f"ok={ok}; err={error}")
        if ok and _import_check("sync transformers+accelerate"):
            return True
        if not ok:
            install_errors.append(f"sync transformers+accelerate: {error}")

        # 3) Fallback: git-версія qwen-tts без deps.
        git_cmd = [
            sys.executable, "-m", "pip", "install",
            "--no-cache-dir", "--upgrade", "--force-reinstall", "--no-deps",
            "git+https://github.com/QwenLM/Qwen3-TTS.git"
        ]
        ok, error = self._run_install_step(
            git_cmd,
            "qwen-tts git fallback",
            install_timeout_sec,
            stream_output=stream_output
        )
        self._log_qwen_runtime_diagnostics("after_git_fallback", extra=f"ok={ok}; err={error}")
        if ok and _import_check("qwen-tts git fallback"):
            return True
        if not ok:
            install_errors.append(f"qwen-tts git fallback: {error}")

        self._log_qwen_runtime_diagnostics("install_failed", extra=" | ".join(install_errors))
        self._qwen_runtime_install_error = f"Failed to install qwen-tts runtime: {' | '.join(install_errors)}"
        raise RuntimeError(self._qwen_runtime_install_error)

    def _create_voice_from_reference_fallback(
        self,
        character_name: str,
        reference_audio_path: str,
        reference_text: str,
        voice_preset: str = "custom",
        gender: str = "male",
        language: str = "russian"
    ) -> Tuple[Optional[Dict[str, Any]], Optional[str]]:
        """Fallback-реалізація create_voice_from_reference для модульного стеку."""
        if self._clone_engine is None:
            return None, "VoiceClone engine not initialized"
        
        try:
            voice_prompt, _ = self._clone_engine.create_voice_prompt(
                reference_audio_path=reference_audio_path,
                reference_text=reference_text,
                use_cache=True,
                validate=True
            )
            if voice_prompt is None:
                return None, "Failed to create voice prompt from reference"
            
            if self._library_manager is None:
                return voice_prompt, None
            
            # Lightweight manager path (SimpleVoiceLibraryManager)
            if hasattr(self._library_manager, "register_reference"):
                self._library_manager.register_reference(
                    character_name=character_name,
                    reference_audio_path=reference_audio_path,
                    reference_text=reference_text,
                    voice_preset=voice_preset,
                    gender=gender,
                    language=language
                )
                return voice_prompt, None
            
            if self._voice_library_entry_cls is None:
                if hasattr(self._library_manager, "save_library"):
                    self._library_manager.save_library()
                return voice_prompt, None
            
            quality_score = 0.85
            try:
                if hasattr(self._library_manager, "validate_reference"):
                    is_valid, error_msg, quality = self._library_manager.validate_reference(
                        reference_audio_path,
                        strict=False
                    )
                    quality_score = float(quality) if quality is not None else quality_score
                    if not is_valid and error_msg:
                        logger.warning(f"[Orchestrator] Reference validation warning for {character_name}: {error_msg}")
            except Exception as e:
                logger.warning(f"[Orchestrator] Reference validation failed for {character_name}: {e}")
            
            duration = 0.0
            if hasattr(self._library_manager, "_calculate_duration"):
                try:
                    duration = float(self._library_manager._calculate_duration(reference_audio_path))
                except Exception:
                    duration = 0.0
            
            drive_path = getattr(self._library_manager, "_drive_path", "")
            rel_path = reference_audio_path
            if drive_path:
                try:
                    rel_path = os.path.relpath(reference_audio_path, drive_path)
                except Exception:
                    rel_path = reference_audio_path
            
            entry = self._voice_library_entry_cls(
                character_name=character_name,
                voice_preset=voice_preset,
                reference_audio_path=rel_path,
                reference_text=reference_text,
                duration=duration,
                quality_score=quality_score,
                created_at=datetime.now().isoformat(),
                updated_at=datetime.now().isoformat(),
                metadata={"gender": gender, "language": language, "source": "fallback_reference"}
            )
            
            if hasattr(self._library_manager, "_library"):
                self._library_manager._library[character_name] = entry
            if hasattr(self._library_manager, "save_library"):
                self._library_manager.save_library()
            
            return voice_prompt, None
        except Exception as e:
            return None, str(e)
    
    def _create_fallback_engines(self) -> None:
        """Створює fallback TTS двигуни."""
        # Простий fallback без зовнішніх залежностей
        self._design_engine = None
        self._clone_engine = None
        self._library_manager = None
        self._cache = None
        self._create_voice_from_reference_fn = None
        self._voice_library_entry_cls = None
        self._voice_manager = VoiceManager(
            create_from_reference_fn=None,
            text_preprocessor=self._prepare_tts_text
        )
        logger.warning("[Orchestrator] Using fallback voice manager (no TTS)")
        logger.warning(
            "[Orchestrator] Ensure internet access and qwen-tts/transformers install, or upload "
            "voice_library_all_in_one.py / voicebox_adapter_colab.py + voice_library_manager.py to Colab/Drive."
        )
    
    def _resolve_project_title(self, file_path: str, source_name: Optional[str] = None) -> str:
        """Повертає читабельну назву проекту з оригінального імені файлу."""
        raw_name = (source_name or "").strip()
        candidate = Path(raw_name).stem if raw_name else Path(file_path).stem
        candidate = re.sub(r"\s+", " ", candidate).strip(" _-.")

        # Якщо назва схожа на тимчасовий файл, пробуємо fallback з file_path.
        if not candidate or re.fullmatch(r"tmp[\w-]{4,}", candidate.lower()):
            fallback = Path(file_path).stem
            candidate = re.sub(r"\s+", " ", fallback).strip(" _-.")

        if not candidate:
            candidate = f"audio_drama_{datetime.now().strftime('%Y%m%d')}"
        return candidate

    async def process_uploaded_file(
        self,
        file_path: str,
        user_id: int,
        source_name: Optional[str] = None
    ) -> Optional[str]:
        """Обробляє завантажений файл та створює аудіоп'єсу."""
        if not self._is_initialized:
            logger.error("[Orchestrator] Not initialized")
            return None
        
        logger.info(f"[Orchestrator] Processing file: {file_path}")
        self._stop_requested = False
        self._checkpoint_targets_sent.clear()
        self._checkpoints.clear()
        self._checkpoint_resume_event.set()
        self._waiting_for_checkpoint_resume = False
        self._pre_voice_reference_event.set()
        self._waiting_for_pre_voice_reference = False
        self._voice_approvals.clear()
        self._voice_approval_event.clear()
        self._voice_update_event.clear()
        self._current_approval_character = None
        for task in self._reference_render_tasks.values():
            if task and not task.done():
                task.cancel()
        self._reference_render_tasks.clear()
        self._reference_render_errors.clear()
        self._preview_generation_active = False
        self._last_progress_percent = -1
        self._last_progress_min_bucket = -1
        self._dirty_characters.clear()
        self._last_generation_stats = {
            "total_dialogues": 0,
            "reused_dialogues": 0,
            "regenerated_dialogues": 0,
        }
        if self._bot:
            self._bot.reset_progress_message()
        
        try:
            # 1. Парсимо файл
            text = FileParser.auto_detect_and_parse(file_path)
            title = self._resolve_project_title(file_path, source_name=source_name)
            self._current_language = detect_text_language(text, normalize_language_code(self.config.language, "russian"))
            await self._send_progress(f"🌐 Визначена мова озвучки: {self._current_language}")
            
            # Створюємо робочу директорію
            self._current_work_dir = self._create_work_directory(title)
            self._save_generation_state(status="running", reason="new_file_processing")
            self._write_resume_pointer(self._current_work_dir, status="running", reason="new_file_processing")
            
            # 2. Генеруємо сценарій
            await self._send_progress("📝 Конвертація у сценарій...")
            scenario = await self.generate_scenario(text, title)
            self._current_scenario = scenario
            
            # Зберігаємо сценарій
            self._save_scenario(scenario)
            
            # 3. Витягуємо персонажів
            await self._send_progress("👥 Аналіз персонажів...")
            if self._gemini_client:
                characters = await self._gemini_client.extract_characters(
                    text,
                    language=self.get_current_language()
                )
            else:
                characters = []
            
            characters = self._ensure_characters_for_scenario(scenario, characters)
            scenario.characters = characters
            require_manual_approval = bool(self._bot and self.config.auto_send_telegram)
            self.initialize_voice_approvals(characters, require_manual=require_manual_approval)
            self.sync_approvals_with_ready_references(characters)
            self._save_scenario(scenario)
            
            # Реєструємо персонажів у voice manager
            if self._voice_manager:
                for char in characters:
                    self._voice_manager.register_character(char)

            # 3.5. Даємо вікно для завантаження своїх референсів до генерації тестових голосів.
            if require_manual_approval and characters and self.config.pre_voice_reference_stage_enabled:
                await self._send_progress(
                    "📥 Можеш завантажити свої голоси перед генерацією тестових референсів.\n"
                    "Після завантаження натисни '▶️ Продовжити'."
                )
                continue_from_pre_stage = await self._wait_for_pre_voice_reference_stage(characters)
                if not continue_from_pre_stage:
                    await self._send_progress(
                        "⏸ Пауза на етапі завантаження референсів.\n"
                        "Для продовження натисни /resume або кнопку '▶️ Продовжити' у попередньому повідомленні."
                    )
                    return None
            
            # 4. Генеруємо тестові голоси
            await self._send_progress("🎤 Створення тестових голосів...")
            await self.generate_test_voices(characters)

            missing_references = self._collect_characters_missing_any_reference(characters)
            if missing_references:
                preview_missing = ", ".join(missing_references[:8])
                if len(missing_references) > 8:
                    preview_missing += f", ... (+{len(missing_references) - 8})"
                await self._send_progress(
                    "⏸ Не вдалося створити прев'ю/референси для частини героїв.\n"
                    f"⌛ Без аудіо: {preview_missing}\n"
                    "Перевір Qwen3-TTS runtime (qwen_tts + transformers), або завантаж свої голоси, потім /resume."
                )
                self._save_generation_state(
                    status="paused_manual",
                    reason="test_voice_generation_failed",
                    extra={"missing_characters": missing_references}
                )
                if self._current_work_dir:
                    self._write_resume_pointer(
                        self._current_work_dir,
                        status="paused_manual",
                        reason="test_voice_generation_failed"
                    )
                return None
            
            if self._bot and self._current_work_dir:
                test_voices_dir = os.path.join(self._current_work_dir, "test_voices")
                preview_dir = os.path.join(self._current_work_dir, "preview_voices")
                await self._bot.send_message(
                    (
                        f"📁 Папка з preview-голосами:\n{preview_dir}\n{self.format_drive_hint(preview_dir)}\n\n"
                        f"📁 Папка з повними референсами:\n{test_voices_dir}\n{self.format_drive_hint(test_voices_dir)}"
                    ).strip()
                )
            
            if require_manual_approval and characters:
                approved = await self.wait_for_all_voice_approvals()
                if not approved:
                    await self._send_progress("⏸ Генерацію зупинено до затвердження голосів.")
                    self._save_generation_state(status="paused_manual", reason="voice_approval_pending")
                    if self._current_work_dir:
                        self._write_resume_pointer(
                            self._current_work_dir,
                            status="paused_manual",
                            reason="voice_approval_pending"
                        )
                    return None
                refs_ready = await self.wait_for_reference_renders()
                if not refs_ready:
                    await self._send_progress(
                        "⏸ Генерацію зупинено: не всі повні референси голосів готові.\n"
                        "Перегенеруй проблемні голоси або завантаж свій референс."
                    )
                    self._save_generation_state(status="paused_manual", reason="reference_render_pending")
                    if self._current_work_dir:
                        self._write_resume_pointer(
                            self._current_work_dir,
                            status="paused_manual",
                            reason="reference_render_pending"
                        )
                    return None
                await self._send_progress("✅ Усі голоси затверджені і повні референси готові. Починаю генерацію фінального аудіо.")
            
            if self._stop_requested:
                await self._send_progress("⏹ Генерацію зупинено до фінального рендеру.")
                return None
            
            # 5. Генеруємо аудіоп'єсу
            await self._send_progress("🎬 Генерація аудіоп'єси...")
            final_audio = await self.generate_audio_drama(scenario)
            if final_audio:
                self._save_generation_state(status="completed", reason="done")
                self._clear_resume_pointer()
                await self._send_progress("⏸ Генерацію завершено. Очікую новий файл для наступного запуску.")
            
            return final_audio
        except ResourceLimitPauseError as e:
            logger.warning(f"[Orchestrator] Processing paused due to resource limit: {e}")
            if self._current_scenario:
                self._save_scenario(self._current_scenario)
            self._save_generation_state(status="paused_resource", reason=str(e))
            if self._current_work_dir:
                self._write_resume_pointer(self._current_work_dir, status="paused_resource", reason=str(e))
            await self._send_progress(
                "⏸ Генерацію поставлено на паузу через ліміт ресурсів Colab/GPU.\n"
                "Після відновлення лімітів використай /resume або 'продовжити генерацію'."
            )
            return None
            
        except Exception as e:
            logger.error(f"[Orchestrator] Processing failed: {e}")
            self._save_generation_state(status="failed", reason=str(e))
            await self._send_progress(f"❌ Помилка: {str(e)}")
            return None
    
    async def generate_scenario(self, text: str, title: str = "Untitled") -> Scenario:
        """Генерує сценарій з тексту."""
        if self._gemini_client:
            try:
                return await self._gemini_client.convert_to_scenario(
                    text,
                    title,
                    language=self.get_current_language()
                )
            except Exception as e:
                err_text = str(e or "")
                lowered = err_text.lower()
                if "429" in lowered or "quota" in lowered or "rate limit" in lowered:
                    logger.warning(
                        "[Orchestrator] Gemini scenario conversion rate-limited. "
                        "Using simple local scenario fallback."
                    )
                    await self._send_progress(
                        "⚠️ Gemini тимчасово недоступний (quota/rate limit). "
                        "Використовую спрощений локальний сценарій."
                    )
                else:
                    logger.warning(
                        "[Orchestrator] Gemini scenario conversion failed, using fallback: "
                        f"{e}"
                    )
                return self._create_simple_scenario(text, title)
        else:
            # Fallback - простий сценарій
            return self._create_simple_scenario(text, title)
    
    def _create_simple_scenario(self, text: str, title: str) -> Scenario:
        """Створює простий сценарій без Gemini."""
        # Розбиваємо на абзаци
        paragraphs = [p.strip() for p in text.split('\n\n') if p.strip()]
        
        scenes = []
        current_scene_lines = []
        scene_id = 1
        
        for para in paragraphs:
            # Проста евристика: нова сцена при довгому абзаці
            if len(para) > 500 and current_scene_lines:
                scenes.append(Scene(
                    id=scene_id,
                    title=f"Сцена {scene_id}",
                    setting="",
                    dialogue_lines=current_scene_lines
                ))
                scene_id += 1
                current_scene_lines = []
            
            current_scene_lines.append(DialogueLine(
                character_name=NARRATOR_NAME,
                text=para,
                emotion="neutral",
                scene_id=scene_id,
                order_in_scene=len(current_scene_lines) + 1
            ))
        
        # Додаємо останню сцену
        if current_scene_lines:
            scenes.append(Scene(
                id=scene_id,
                title=f"Сцена {scene_id}",
                setting="",
                dialogue_lines=current_scene_lines
            ))
        
        return Scenario(
            title=title,
            author="Unknown",
            scenes=scenes,
            characters=[Character(name=NARRATOR_NAME, role="narrator", gender="male")],
            total_duration_estimate=len(text.split()) / 2.5  # Приблизно
        )

    def _ensure_characters_for_scenario(
        self,
        scenario: Scenario,
        extracted_characters: List[Character]
    ) -> List[Character]:
        """Гарантує, що всі спікери сценарію мають персонажа з валідним описом голосу."""
        characters_map: Dict[str, Character] = {}
        
        def _resolve_existing(name: str) -> str:
            return resolve_name_alias(name, [c.name for c in characters_map.values()])
        
        # 1) Враховуємо персонажів з Gemini, але мерджимо близькі дублікати імен.
        for character in extracted_characters or []:
            if not character.name:
                continue
            normalized_name = canonical_character_name(character.name) or character.name
            normalized_name = _resolve_existing(normalized_name)
            if normalized_name == NARRATOR_NAME:
                character.role = "narrator"
                character.gender_locked = False
            character.name = normalized_name
            if not character.voice_description:
                character.voice_description = self._default_voice_description(character.name, character.role, character.gender)
            desc_gender = extract_gender_from_text(character.voice_description)
            character.gender = resolve_character_gender(character.name, character.role, desc_gender or character.gender)
            key = character.name.strip().lower()
            if not key:
                continue
            if key in characters_map:
                # Якщо дублікат: доповнюємо тільки відсутні поля.
                base = characters_map[key]
                if not base.voice_description and character.voice_description:
                    base.voice_description = character.voice_description
                if not base.personality_traits and character.personality_traits:
                    base.personality_traits = character.personality_traits
                continue
            characters_map[key] = character
        
        # 2) Нормалізуємо спікерів у репліках і теж мерджимо близькі варіанти імен.
        speaker_names: List[str] = []
        for scene in scenario.scenes:
            for dialogue in scene.dialogue_lines:
                if dialogue.character_name and dialogue.character_name.strip():
                    canonical_name = canonical_character_name(dialogue.character_name.strip())
                    if canonical_name:
                        canonical_name = _resolve_existing(canonical_name)
                        dialogue.character_name = canonical_name
                        speaker_names.append(canonical_name)
        
        if not speaker_names:
            speaker_names = [NARRATOR_NAME]
        
        for speaker_name in speaker_names:
            resolved_speaker = _resolve_existing(speaker_name)
            key = resolved_speaker.lower()
            if key in characters_map:
                continue
            
            role = "narrator" if is_narrator_alias(resolved_speaker) else "supporting"
            gender = resolve_character_gender(resolved_speaker, role, None)
            characters_map[key] = Character(
                name=resolved_speaker,
                role=role,
                voice_description=self._default_voice_description(resolved_speaker, role, gender),
                gender=gender,
                gender_locked=False
            )
        
        # 3) Додатково уніфікуємо імена в діалогах згідно фінального списку.
        final_names = [c.name for c in characters_map.values()]
        for scene in scenario.scenes:
            for dialogue in scene.dialogue_lines:
                dialogue.character_name = resolve_name_alias(dialogue.character_name, final_names)
        
        return list(characters_map.values())

    def _default_voice_description(self, character_name: str, role: str = "supporting", gender: str = "male") -> str:
        """Створює дефолтний опис голосу, який відповідає вимозі 20+ слів."""
        return localized_voice_description(
            character_name=character_name,
            role=role,
            gender=gender,
            language=self.get_current_language()
        )

    def get_current_language(self) -> str:
        """Повертає поточну мову озвучки."""
        return normalize_language_code(self._current_language, normalize_language_code(self.config.language, "russian"))

    def _resolve_language_for_text(self, text: str) -> str:
        """Визначає мову для конкретного фрагмента тексту."""
        # Для цього проєкту фіксуємо одну мову озвучки на весь файл (за мовою вхідного тексту).
        return self.get_current_language()

    def _prepare_tts_text(self, text: str, language: Optional[str] = None) -> str:
        """Підготовка тексту для TTS (словник вимови, наголоси, ё тощо)."""
        if not text:
            return text
        lang = normalize_language_code(language or self.get_current_language(), self.get_current_language())
        prepared = apply_inline_stress_markers(text, lang)
        prepared = expand_numbers_for_tts(prepared, lang)
        if self._pronunciation_manager:
            prepared = self._pronunciation_manager.apply(prepared, lang)
        return prepared

    def _has_strong_emotion_cue(self, text: str) -> bool:
        """Визначає явні маркери сильної емоції у тексті."""
        value = (text or "").strip()
        if not value:
            return False
        lowered = value.lower()
        if "!!" in value or "??" in value or "?!" in value or "!?" in value:
            return True
        if re.search(r"\b(крич|ор[её]т|вопит|рычит|шепч|scream|shout|yell|whisper)\w*", lowered):
            return True
        if re.search(r"[A-ZА-ЯІЇЄҐЁ]{4,}", value):
            return True
        return False

    def _stabilize_scenario_emotions(self, scenario: Optional[Scenario]) -> None:
        """Прибирає випадкові різкі емоційні стрибки між репліками одного героя."""
        if not scenario or not scenario.scenes:
            return

        for scene in scenario.scenes:
            if not scene or not scene.dialogue_lines:
                continue
            ordered_lines = sorted(
                scene.dialogue_lines,
                key=lambda dl: (int(getattr(dl, "order_in_scene", 0) or 0), str(getattr(dl, "character_name", "")))
            )
            last_emotion_by_speaker: Dict[str, str] = {}
            for line in ordered_lines:
                speaker = canonical_character_name(line.character_name) or line.character_name
                current = normalize_emotion_label(line.emotion, line.text)
                prev = last_emotion_by_speaker.get(speaker)
                strong_cue = self._has_strong_emotion_cue(line.text)
                is_narrator = is_narrator_alias(speaker) or speaker == NARRATOR_NAME

                if is_narrator and current in {"shout", "angry", "excited"}:
                    current = "neutral"

                if current in {"shout", "angry"} and not strong_cue:
                    if prev in {"calm", "neutral", "happy", "sad", "dramatic"}:
                        current = prev
                    elif "!" in (line.text or ""):
                        current = "dramatic"
                    else:
                        current = "neutral"
                elif current == "excited" and not strong_cue:
                    if prev in {"calm", "neutral", "sad"}:
                        current = prev
                    elif "!" not in (line.text or ""):
                        current = "happy"

                line.emotion = normalize_emotion_label(current, line.text)
                last_emotion_by_speaker[speaker] = line.emotion

    def add_pronunciation_rule(self, language: str, source: str, target: str) -> bool:
        """Додає/оновлює правило вимови."""
        if not self._pronunciation_manager:
            return False
        return self._pronunciation_manager.upsert_rule(language, source, target)

    def remove_pronunciation_rule(self, language: str, source: str) -> bool:
        """Видаляє правило вимови."""
        if not self._pronunciation_manager:
            return False
        return self._pronunciation_manager.remove_rule(language, source)

    def list_pronunciation_rules(self, language: str) -> Dict[str, str]:
        """Повертає правила вимови для мови."""
        if not self._pronunciation_manager:
            return {}
        return self._pronunciation_manager.list_rules(language)

    def get_pronunciation_dictionary_path(self) -> str:
        """Повертає шлях до файла словника вимови."""
        if not self._pronunciation_manager:
            return ""
        return self._pronunciation_manager.storage_path

    def mark_character_audio_dirty(self, character_name: str) -> None:
        """Позначає персонажа як зміненого, щоб його репліки перегенерувались інкрементально."""
        if not character_name:
            return
        candidate = canonical_character_name(character_name) or character_name
        resolved = self._resolve_voice_approval_name(candidate) or candidate
        if self._current_scenario:
            char = self._current_scenario.get_character(resolved)
            if char:
                resolved = char.name
        if not resolved:
            return
        self._dirty_characters.add(resolved)

    def clear_dirty_characters(self) -> None:
        """Очищає список змінених персонажів після успішної фіналізації."""
        self._dirty_characters.clear()

    def get_dirty_characters(self) -> List[str]:
        """Повертає список персонажів, чиї репліки треба регенерувати."""
        return sorted(self._dirty_characters)

    def get_last_generation_stats(self) -> Dict[str, int]:
        """Повертає статистику останньої генерації аудіо."""
        return dict(self._last_generation_stats)

    def _is_dialogue_character_dirty(self, dialogue: DialogueLine, dirty_characters: set[str]) -> bool:
        if not dialogue or not dirty_characters:
            return False
        speaker = canonical_character_name(dialogue.character_name) or dialogue.character_name
        for dirty_name in dirty_characters:
            if are_names_equivalent(speaker, dirty_name):
                return True
        return False

    def _resolve_voice_approval_name(self, requested_name: str) -> Optional[str]:
        """Повертає канонічне ім'я персонажа зі списку approval."""
        if not requested_name:
            return None
        candidate = canonical_character_name(requested_name) or requested_name
        normalized = candidate.strip().lower()
        if is_narrator_alias(normalized):
            normalized = NARRATOR_NAME.lower()
        for name in self._voice_approvals.keys():
            if name.strip().lower() == normalized:
                return name
        for name in self._voice_approvals.keys():
            if are_names_equivalent(candidate, name):
                return name
        return None

    def initialize_voice_approvals(self, characters: List[Character], require_manual: bool) -> None:
        """Ініціалізує статуси затвердження голосів."""
        self._voice_approvals = {}
        for character in characters or []:
            if character and character.name:
                self._voice_approvals[character.name] = (not require_manual)
        self._current_approval_character = None
        self._voice_update_event.clear()
        
        if not self._voice_approvals or all(self._voice_approvals.values()):
            self._voice_approval_event.set()
        else:
            self._voice_approval_event.clear()

    def sync_approvals_with_ready_references(self, characters: List[Character]) -> None:
        """Автоматично затверджує персонажів із готовими референсами."""
        for character in characters or []:
            if not character or not character.name:
                continue
            if self._is_character_reference_ready(character):
                self._voice_approvals[character.name] = True
        if not self._voice_approvals or all(self._voice_approvals.values()):
            self._voice_approval_event.set()
        else:
            self._voice_approval_event.clear()

    def update_voice_approval(self, character_name: str, approved: bool) -> Tuple[Optional[str], int, int, bool]:
        """Оновлює затвердження голосу персонажа."""
        resolved_name = self._resolve_voice_approval_name(character_name)
        if not resolved_name:
            total = len(self._voice_approvals)
            approved_count = sum(1 for value in self._voice_approvals.values() if value)
            return None, approved_count, total, (total > 0 and approved_count == total)
        
        self._voice_approvals[resolved_name] = bool(approved)
        total = len(self._voice_approvals)
        approved_count = sum(1 for value in self._voice_approvals.values() if value)
        all_approved = (total > 0 and approved_count == total)
        if not approved and not self._current_approval_character:
            self._current_approval_character = resolved_name
        if approved and self._current_approval_character and are_names_equivalent(self._current_approval_character, resolved_name):
            self._current_approval_character = None
        
        if all_approved:
            self._voice_approval_event.set()
        else:
            self._voice_approval_event.clear()
        self._voice_update_event.set()
        
        return resolved_name, approved_count, total, all_approved

    def get_voice_approval_summary(self) -> Tuple[int, int, List[str]]:
        """Повертає (approved_count, total, pending_names)."""
        total = len(self._voice_approvals)
        approved_count = sum(1 for value in self._voice_approvals.values() if value)
        pending = [name for name, value in self._voice_approvals.items() if not value]
        return approved_count, total, pending

    def is_waiting_for_pre_voice_reference(self) -> bool:
        """Повертає True, коли активний pre-stage завантаження референсів."""
        return bool(self._waiting_for_pre_voice_reference)

    def _is_character_reference_ready(self, character: Optional[Character]) -> bool:
        """Повертає True, якщо для персонажа вже є готовий референс, який можна не перегенеровувати."""
        if not character:
            return False
        if self._uses_user_uploaded_reference(character) and self._has_ready_full_reference(character):
            return True
        if self._has_ready_full_reference(character) and self._voice_approvals.get(character.name, False):
            return True
        return False

    def get_reference_upload_summary(self) -> Tuple[int, int, List[str]]:
        """Повертає (ready_count, total, pending_names) для pre-stage референсів."""
        if not self._current_scenario or not self._current_scenario.characters:
            return 0, 0, []
        ready_names: List[str] = []
        pending_names: List[str] = []
        for character in self._current_scenario.characters:
            if self._is_character_reference_ready(character):
                ready_names.append(character.name)
            else:
                pending_names.append(character.name)
        return len(ready_names), len(self._current_scenario.characters), pending_names

    def _collect_characters_missing_any_reference(self, characters: List[Character]) -> List[str]:
        """Повертає персонажів, для яких немає жодного наявного audio-референсу."""
        missing: List[str] = []
        for character in characters or []:
            if not character:
                continue
            ref_path = (character.reference_audio_path or "").strip()
            if ref_path and os.path.exists(ref_path):
                continue
            missing.append(character.name)
        return missing

    def _build_pre_voice_reference_buttons(self, characters: List[Character]) -> List[List[Dict[str, str]]]:
        """Формує inline-кнопки для pre-stage завантаження референсів."""
        if not self._bot:
            return []
        buttons: List[List[Dict[str, str]]] = []
        row: List[Dict[str, str]] = []
        for character in characters or []:
            token = self._bot._register_character_token(character.name)
            row.append({
                "text": f"📥 {character.name[:22]}",
                "callback_data": f"upl:{token}"
            })
            if len(row) >= 2:
                buttons.append(row)
                row = []
        if row:
            buttons.append(row)
        buttons.append([{"text": "↩️ Скасувати вибір героя", "callback_data": "upl_cancel"}])
        buttons.append([
            {"text": "▶️ Продовжити", "callback_data": "prevoice_continue"},
            {"text": "⏭ Пропустити свої голоси", "callback_data": "prevoice_skip"},
        ])
        buttons.append([{"text": "⏸ Пауза", "callback_data": "prevoice_pause"}])
        return buttons

    async def _wait_for_pre_voice_reference_stage(self, characters: List[Character]) -> bool:
        """Ставить паузу перед тестовими голосами, щоб користувач міг завантажити свої референси."""
        if not characters:
            return True
        if not self._bot or not self.config.auto_send_telegram:
            return True
        if not self.config.pre_voice_reference_stage_enabled:
            return True
        if self._stop_requested:
            return False

        self._waiting_for_pre_voice_reference = True
        self._pre_voice_reference_event.clear()
        self._save_generation_state(status="paused_manual", reason="pre_voice_reference_upload_pending")
        if self._current_work_dir:
            self._write_resume_pointer(
                self._current_work_dir,
                status="paused_manual",
                reason="pre_voice_reference_upload_pending"
            )

        ready_count, total_count, pending_names = self.get_reference_upload_summary()
        pending_preview = ", ".join(pending_names[:8]) if pending_names else "-"
        if len(pending_names) > 8:
            pending_preview += f", ... (+{len(pending_names) - 8})"
        wait_message = (
            "📥 Перед генерацією тестових голосів можна завантажити свої референси.\n"
            "Натисни кнопку героя, надішли audio/voice (і, за потреби, .md опис), потім натисни '▶️ Продовжити'.\n"
            "Якщо свої голоси не потрібні, натисни '⏭ Пропустити свої голоси'.\n"
            f"📦 Готово референсів: {ready_count}/{total_count}\n"
            f"⌛ Очікують: {pending_preview}"
        )
        buttons = self._build_pre_voice_reference_buttons(characters)
        if buttons:
            await self._bot.send_inline_buttons(wait_message, buttons)
        else:
            await self._send_progress(wait_message)

        timeout_sec = max(0, int(self.config.pre_voice_reference_timeout_sec or 0))
        try:
            if timeout_sec > 0:
                await asyncio.wait_for(self._pre_voice_reference_event.wait(), timeout=timeout_sec)
            else:
                await self._pre_voice_reference_event.wait()
        except asyncio.TimeoutError:
            await self._send_progress(
                f"⌛ Час очікування pre-stage ({timeout_sec}с) вичерпано. "
                "Продовжую генерацію тестових голосів автоматично."
            )
            self._stop_requested = False
        finally:
            self._waiting_for_pre_voice_reference = False
            self._pre_voice_reference_event.clear()

        if self._stop_requested:
            self._save_generation_state(status="paused_manual", reason="pre_voice_reference_upload_pending")
            if self._current_work_dir:
                self._write_resume_pointer(
                    self._current_work_dir,
                    status="paused_manual",
                    reason="pre_voice_reference_upload_pending"
                )
            return False

        ready_count, total_count, pending_names = self.get_reference_upload_summary()
        pending_hint = ", ".join(pending_names[:6]) if pending_names else "-"
        if len(pending_names) > 6:
            pending_hint += f", ... (+{len(pending_names) - 6})"
        await self._send_progress(
            "▶️ Pre-stage завершено, запускаю генерацію тестових голосів.\n"
            f"📦 Готово референсів: {ready_count}/{total_count}\n"
            f"⌛ Без власного референсу: {pending_hint}"
        )
        self._save_generation_state(status="running", reason="pre_voice_reference_upload_done")
        if self._current_work_dir:
            self._write_resume_pointer(self._current_work_dir, status="running", reason="pre_voice_reference_upload_done")
        return True

    async def handle_pre_voice_reference_decision(self, continue_generation: bool) -> None:
        """Обробляє рішення користувача для pre-stage перед тестовими голосами."""
        if continue_generation:
            logger.info("[Orchestrator] User chose to continue from pre-voice-reference stage")
            self._stop_requested = False
            self._pre_voice_reference_event.set()
            self._save_generation_state(status="running", reason="pre_voice_reference_continue")
        else:
            logger.info("[Orchestrator] User chose to pause on pre-voice-reference stage")
            self._stop_requested = True
            self._pre_voice_reference_event.set()
            self._save_generation_state(status="paused_manual", reason="pre_voice_reference_upload_pending")
            if self._current_work_dir:
                self._write_resume_pointer(
                    self._current_work_dir,
                    status="paused_manual",
                    reason="pre_voice_reference_upload_pending"
                )

    def get_voice_approval_progress_text(self) -> str:
        """Повертає рядок прогресу затвердження голосів."""
        approved_count, total, pending = self.get_voice_approval_summary()
        if total <= 0:
            return "📊 Затверджено голосів: 0/0"
        percent = int(round((approved_count / max(1, total)) * 100))
        width = 10
        filled = int(round((percent / 100.0) * width))
        bar = f"[{'█' * filled}{'░' * (width - filled)}]"
        pending_preview = ", ".join(pending[:4]) if pending else "-"
        return f"📊 Затверджено голосів: {approved_count}/{total} {bar} {percent}%\n⌛ Очікують: {pending_preview}"

    def approve_next_pending_voice(self) -> Tuple[Optional[str], int, int, bool]:
        """Затверджує наступний голос у черзі (для короткої відповіді 'ок')."""
        _, _, pending = self.get_voice_approval_summary()
        if not pending:
            approved_count, total, _ = self.get_voice_approval_summary()
            return None, approved_count, total, (total > 0 and approved_count == total)
        target = pending[0]
        if self._current_approval_character:
            for name in pending:
                if are_names_equivalent(name, self._current_approval_character):
                    target = name
                    break
        return self.update_voice_approval(target, approved=True)

    async def wait_for_character_approval(self, character_name: str) -> bool:
        """Чекає затвердження конкретного персонажа перед генерацією наступного."""
        resolved = self._resolve_voice_approval_name(character_name) or character_name
        self._current_approval_character = resolved
        while True:
            _, _, pending = self.get_voice_approval_summary()
            if not any(are_names_equivalent(resolved, item) for item in pending):
                return True
            if self._stop_requested:
                return False
            try:
                await asyncio.wait_for(self._voice_update_event.wait(), timeout=3.0)
            except asyncio.TimeoutError:
                pass
            self._voice_update_event.clear()

    async def wait_for_all_voice_approvals(self) -> bool:
        """Чекає поки всі голоси будуть затверджені."""
        approved_count, total, pending = self.get_voice_approval_summary()
        if total == 0 or approved_count == total:
            return True
        
        pending_preview = ", ".join(pending[:8])
        if len(pending) > 8:
            pending_preview += f", ... (+{len(pending) - 8})"
        await self._send_progress(
            "🛑 Перед генерацією фінального аудіо потрібно затвердити всі голоси.\n"
            f"✅ Затверджено: {approved_count}/{total}\n"
            f"⌛ Очікують: {pending_preview}"
        )
        
        while True:
            if self._stop_requested:
                return False
            approved_count, total, _ = self.get_voice_approval_summary()
            if total == 0 or approved_count == total:
                return True
            try:
                await asyncio.wait_for(self._voice_approval_event.wait(), timeout=3.0)
            except asyncio.TimeoutError:
                continue
    
    def _save_voice_profile_markdown(
        self,
        character: Character,
        audio_path: str,
        stage: str = "preview"
    ) -> Optional[str]:
        """Зберігає markdown-паспорт голосу поруч з audio файлом."""
        if not audio_path:
            return None
        try:
            md_path = os.path.splitext(audio_path)[0] + ".md"
            content = [
                f"# Voice Profile: {character.name}",
                "",
                f"- stage: {stage}",
                f"- role: {character.role}",
                f"- gender: {character.gender}",
                f"- language: {self.get_current_language()}",
                f"- audio_path: {audio_path}",
                f"- updated_at: {datetime.now().isoformat()}",
                "",
                "## Voice Description",
                character.voice_description or "",
                "",
                "## Reference Text",
                character.reference_text or "",
            ]
            with open(md_path, "w", encoding="utf-8") as f:
                f.write("\n".join(content).strip() + "\n")
            return md_path
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to save voice markdown for {character.name}: {e}")
            return None
    
    def _uses_user_uploaded_reference(self, character: Optional[Character]) -> bool:
        """Повертає True, якщо референс голосу був завантажений користувачем."""
        if not character or not character.reference_audio_path:
            return False
        ref_path = str(character.reference_audio_path).replace("\\", "/").lower()
        return "/references/" in ref_path or "_reference_v" in os.path.basename(ref_path)
    
    def _has_ready_full_reference(self, character: Optional[Character]) -> bool:
        """Перевіряє, чи вже існує готовий (не preview) референс голосу."""
        if not character or not character.reference_audio_path:
            return False
        ref_path = str(character.reference_audio_path).replace("\\", "/").lower()
        if "/preview_voices/" in ref_path:
            return False
        return os.path.exists(character.reference_audio_path)
    
    def _schedule_full_reference_render(self, character_name: str, full_output_dir: str) -> None:
        """Ставить у фон повний рендер референсу для вже затвердженого героя."""
        if not self._voice_manager:
            return
        character = self._voice_manager.get_character(character_name)
        if not character:
            return
        if self._uses_user_uploaded_reference(character):
            logger.info(f"[Orchestrator] User reference already set for {character.name}, skipping full render task")
            return
        if self._has_ready_full_reference(character):
            return
        existing = self._reference_render_tasks.get(character.name)
        if existing and not existing.done():
            return
        task = asyncio.create_task(
            self._render_full_reference_voice(character.name, full_output_dir)
        )
        self._reference_render_tasks[character.name] = task
    
    async def _render_full_reference_voice(self, character_name: str, full_output_dir: str) -> Optional[str]:
        """Генерує повний (довший) референс голосу у фоні."""
        if not self._voice_manager:
            return None
        try:
            while self._preview_generation_active and not self._stop_requested:
                await asyncio.sleep(0.2)
            async with self._reference_render_lock:
                if self._stop_requested:
                    return None
                ready = await self._ensure_design_engine_loaded()
                if not ready:
                    raise RuntimeError("VoiceDesign engine unavailable for full reference render")
                character = self._voice_manager.get_character(character_name)
                if not character:
                    return None
                if self._uses_user_uploaded_reference(character):
                    return character.reference_audio_path
                if self._has_ready_full_reference(character):
                    return character.reference_audio_path
                
                voice_path = await self._voice_manager.create_test_voice(
                    character,
                    self.config.test_voice_duration,
                    full_output_dir,
                    language=self.get_current_language()
                )
                persisted_voice_path = self._ensure_file_on_drive(voice_path, "test_voices")
                self._sync_character_reference_path(
                    character.name,
                    old_path=voice_path,
                    new_path=persisted_voice_path
                )
                character.reference_audio_path = persisted_voice_path
                md_path = self._save_voice_profile_markdown(character, persisted_voice_path, stage="approved_reference")
                if self._bot and self.config.auto_send_telegram:
                    md_hint = f"\n📝 {md_path}" if md_path else ""
                    await self._bot.send_message(
                        (
                            f"✅ Повний референс для '{character.name}' готовий.\n"
                            f"📁 {persisted_voice_path}{md_hint}\n"
                            f"{self.format_drive_hint(persisted_voice_path)}"
                        ).strip()
                    )
                return persisted_voice_path
        except Exception as e:
            self._reference_render_errors[character_name] = str(e)
            logger.error(f"[Orchestrator] Full reference render failed for {character_name}: {e}")
            return None
    
    async def wait_for_reference_renders(self) -> bool:
        """Чекає завершення фонового повного рендеру референсів."""
        if not self._reference_render_tasks:
            return True
        total = len(self._reference_render_tasks)
        await self._send_progress(f"⏳ Дорендер повних референсів голосів: 0/{total}")
        
        last_done = -1
        while True:
            if self._stop_requested:
                return False
            pending = [t for t in self._reference_render_tasks.values() if not t.done()]
            done = total - len(pending)
            if done != last_done:
                last_done = done
                await self._send_progress(f"⏳ Дорендер повних референсів голосів: {done}/{total}")
            if not pending:
                break
            await asyncio.sleep(1.0)
        
        # Після завершення збираємо помилки.
        errors: List[str] = []
        for name, task in self._reference_render_tasks.items():
            try:
                result = task.result()
                if not result and name not in self._reference_render_errors:
                    self._reference_render_errors[name] = "unknown render error"
            except Exception as e:
                self._reference_render_errors[name] = str(e)
            if name in self._reference_render_errors:
                errors.append(f"{name}: {self._reference_render_errors[name]}")
        
        if errors:
            preview = "\n".join(errors[:6])
            await self._send_progress(
                "⚠️ Повний рендер деяких референсів завершився з помилками.\n"
                f"{preview}\n"
                "Можна перегенерувати голос або завантажити свій референс."
            )
            return False
        
        await self._send_progress("✅ Усі повні референси голосів готові.")
        return True
    
    async def generate_test_voices(self, characters: List[Character]) -> Dict[str, str]:
        """Генерує тестові голоси для персонажів."""
        if not self._voice_manager:
            logger.warning("[Orchestrator] Voice manager not available")
            return {}
        
        design_ready = await self._ensure_design_engine_loaded()
        if not design_ready:
            logger.error("[Orchestrator] VoiceDesign engine is unavailable. Test voices cannot be generated.")
            return {}
        
        test_voices: Dict[str, str] = {}
        manual_approval_flow = any(not approved for approved in self._voice_approvals.values())
        preview_dir = os.path.join(self._current_work_dir, "preview_voices")
        full_dir = os.path.join(self._current_work_dir, "test_voices")
        os.makedirs(preview_dir, exist_ok=True)
        os.makedirs(full_dir, exist_ok=True)
        self._preview_generation_active = manual_approval_flow
        try:
            for char in characters:
                if self._stop_requested:
                    logger.info("[Orchestrator] Stop requested during test voice generation")
                    break
                if self._is_character_reference_ready(char):
                    ready_path = char.reference_audio_path or ""
                    if ready_path:
                        test_voices[char.name] = ready_path
                    logger.info(
                        f"[Orchestrator] Reusing ready reference for {char.name}: "
                        f"{ready_path or 'n/a'}"
                    )
                    if self._bot and self.config.auto_send_telegram and ready_path:
                        await self._bot.send_message(
                            (
                                f"📌 Використовую готовий референс для '{char.name}'.\n"
                                f"📁 {ready_path}\n"
                                f"{self.format_drive_hint(ready_path)}"
                            ).strip()
                        )
                    continue
                try:
                    generation_duration = self.config.preview_voice_duration if manual_approval_flow else self.config.test_voice_duration
                    output_dir = preview_dir if manual_approval_flow else full_dir
                    voice_path = await self._voice_manager.create_test_voice(
                        char,
                        generation_duration,
                        output_dir,
                        language=self.get_current_language()
                    )
                    persisted_voice_path = self._ensure_file_on_drive(
                        voice_path,
                        "preview_voices" if manual_approval_flow else "test_voices"
                    )
                    self._sync_character_reference_path(char.name, old_path=voice_path, new_path=persisted_voice_path)
                    voice_path = persisted_voice_path
                    test_voices[char.name] = voice_path
                    md_path = self._save_voice_profile_markdown(
                        char,
                        voice_path,
                        stage="preview" if manual_approval_flow else "approved_reference"
                    )
                    
                    # Відправляємо тестовий голос в Telegram
                    if self._bot and self.config.auto_send_telegram:
                        progress_text = self.get_voice_approval_progress_text()
                        voice_card_message = await self._bot.send_audio(
                            voice_path,
                            (
                                f"🎤 {'3с прев' if manual_approval_flow else 'Тестовий'} голос: {char.name}\n"
                                f"{char.voice_description[:100]}...\n\n"
                                f"{progress_text}"
                            ),
                            reply_markup=self._bot._get_voice_keyboard(char.name)
                        )
                        if voice_card_message:
                            self._bot.track_voice_card_message(char.name, voice_card_message)
                        md_hint = f"\n📝 Опис (.md): {md_path}" if md_path else ""
                        await self._bot.send_message(
                            (
                                f"🔗 Референс {char.name}:\n{voice_path}\n{self.format_drive_hint(voice_path)}\n"
                                f"{md_hint}\n"
                                f"✅ Натисни 'Прийняти' для {char.name} або 'Перегенерувати'\n"
                                f"📥 Або натисни 'Свій голос' і надішли audio/voice + .md опис\n"
                                f"{progress_text}"
                            ).strip()
                        )
                        if manual_approval_flow and char.name in self._voice_approvals:
                            approved = await self.wait_for_character_approval(char.name)
                            if not approved:
                                logger.info(f"[Orchestrator] Stopped while waiting approval for {char.name}")
                                return test_voices
                            # Після approve запускаємо повний рендер у фоні, не блокуючи наступного героя.
                            self._schedule_full_reference_render(char.name, full_dir)
                    
                except Exception as e:
                    if is_resource_limit_error(e):
                        raise ResourceLimitPauseError(str(e)) from e
                    logger.error(f"[Orchestrator] Failed to create voice for {char.name}: {e}")
            
            return test_voices
        finally:
            self._preview_generation_active = False

    def _build_progress_bar(self, percent: int, width: int = 20) -> str:
        """Формує текстовий progress bar."""
        safe_percent = max(0, min(100, int(percent)))
        filled = int(round((safe_percent / 100.0) * width))
        return f"[{'█' * filled}{'░' * (width - filled)}]"

    async def _send_generation_status(
        self,
        scenario: Scenario,
        processed_dialogues: int,
        total_dialogues: int,
        stage: str = "",
        force: bool = False
    ) -> None:
        """Оновлює статус-бар генерації основного аудіо."""
        if not self._bot:
            return
        
        total = max(1, total_dialogues)
        display_total = total_dialogues if total_dialogues > 0 else 1
        percent = int(round((processed_dialogues / total) * 100))
        minute_bucket = int(self._generated_duration // 60)
        should_update = (
            force
            or percent >= self._last_progress_percent + 2
            or minute_bucket > self._last_progress_min_bucket
            or processed_dialogues >= total_dialogues
        )
        if not should_update:
            return
        
        self._last_progress_percent = percent
        self._last_progress_min_bucket = minute_bucket
        est_minutes = max(0.1, float(scenario.total_duration_estimate or 0.0) / 60.0)
        generated_minutes = max(0.0, self._generated_duration / 60.0)
        bar = self._build_progress_bar(percent)
        status_message = (
            f"🎬 Генерація аудіо\n"
            f"{bar} {percent}%\n"
            f"🧩 Реплік: {processed_dialogues}/{display_total}\n"
            f"⏱ Згенеровано: {generated_minutes:.1f} хв із ~{est_minutes:.1f} хв\n"
            f"🔹 Етап: {stage or 'Озвучування'}"
        )
        await self._bot.send_or_update_progress(status_message)
    
    async def generate_audio_drama(self, scenario: Scenario) -> Optional[str]:
        """Генерує фінальну аудіоп'єсу."""
        if not self._voice_manager or not self._audio_assembler:
            logger.error("[Orchestrator] Audio components not initialized")
            return None
        self._last_final_audio_sent_ok = False
        self._stabilize_scenario_emotions(scenario)
        
        # Перед генерацією реплік намагаємось підняти clone engine (lazy-load)
        clone_ready = await self._ensure_clone_engine_loaded()
        if not clone_ready:
            logger.warning("[Orchestrator] VoiceClone unavailable, using VoiceDesign fallback for dialogues")
        
        dialogue_dir = os.path.join(self._current_work_dir, "dialogue_chunks")
        background_dir = os.path.join(self._current_work_dir, "background_sounds")
        os.makedirs(dialogue_dir, exist_ok=True)
        os.makedirs(background_dir, exist_ok=True)
        
        mixing_config = MixingConfig()
        self._generated_duration = 0.0
        total_dialogues = sum(len(scene.dialogue_lines) for scene in scenario.scenes)
        processed_dialogues = 0
        reused_dialogues = 0
        regenerated_dialogues = 0
        dirty_characters = set(self._dirty_characters)
        current_scene_id: Optional[int] = None
        current_order_in_scene: Optional[int] = None

        def _flush_generation_stats() -> None:
            self._last_generation_stats = {
                "total_dialogues": int(total_dialogues),
                "reused_dialogues": int(reused_dialogues),
                "regenerated_dialogues": int(regenerated_dialogues),
            }
        
        try:
            self._save_generation_state(
                status="running",
                reason="audio_generation_started",
                processed_dialogues=processed_dialogues,
                total_dialogues=total_dialogues
            )
            await self._send_generation_status(
                scenario,
                processed_dialogues,
                total_dialogues,
                stage="Підготовка",
                force=True
            )
            if dirty_characters:
                dirty_preview = ", ".join(sorted(dirty_characters)[:6])
                if len(dirty_characters) > 6:
                    dirty_preview += f", ... (+{len(dirty_characters) - 6})"
                await self._send_progress(
                    "♻️ Інкрементальний режим: перегенерація лише змінених голосів.\n"
                    f"🎭 Змінені персонажі: {dirty_preview}"
                )
            
            # Генеруємо аудіо для кожної сцени
            for scene in scenario.scenes:
                current_scene_id = scene.id
                if self._stop_requested:
                    await self._send_progress("⏹ Генерацію зупинено користувачем.")
                    await self._send_generation_status(
                        scenario,
                        processed_dialogues,
                        total_dialogues,
                        stage="Зупинено",
                        force=True
                    )
                    _flush_generation_stats()
                    return None
                
                await self._send_progress(f"🎬 Сцена {scene.id}: {scene.title}")
                
                # Генеруємо репліки
                for dialogue in scene.dialogue_lines:
                    current_order_in_scene = dialogue.order_in_scene
                    if self._stop_requested:
                        await self._send_progress("⏹ Генерацію зупинено користувачем.")
                        await self._send_generation_status(
                            scenario,
                            processed_dialogues,
                            total_dialogues,
                            stage="Зупинено",
                            force=True
                        )
                        self._save_generation_state(
                            status="paused_manual",
                            reason="stopped_by_user",
                            current_scene_id=current_scene_id,
                            current_order_in_scene=current_order_in_scene,
                            processed_dialogues=processed_dialogues,
                            total_dialogues=total_dialogues
                        )
                        _flush_generation_stats()
                        return None
                    
                    dialogue_is_dirty = self._is_dialogue_character_dirty(dialogue, dirty_characters)
                    if dialogue_is_dirty:
                        # Для зміненого голосу примусово скидаємо старий chunk.
                        if dialogue.audio_path and os.path.exists(dialogue.audio_path):
                            try:
                                os.remove(dialogue.audio_path)
                            except Exception:
                                pass
                        dialogue.audio_path = None
                        dialogue.duration = 0.0

                    # Resume path: якщо файл репліки вже існує, пропускаємо регенерацію.
                    if not dialogue.audio_path:
                        safe_name = dialogue.character_name.replace(' ', '_').lower()
                        candidate_path = os.path.join(
                            dialogue_dir,
                            f"scene{dialogue.scene_id}_{safe_name}_{dialogue.order_in_scene}.wav"
                        )
                        if dialogue_is_dirty and os.path.exists(candidate_path):
                            try:
                                os.remove(candidate_path)
                            except Exception:
                                pass
                        if os.path.exists(candidate_path):
                            dialogue.audio_path = candidate_path
                    
                    if dialogue.audio_path and os.path.exists(dialogue.audio_path):
                        if not dialogue.duration or dialogue.duration <= 0:
                            dialogue.duration = self._audio_assembler.get_duration(dialogue.audio_path)
                        self._generated_duration += max(0.0, dialogue.duration or 0.0)
                        reused_dialogues += 1
                        processed_dialogues += 1
                        self._save_generation_state(
                            status="running",
                            reason="resume_skip_existing_dialogue",
                            current_scene_id=current_scene_id,
                            current_order_in_scene=current_order_in_scene,
                            processed_dialogues=processed_dialogues,
                            total_dialogues=total_dialogues
                        )
                        await self._send_generation_status(
                            scenario,
                            processed_dialogues,
                            total_dialogues,
                            stage=f"Resume: сцена {scene.id}"
                        )
                        continue
                    
                    audio_path = await self._generate_dialogue_audio(
                        dialogue,
                        scenario,
                        dialogue_dir
                    )
                    processed_dialogues += 1
                    
                    if audio_path:
                        dialogue.audio_path = audio_path
                        dialogue.duration = self._audio_assembler.get_duration(audio_path)
                        self._generated_duration += dialogue.duration
                        regenerated_dialogues += 1
                        
                        # Перевіряємо контрольні точки
                        await self._check_checkpoints()
                        self._save_scenario(scenario)
                    
                    self._save_generation_state(
                        status="running",
                        reason="dialogue_processed",
                        current_scene_id=current_scene_id,
                        current_order_in_scene=current_order_in_scene,
                        processed_dialogues=processed_dialogues,
                        total_dialogues=total_dialogues
                    )
                    
                    await self._send_generation_status(
                        scenario,
                        processed_dialogues,
                        total_dialogues,
                        stage=f"Сцена {scene.id}: {scene.title}"
                    )
                
                # Генеруємо фонові звуки (Pexels або fallback synthetic ambience)
                await self._generate_background_sounds(scene, background_dir)
            
            # Збираємо фінальне аудіо
            final_path = os.path.join(self._current_work_dir, "final_audio_drama.wav")
            
            await self._send_progress("🔧 Збірка фінального аудіо...")
            await self._send_generation_status(
                scenario,
                total_dialogues,
                total_dialogues,
                stage="Збірка фінального аудіо",
                force=True
            )
            
            self._audio_assembler.assemble_from_scenes(
                scenario.scenes,
                dialogue_dir,
                mixing_config,
                final_path
            )
            await self._send_generation_status(
                scenario,
                total_dialogues,
                total_dialogues,
                stage="Готово",
                force=True
            )
            self._save_scenario(scenario)
            self._save_generation_state(
                status="completed",
                reason="audio_generation_completed",
                current_scene_id=current_scene_id,
                current_order_in_scene=current_order_in_scene,
                processed_dialogues=total_dialogues,
                total_dialogues=total_dialogues
            )
            self._clear_resume_pointer()
            _flush_generation_stats()
            self.clear_dirty_characters()
            
            # Відправляємо результат
            if self._bot:
                sent_final = await self._bot.send_audio(
                    final_path,
                    (
                        f"🎭 Аудіоп'єса: {scenario.title}\n"
                        f"⏱ Тривалість: {self._generated_duration:.1f} сек\n"
                        f"{self.format_drive_hint(final_path)}"
                    )
                )
                self._last_final_audio_sent_ok = bool(sent_final)
                if not sent_final:
                    await self._bot.send_message(
                        (
                            "⚠️ Не вдалося автоматично відправити фінальний файл як audio/document.\n"
                            "Перевір ліміт Telegram для розміру файла або мережеві помилки.\n"
                            f"{self.format_drive_hint(final_path)}"
                        ).strip()
                    )
            
            return final_path
        except MissingReferenceVoiceError as e:
            logger.warning(f"[Orchestrator] Pausing generation due strict reference mode: {e}")
            self._save_scenario(scenario)
            self._save_generation_state(
                status="paused_manual",
                reason="reference_voice_prompt_missing",
                current_scene_id=current_scene_id,
                current_order_in_scene=current_order_in_scene,
                processed_dialogues=processed_dialogues,
                total_dialogues=total_dialogues,
                extra={
                    "character_name": e.character_name,
                    "scene_id": e.scene_id,
                    "order_in_scene": e.order_in_scene,
                    "reference_audio_path": e.reference_audio_path,
                    "error_reason": e.reason,
                    "error_message": str(e),
                }
            )
            if self._current_work_dir:
                self._write_resume_pointer(
                    self._current_work_dir,
                    status="paused_manual",
                    reason="reference_voice_prompt_missing"
                )
            await self._send_generation_status(
                scenario,
                processed_dialogues,
                total_dialogues,
                stage=f"Пауза: референс '{e.character_name}'",
                force=True
            )
            ref_hint = f"\n📁 Reference: {e.reference_audio_path}" if e.reference_audio_path else ""
            await self._send_progress(
                "⏸ Генерацію поставлено на паузу (strict reference mode).\n"
                f"Не вдалося зібрати voice prompt для героя '{e.character_name}'.{ref_hint}\n"
                "Завантаж/онови референс цього героя або перегенеруй його голос, потім надішли /resume."
            )
            _flush_generation_stats()
            return None
        except ResourceLimitPauseError as e:
            logger.warning(f"[Orchestrator] Pausing generation due to resource limit: {e}")
            self._save_scenario(scenario)
            partial_path = self._save_partial_audio_checkpoint(scenario, "auto_pause_resource")
            self._save_generation_state(
                status="paused_resource",
                reason=str(e),
                current_scene_id=current_scene_id,
                current_order_in_scene=current_order_in_scene,
                processed_dialogues=processed_dialogues,
                total_dialogues=total_dialogues,
                extra={"partial_checkpoint": partial_path}
            )
            if self._current_work_dir:
                self._write_resume_pointer(self._current_work_dir, status="paused_resource", reason=str(e))
            await self._send_generation_status(
                scenario,
                processed_dialogues,
                total_dialogues,
                stage="Пауза: ліміт GPU/ресурсів",
                force=True
            )
            partial_hint = ""
            if partial_path:
                partial_hint = (
                    f"\n🎧 Partial-аудіо: {partial_path}\n"
                    f"{self.format_drive_hint(partial_path)}"
                )
            await self._send_progress(
                "⏸ Генерацію автоматично поставлено на паузу через обмеження ресурсів Colab/GPU.\n"
                "Після відновлення лімітів надішли команду /resume або 'продовжити генерацію'."
                f"{partial_hint}"
            )
            _flush_generation_stats()
            return None
            
        except Exception as e:
            logger.error(f"[Orchestrator] Audio drama generation failed: {e}")
            self._save_generation_state(
                status="failed",
                reason=str(e),
                current_scene_id=current_scene_id,
                current_order_in_scene=current_order_in_scene,
                processed_dialogues=processed_dialogues,
                total_dialogues=total_dialogues
            )
            _flush_generation_stats()
            return None
    
    async def _generate_dialogue_audio(
        self,
        dialogue: DialogueLine,
        scenario: Scenario,
        output_dir: str
    ) -> Optional[str]:
        """Генерує аудіо для однієї репліки."""
        character = scenario.get_character(dialogue.character_name)
        
        if not character:
            # Створюємо fallback персонажа
            normalized_name = canonical_character_name(dialogue.character_name) or dialogue.character_name
            role = "narrator" if is_narrator_alias(normalized_name) else "supporting"
            gender = resolve_character_gender(normalized_name, role, None)
            character = Character(
                name=normalized_name,
                role=role,
                gender=gender,
                voice_description=self._default_voice_description(normalized_name, role, gender)
            )
            dialogue.character_name = character.name
            self._voice_manager.register_character(character)
            scenario.characters.append(character)
        elif self._voice_manager.get_character(character.name) is None:
            self._voice_manager.register_character(character)
        
        # Отримуємо voice reference
        if self._clone_engine and not self._engine_is_loaded(self._clone_engine):
            await self._ensure_clone_engine_loaded()
        dialogue_language = self._resolve_language_for_text(dialogue.text)
        voice_prompt = await self._voice_manager.get_voice_reference(character.name, dialogue_language)
        strict_reference_mode = bool(self.config.strict_reference_mode_enabled)
        has_reference_declared = bool(character.reference_audio_path)
        has_reference_file = bool(has_reference_declared and os.path.exists(character.reference_audio_path))
        if strict_reference_mode and has_reference_declared and not voice_prompt:
            reason = "reference_file_missing" if not has_reference_file else "voice_prompt_build_failed"
            detail = (
                f"Strict reference mode: cannot use fallback voice for '{character.name}'. "
                f"reason={reason}, reference={character.reference_audio_path}"
            )
            raise MissingReferenceVoiceError(
                character_name=character.name,
                scene_id=dialogue.scene_id,
                order_in_scene=dialogue.order_in_scene,
                reference_audio_path=character.reference_audio_path or "",
                reason=reason,
                message=detail
            )
        dialogue_role = "narrator" if (character.role == "narrator" or is_narrator_alias(character.name)) else character.role
        dialogue_emotion = normalize_emotion_label(dialogue.emotion, dialogue.text)
        dialogue.emotion = dialogue_emotion
        source_tts_text = ensure_terminal_punctuation(dialogue.text, dialogue_role)
        dialogue.text = source_tts_text
        tts_text_candidate = self._prepare_tts_text(source_tts_text, dialogue_language)
        speech_speed = emotion_speed_multiplier(dialogue_emotion, dialogue_role)
        # Keep fallback VoiceDesign prompt stable per character.
        # Per-line prompt mutations cause timbre drift across dialogues.
        stable_design_prompt = build_gender_aware_prompt(character)
        quality_guard_enabled = bool(self._quality_guard and self._quality_guard.enabled)
        max_attempts = 1
        if quality_guard_enabled and self._quality_guard:
            max_attempts = max(1, int(self._quality_guard.max_attempts))
        
        # Формуємо ім'я файлу
        safe_name = dialogue.character_name.replace(' ', '_').lower()
        output_path = os.path.join(
            output_dir,
            f"scene{dialogue.scene_id}_{safe_name}_{dialogue.order_in_scene}.wav"
        )

        for attempt in range(1, max_attempts + 1):
            tts_text = tts_text_candidate
            try:
                # Генеруємо через clone engine
                if self._clone_engine and voice_prompt:
                    audio, sr = await asyncio.to_thread(
                        self._clone_engine.generate,
                        tts_text,
                        {
                            "voice_prompt": voice_prompt,
                            "speed": speech_speed,
                            "seed": character.seed
                        },
                        dialogue_language
                    )
                # Fallback на design engine
                elif self._design_engine:
                    if not self._engine_is_loaded(self._design_engine):
                        await self._ensure_design_engine_loaded()
                    audio, sr = await asyncio.to_thread(
                        self._design_engine.generate,
                        tts_text,
                        {
                            "seed": character.seed,
                            "prompt": stable_design_prompt,
                            "gender": character.gender,
                            "speed": speech_speed
                        },
                        dialogue_language
                    )
                else:
                    logger.warning(f"[Orchestrator] No TTS engine available for {dialogue.character_name}")
                    return None

                if audio is None:
                    logger.warning(
                        f"[Orchestrator] Empty audio for {dialogue.character_name} "
                        f"(scene={dialogue.scene_id}, order={dialogue.order_in_scene}, attempt={attempt})"
                    )
                    continue

                # Зберігаємо
                self._save_audio(audio, sr, output_path)

                # Якщо quality-guard вимкнений - завершуємо одразу.
                if not quality_guard_enabled or not self._quality_guard:
                    return output_path

                report = await asyncio.to_thread(
                    self._quality_guard.analyze_dialogue,
                    output_path,
                    tts_text,
                    dialogue_language,
                    attempt
                )
                report_path = self._write_dialogue_quality_report(
                    dialogue=dialogue,
                    character_name=dialogue.character_name,
                    report=report,
                    expected_text=dialogue.text,
                    tts_text=tts_text
                )
                issues_preview = ", ".join(issue.code for issue in report.issues[:5]) if report.issues else "none"
                logger.info(
                    f"[QualityGuard] scene={dialogue.scene_id} line={dialogue.order_in_scene} "
                    f"attempt={attempt}/{max_attempts} passed={report.passed} "
                    f"sim={report.similarity:.2f} issues={issues_preview}"
                )
                if report_path:
                    logger.info(f"[QualityGuard] Report saved: {report_path}")

                if report.passed or attempt >= max_attempts:
                    if report.issues:
                        logger.warning(
                            f"[QualityGuard] Accepting line with issues after attempt {attempt}: "
                            f"{', '.join(issue.code for issue in report.issues)}"
                        )
                    return output_path

                learned_rules = self._quality_guard.learn_pronunciation_rules(
                    self._pronunciation_manager,
                    dialogue_language,
                    report.issues
                )
                if learned_rules:
                    logger.info(
                        f"[QualityGuard] Learned pronunciation rules: {learned_rules} "
                        f"for scene={dialogue.scene_id} line={dialogue.order_in_scene}"
                    )

                next_text = (report.corrected_text or "").strip()
                if not next_text:
                    logger.warning(
                        f"[QualityGuard] No corrected text for retry, using current audio "
                        f"(scene={dialogue.scene_id}, line={dialogue.order_in_scene})"
                    )
                    return output_path

                next_prepared = self._prepare_tts_text(next_text, dialogue_language)
                if next_prepared.strip() == tts_text.strip():
                    logger.warning(
                        f"[QualityGuard] Corrected text unchanged, skipping retry "
                        f"(scene={dialogue.scene_id}, line={dialogue.order_in_scene})"
                    )
                    return output_path
                tts_text_candidate = next_prepared
            except Exception as e:
                if is_resource_limit_error(e):
                    raise ResourceLimitPauseError(str(e)) from e
                logger.error(f"[Orchestrator] Dialogue generation failed: {e}")
                return None
        
        return None
    
    async def _generate_background_sounds(
        self,
        scene: Scene,
        output_dir: str
    ) -> None:
        """Генерує фонові звуки для сцени."""
        if any(c.local_path and os.path.exists(c.local_path) for c in scene.sound_cues):
            return
        scene.sound_cues = []
        
        sound_cues: List[SoundCue] = []
        
        # 1) Описи звуків через Gemini.
        if self._gemini_client:
            try:
                sound_cues = await self._gemini_client.describe_sounds(
                    Scenario(
                        title="temp",
                        scenes=[scene],
                        characters=[]
                    ),
                    language=self.get_current_language()
                )
            except Exception as e:
                logger.warning(f"[Orchestrator] Gemini sound-cues failed for scene {scene.id}: {e}")
        
        # 2) Fallback, якщо Gemini не повернула sound cues.
        if not sound_cues:
            setting_text = (scene.setting or scene.title or "ambient background").strip()
            sound_cues = [
                SoundCue(
                    scene_id=scene.id,
                    description=setting_text,
                    keywords=["ambient", "atmosphere", "background"],
                    start_time=0.0,
                    duration=max(15.0, scene.get_total_duration() if hasattr(scene, "get_total_duration") else 30.0),
                    volume=0.18
                )
            ]
        
        # 3) Пошук у Pexels + fallback на синтетичний фон.
        for idx, cue in enumerate(sound_cues, start=1):
            output_path = os.path.join(output_dir, f"scene{scene.id}_bg_{idx}.wav")
            local_path: Optional[str] = None
            if self._pexels_client:
                queries: List[str] = []
                if cue.keywords:
                    queries.append(" ".join(cue.keywords[:3]))
                scene_hint = re.sub(r"\s+", " ", (scene.setting or scene.title or "").strip())
                if scene_hint:
                    queries.append(scene_hint)
                queries.extend(["ambient background", "cinematic ambience"])
                
                for query in queries:
                    try:
                        results = await self._pexels_client.search_sounds(query)
                    except Exception as e:
                        logger.warning(f"[Orchestrator] Pexels search failed for '{query}': {e}")
                        results = []
                    if not results:
                        continue
                    try:
                        local_path = await self._pexels_client.get_video_audio(results[0].url, output_path)
                    except Exception as e:
                        logger.warning(f"[Orchestrator] Pexels download/extract failed for '{query}': {e}")
                        local_path = None
                    if local_path and os.path.exists(local_path):
                        break
            
            if not local_path:
                local_path = self._create_synthetic_background(cue, output_path)
            
            if local_path and os.path.exists(local_path):
                cue.local_path = local_path
                scene.sound_cues.append(cue)
    
    def _create_synthetic_background(self, cue: SoundCue, output_path: str) -> Optional[str]:
        """Створює легкий fallback-фон, якщо Pexels недоступний."""
        try:
            sr = 24000
            duration = float(cue.duration or 20.0)
            duration = max(8.0, min(duration, 180.0))
            samples = int(sr * duration)
            if samples <= 0:
                return None
            
            # М'який low-passed noise + слабкий тон для "атмосфери".
            noise = np.random.normal(0.0, 1.0, samples).astype(np.float32)
            kernel = np.ones(512, dtype=np.float32) / 512.0
            smooth = np.convolve(noise, kernel, mode="same")
            t = np.arange(samples, dtype=np.float32) / float(sr)
            tone = np.sin(2.0 * np.pi * 120.0 * t).astype(np.float32)
            volume = max(0.05, min(float(cue.volume or 0.18), 0.4))
            audio = (smooth * (0.12 * volume)) + (tone * (0.02 * volume))
            peak = float(np.max(np.abs(audio))) if audio.size else 0.0
            if peak > 0:
                audio = audio / peak * 0.85
            self._save_audio(audio, sr, output_path)
            return output_path if os.path.exists(output_path) else None
        except Exception as e:
            logger.warning(f"[Orchestrator] Synthetic background failed: {e}")
            return None
    
    async def _check_checkpoints(self) -> None:
        """Перевіряє та відправляє контрольні точки."""
        # Контрольна точка 1 хвилина
        if (self.config.checkpoint_1min and 
            self._generated_duration >= 60 and 
            60 not in self._checkpoint_targets_sent):
            
            created = await self._create_checkpoint(60)
            if created:
                await self._wait_for_checkpoint_resume(60)
        
        # Контрольна точка 30 хвилин
        if (self.config.checkpoint_30min and 
            self._generated_duration >= 1800 and 
            1800 not in self._checkpoint_targets_sent):
            
            await self._create_checkpoint(1800)
    
    async def _create_checkpoint(self, duration: float) -> bool:
        """Створює контрольну точку."""
        if not self._current_work_dir:
            return False
        
        checkpoint_dir = os.path.join(self._current_work_dir, "checkpoints")
        os.makedirs(checkpoint_dir, exist_ok=True)
        
        checkpoint_path = os.path.join(
            checkpoint_dir,
            f"checkpoint_{duration/60:.0f}min.wav"
        )
        
        # Збираємо поточний стан
        if self._audio_assembler and self._current_scenario:
            mixing_config = MixingConfig()
            
            # Збираємо тільки завершені сцени
            completed_scenes = [
                s for s in self._current_scenario.scenes
                if all(dl.audio_path for dl in s.dialogue_lines)
            ]
            
            if completed_scenes:
                self._audio_assembler.assemble_from_scenes(
                    completed_scenes,
                    os.path.join(self._current_work_dir, "dialogue_chunks"),
                    mixing_config,
                    checkpoint_path
                )
                
                checkpoint = Checkpoint(
                    timestamp=datetime.now().isoformat(),
                    duration=duration,
                    audio_path=checkpoint_path,
                    scene_id=completed_scenes[-1].id,
                    message=(
                        f"Контрольна точка: {duration/60:.0f} хв\n"
                        f"Згенеровано: {self._generated_duration:.1f} сек"
                    )
                )
                
                self._checkpoints.append(checkpoint)
                self._checkpoint_targets_sent.add(int(duration))
                
                # Відправляємо в Telegram
                if self._bot:
                    checkpoint_dir = os.path.dirname(checkpoint_path)
                    await self._bot.send_checkpoint(
                        checkpoint_path,
                        self._generated_duration,
                        checkpoint.message,
                        self.format_drive_hint(checkpoint_dir)
                    )
                return True
        return False

    async def _wait_for_checkpoint_resume(self, duration: float) -> None:
        """Ставить генерацію на паузу до рішення користувача на checkpoint."""
        if int(duration) != 60:
            return
        if not self._bot:
            logger.info("[Orchestrator] No Telegram bot instance, checkpoint pause skipped")
            return
        if self._stop_requested:
            return
        
        self._waiting_for_checkpoint_resume = True
        self._checkpoint_resume_event.clear()
        await self._send_progress("⏸ Генерацію поставлено на паузу на 1 хвилині. Натисни кнопку 'Генерація' для продовження.")
        await self._checkpoint_resume_event.wait()
        self._waiting_for_checkpoint_resume = False
    
    async def import_reference_voice(
        self,
        character_name: str,
        source_audio_path: str,
        description_override: str = "",
        gender_override: Optional[str] = None
    ) -> Tuple[Optional[str], Optional[str]]:
        """Імпортує користувацький референс голосу (audio + markdown) і прив'язує до персонажа."""
        if not self._voice_manager or not self._current_work_dir:
            return None, None
        if not source_audio_path or not os.path.exists(source_audio_path):
            return None, None
        
        resolved_name = character_name
        character = self._voice_manager.get_character(character_name)
        if character is None:
            if self._current_scenario:
                character = self._current_scenario.get_character(character_name)
            if character is None:
                role = "narrator" if is_narrator_alias(character_name) else "supporting"
                gender = resolve_character_gender(character_name, role, gender_override)
                character = Character(
                    name=canonical_character_name(character_name) or character_name,
                    role=role,
                    gender=gender,
                    voice_description=self._default_voice_description(character_name, role, gender)
                )
                if self._current_scenario:
                    self._current_scenario.characters.append(character)
            self._voice_manager.register_character(character)
        if character:
            resolved_name = character.name
        
        existing_task = self._reference_render_tasks.get(resolved_name)
        if existing_task and not existing_task.done():
            existing_task.cancel()
        self._reference_render_errors.pop(resolved_name, None)
        
        references_dir = os.path.join(self._current_work_dir, "references")
        os.makedirs(references_dir, exist_ok=True)
        safe_name = re.sub(r"[^a-zA-Z0-9а-яА-ЯіїєґІЇЄҐ_\\-]+", "_", resolved_name).strip("_").lower() or "character"
        ext = Path(source_audio_path).suffix.lower() or ".wav"
        
        idx = 1
        while True:
            base_name = f"{safe_name}_reference_v{idx}"
            target_audio = os.path.join(references_dir, f"{base_name}{ext}")
            target_md = os.path.join(references_dir, f"{base_name}.md")
            if not os.path.exists(target_audio) and not os.path.exists(target_md):
                break
            idx += 1
        
        shutil.copy2(source_audio_path, target_audio)
        persisted_audio = self._ensure_file_on_drive(target_audio, "references")
        
        character = self._voice_manager.get_character(resolved_name)
        if character is None:
            return None, None
        
        if gender_override:
            character.gender = normalize_gender(gender_override, character.gender or "male")
            character.gender_locked = True
        else:
            desc_gender = extract_gender_from_text(description_override) if description_override else None
            if not character.gender_locked:
                character.gender = resolve_character_gender(
                    character.name,
                    character.role,
                    desc_gender or character.gender
                )
        
        if description_override:
            desc_clean = description_override.strip()
            if len(desc_clean.split()) >= 8:
                # Для .md/детального опису користувача беремо опис як основний.
                character.voice_description = desc_clean
            else:
                # Короткі інструкції типу "кавказький акцент" додаємо як модифікатор.
                character.voice_description = f"{character.voice_description}. {desc_clean}".strip(". ")
        if not character.voice_description:
            character.voice_description = self._default_voice_description(character.name, character.role, character.gender)
        
        reference_language = self.get_current_language()
        character.reference_audio_path = persisted_audio
        character.reference_text = character.reference_text or self._voice_manager._generate_test_text(
            character,
            15.0,
            reference_language
        )
        character.voice_prompt = None
        await self._voice_manager.get_voice_reference(character.name, reference_language)
        self._voice_manager.save_voice_to_library(character, persisted_audio)
        self.mark_character_audio_dirty(character.name)
        
        md_content = [
            f"# Voice Reference: {character.name}",
            "",
            f"- role: {character.role}",
            f"- gender: {character.gender}",
            f"- language: {reference_language}",
            f"- source_audio: {persisted_audio}",
            f"- created_at: {datetime.now().isoformat()}",
            "",
            "## Voice Description",
            character.voice_description or "",
            "",
            "## Reference Text",
            character.reference_text or "",
        ]
        
        target_md_abs = os.path.join(os.path.dirname(persisted_audio), f"{Path(persisted_audio).stem}.md")
        with open(target_md_abs, "w", encoding="utf-8") as f:
            f.write("\n".join(md_content).strip() + "\n")
        
        self._save_scenario(self._current_scenario) if self._current_scenario else None
        return persisted_audio, target_md_abs
    
    async def regenerate_character_voice(
        self,
        character_name: str,
        modification: str = "",
        gender_override: Optional[str] = None,
        preview: Optional[bool] = None
    ) -> Optional[str]:
        """Перегенерує голос персонажа."""
        if not self._voice_manager or not self._current_work_dir:
            return None
        
        if preview is None:
            resolved = self._resolve_voice_approval_name(character_name) or character_name
            preview = not self._voice_approvals.get(resolved, True)
        output_subfolder = "preview_voices" if preview else "test_voices"
        output_dir = os.path.join(self._current_work_dir, output_subfolder)
        target_duration = self.config.preview_voice_duration if preview else self.config.test_voice_duration
        os.makedirs(output_dir, exist_ok=True)
        
        try:
            if not await self._ensure_design_engine_loaded():
                logger.error("[Orchestrator] VoiceDesign engine unavailable for regeneration")
                return None
            new_path = await self._voice_manager.regenerate_voice(
                character_name,
                modification,
                gender_override,
                output_dir,
                language=self.get_current_language(),
                duration_sec=target_duration
            )
            persisted_path = self._ensure_file_on_drive(new_path, output_subfolder)
            self._sync_character_reference_path(character_name, old_path=new_path, new_path=persisted_path)
            character = self._voice_manager.get_character(character_name)
            if character:
                self._save_voice_profile_markdown(
                    character,
                    persisted_path,
                    stage="preview" if preview else "approved_reference"
                )
                self.mark_character_audio_dirty(character.name)
            else:
                self.mark_character_audio_dirty(character_name)
            return persisted_path
        except Exception as e:
            logger.error(f"[Orchestrator] Voice regeneration failed: {e}")
            return None

    def _ensure_file_on_drive(self, source_path: str, subfolder: str = "artifacts") -> str:
        """Гарантує, що файл присутній у Google Drive, і повертає шлях у Drive."""
        if not source_path or not os.path.exists(source_path):
            return source_path
        
        abs_source = os.path.abspath(source_path)
        drive_root = "/content/drive/MyDrive"
        
        # Якщо вже на Drive - додаткових дій не потрібно.
        if abs_source.startswith(drive_root):
            return abs_source
        
        # Якщо Drive не змонтований - повертаємо локальний шлях.
        if not os.path.exists(drive_root):
            return abs_source
        
        target_base = self.config.drive_base_path
        if not str(target_base).startswith(drive_root):
            target_base = os.path.join(drive_root, "audio_drama")
        
        export_dir = os.path.join(target_base, "exports", subfolder)
        os.makedirs(export_dir, exist_ok=True)
        
        filename = os.path.basename(abs_source)
        target_path = os.path.join(export_dir, filename)
        if os.path.exists(target_path):
            stem, ext = os.path.splitext(filename)
            target_path = os.path.join(
                export_dir,
                f"{stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
            )
        
        try:
            shutil.copy2(abs_source, target_path)
            logger.info(f"[Orchestrator] Copied artifact to Drive: {target_path}")
            return target_path
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to copy artifact to Drive: {e}")
            return abs_source

    def _sync_character_reference_path(self, character_name: str, old_path: str, new_path: str) -> None:
        """Синхронізує reference path персонажа після копіювання в Drive."""
        if not self._voice_manager or not new_path:
            return
        if old_path == new_path:
            return
        
        character = self._voice_manager.get_character(character_name)
        if character:
            if character.reference_audio_path == old_path:
                character.reference_audio_path = new_path
        
        registry = getattr(self._voice_manager, "_test_voices", {})
        if character and character.name in registry and registry[character.name]:
            if registry[character.name][-1] == old_path:
                registry[character.name][-1] = new_path
    
    def _create_work_directory(self, title: str) -> str:
        """Створює робочу директорію для проекту."""
        safe_title = re.sub(r'[^\w\s-]', '', title).strip().replace(' ', '_')
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        work_dir = os.path.join(self.config.drive_base_path, f"{safe_title}_{timestamp}")
        
        # Створюємо піддиректорії
        subdirs = [
            "preview_voices",
            "test_voices",
            "references",
            "background_sounds",
            "dialogue_chunks",
            "checkpoints",
            self._quality_reports_dir_name
        ]
        
        for subdir in subdirs:
            os.makedirs(os.path.join(work_dir, subdir), exist_ok=True)
        
        logger.info(f"[Orchestrator] Work directory created: {work_dir}")
        return work_dir

    def _quality_reports_dir(self) -> Optional[str]:
        """Повертає директорію звітів якості."""
        if not self._current_work_dir:
            return None
        path = os.path.join(self._current_work_dir, self._quality_reports_dir_name)
        os.makedirs(path, exist_ok=True)
        return path

    def _write_dialogue_quality_report(
        self,
        dialogue: DialogueLine,
        character_name: str,
        report: DialogueQualityReport,
        expected_text: str,
        tts_text: str
    ) -> Optional[str]:
        """Зберігає звіт перевірки якості репліки у JSON."""
        reports_dir = self._quality_reports_dir()
        if not reports_dir:
            return None

        safe_name = re.sub(r"[^\w-]+", "_", (character_name or "character").strip().lower())
        file_name = (
            f"scene{int(dialogue.scene_id)}_"
            f"{safe_name}_"
            f"{int(dialogue.order_in_scene)}_"
            f"attempt{int(report.attempt)}.json"
        )
        report_path = os.path.join(reports_dir, file_name)
        payload: Dict[str, Any] = {
            "scene_id": int(dialogue.scene_id),
            "order_in_scene": int(dialogue.order_in_scene),
            "character_name": character_name,
            "language": normalize_language_code(self._resolve_language_for_text(dialogue.text), self.get_current_language()),
            "source_text": dialogue.text,
            "expected_text": expected_text,
            "tts_text": tts_text,
            "report": report.to_dict(),
            "saved_at": datetime.now().isoformat(),
        }
        try:
            with open(report_path, "w", encoding="utf-8") as f:
                json.dump(payload, f, ensure_ascii=False, indent=2)
            return report_path
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to write quality report: {e}")
            return None

    @staticmethod
    def _drive_query_escape(value: str) -> str:
        """Екранує значення для q-параметра Google Drive API."""
        return str(value).replace("\\", "\\\\").replace("'", "\\'")

    def _get_google_drive_service(self):
        """Повертає ініціалізований Google Drive API service (best-effort)."""
        if self._drive_service_failed:
            return None
        if self._drive_service is not None:
            return self._drive_service

        try:
            from googleapiclient.discovery import build
            import google.auth

            credentials, _ = google.auth.default(
                scopes=["https://www.googleapis.com/auth/drive.readonly"]
            )
            self._drive_service = build(
                "drive",
                "v3",
                credentials=credentials,
                cache_discovery=False
            )
            return self._drive_service
        except Exception as e:
            self._drive_service_failed = True
            if not self._drive_service_notice_sent:
                logger.info(f"[Orchestrator] Google Drive API unavailable for direct links: {e}")
                self._drive_service_notice_sent = True
            return None

    def _resolve_drive_item_id(self, path: str) -> Optional[str]:
        """Резолвить file/folder ID у Google Drive для абсолютного шляху в /content/drive/MyDrive."""
        if not path:
            return None

        abs_path = os.path.abspath(path).replace("\\", "/").rstrip("/")
        drive_prefix = "/content/drive/MyDrive"
        if not abs_path.startswith(drive_prefix):
            return None

        is_dir = os.path.isdir(path)
        cache_key = (abs_path, is_dir)
        if cache_key in self._drive_item_id_cache:
            return self._drive_item_id_cache[cache_key]

        service = self._get_google_drive_service()
        if service is None:
            return None

        rel_path = abs_path[len(drive_prefix):].strip("/")
        if not rel_path:
            self._drive_item_id_cache[cache_key] = "root"
            return "root"

        segments = [segment for segment in rel_path.split("/") if segment]
        parent_id = "root"

        try:
            for idx, segment in enumerate(segments):
                is_last = idx == len(segments) - 1
                escaped_name = self._drive_query_escape(segment)
                query_parts = [
                    f"'{parent_id}' in parents",
                    f"name = '{escaped_name}'",
                    "trashed = false",
                ]
                if not is_last or is_dir:
                    query_parts.append("mimeType = 'application/vnd.google-apps.folder'")
                else:
                    query_parts.append("mimeType != 'application/vnd.google-apps.folder'")

                response = service.files().list(
                    q=" and ".join(query_parts),
                    spaces="drive",
                    fields="files(id,mimeType,modifiedTime)",
                    pageSize=10,
                    orderBy="modifiedTime desc",
                    includeItemsFromAllDrives=False,
                    supportsAllDrives=False,
                ).execute()
                matches = response.get("files", [])
                if not matches:
                    self._drive_item_id_cache[cache_key] = None
                    return None

                item_id = matches[0].get("id")
                if not item_id:
                    self._drive_item_id_cache[cache_key] = None
                    return None
                parent_id = item_id

            self._drive_item_id_cache[cache_key] = parent_id
            return parent_id
        except Exception as e:
            logger.debug(f"[Orchestrator] Drive ID resolution failed for {abs_path}: {e}")
            self._drive_item_id_cache[cache_key] = None
            return None

    def _build_google_drive_url(self, path: str) -> Optional[str]:
        """Будує пряме URL-посилання на файл/теку в Google Drive."""
        if not path:
            return None

        abs_path = os.path.abspath(path)
        if abs_path in self._drive_url_cache:
            return self._drive_url_cache[abs_path]

        item_id = self._resolve_drive_item_id(abs_path)
        if not item_id:
            self._drive_url_cache[abs_path] = None
            return None

        if item_id == "root":
            url = "https://drive.google.com/drive/u/0/my-drive"
        elif os.path.isdir(abs_path):
            url = f"https://drive.google.com/drive/folders/{item_id}?usp=sharing"
        else:
            url = f"https://drive.google.com/file/d/{item_id}/view?usp=sharing"

        self._drive_url_cache[abs_path] = url
        return url

    def _build_google_drive_fallback_url(self, path: str) -> Optional[str]:
        """Fallback-URL для Drive, коли прямий item-id недоступний."""
        if not path:
            return None
        abs_path = os.path.abspath(path).replace("\\", "/")
        drive_prefix = "/content/drive/MyDrive"
        if not abs_path.startswith(drive_prefix):
            return "https://drive.google.com/drive/u/0/my-drive"

        rel_path = abs_path[len(drive_prefix):].strip("/")
        if not rel_path:
            return "https://drive.google.com/drive/u/0/my-drive"

        leaf = os.path.basename(abs_path.rstrip("/"))
        parent = os.path.basename(os.path.dirname(abs_path.rstrip("/")))
        query_parts: List[str] = []
        if rel_path:
            query_parts.append(f"\"{rel_path}\"")
        if leaf:
            query_parts.append(f"\"{leaf}\"")
        if parent and parent.lower() not in {"mydrive", "drive"}:
            query_parts.append(f"\"{parent}\"")
        query = " ".join(part for part in query_parts if part).strip()
        if not query:
            query = rel_path or leaf or "audio_drama"
        return f"https://drive.google.com/drive/u/0/search?q={quote(query)}"

    def format_drive_hint(self, path: str) -> str:
        """Повертає зручний текст із підказкою для відкриття артефактів у Google Drive."""
        if not path:
            return ""
        abs_path = os.path.abspath(path)
        target_dir = abs_path if os.path.isdir(abs_path) else os.path.dirname(abs_path)
        files_count = 0
        try:
            if target_dir and os.path.isdir(target_dir):
                files_count = len([n for n in os.listdir(target_dir) if not n.startswith(".")])
        except Exception:
            files_count = 0
        drive_prefix = "/content/drive/MyDrive"
        if abs_path.startswith(drive_prefix):
            rel_path = abs_path[len(drive_prefix):].lstrip("/").replace("\\", "/")
            direct_url = self._build_google_drive_url(abs_path)
            link_url = direct_url or self._build_google_drive_fallback_url(abs_path) or "https://drive.google.com/drive/u/0/my-drive"
            if rel_path:
                return (
                    f"📂 Google Drive шлях: MyDrive/{rel_path}\n"
                    f"📦 Файлів у теці: {files_count}\n"
                    "🔗 Відкрити в Drive:\n"
                    f"{link_url}"
                )
            return (
                f"📦 Файлів у теці: {files_count}\n"
                "🔗 Відкрити в Drive:\n"
                f"{link_url}"
            )
        return f"📂 Локальний шлях: {abs_path}"

    def _state_file_path(self, work_dir: Optional[str] = None) -> str:
        """Повертає шлях до файла стану генерації."""
        active_dir = work_dir or self._current_work_dir or self.config.drive_base_path
        return os.path.join(active_dir, self._state_file_name)

    def _save_generation_state(
        self,
        status: str,
        reason: str = "",
        current_scene_id: Optional[int] = None,
        current_order_in_scene: Optional[int] = None,
        processed_dialogues: int = 0,
        total_dialogues: int = 0,
        extra: Optional[Dict[str, Any]] = None
    ) -> None:
        """Зберігає стан поточної генерації у Drive."""
        if not self._current_work_dir:
            return
        
        scenario_path = os.path.join(self._current_work_dir, "scenario.json")
        payload: Dict[str, Any] = {
            "version": 1,
            "status": status,
            "reason": reason,
            "work_dir": self._current_work_dir,
            "scenario_path": scenario_path,
            "current_scene_id": current_scene_id,
            "current_order_in_scene": current_order_in_scene,
            "processed_dialogues": processed_dialogues,
            "total_dialogues": total_dialogues,
            "generated_duration": float(self._generated_duration or 0.0),
            "language": self.get_current_language(),
            "dirty_characters": sorted(self._dirty_characters),
            "updated_at": datetime.now().isoformat(),
        }
        if extra:
            payload["extra"] = extra
        
        try:
            with open(self._state_file_path(), "w", encoding="utf-8") as f:
                json.dump(payload, f, ensure_ascii=False, indent=2)
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to save generation state: {e}")

    def _load_generation_state(self, work_dir: str) -> Dict[str, Any]:
        """Завантажує стан генерації з work_dir."""
        state_path = self._state_file_path(work_dir)
        if not os.path.exists(state_path):
            return {}
        try:
            with open(state_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            return data if isinstance(data, dict) else {}
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to load generation state: {e}")
            return {}

    def _write_resume_pointer(self, work_dir: str, status: str, reason: str = "") -> None:
        """Оновлює глобальний покажчик на останню paused-сесію."""
        try:
            os.makedirs(self.config.drive_base_path, exist_ok=True)
            payload = {
                "work_dir": work_dir,
                "status": status,
                "reason": reason,
                "updated_at": datetime.now().isoformat(),
            }
            with open(self._resume_pointer_file, "w", encoding="utf-8") as f:
                json.dump(payload, f, ensure_ascii=False, indent=2)
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to write resume pointer: {e}")

    def _read_resume_pointer(self) -> Dict[str, Any]:
        """Зчитує покажчик останньої paused-сесії."""
        if not os.path.exists(self._resume_pointer_file):
            return {}
        try:
            with open(self._resume_pointer_file, "r", encoding="utf-8") as f:
                data = json.load(f)
            return data if isinstance(data, dict) else {}
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to read resume pointer: {e}")
            return {}

    def _clear_resume_pointer(self) -> None:
        """Очищає покажчик paused-сесії після успішного завершення."""
        try:
            if os.path.exists(self._resume_pointer_file):
                os.remove(self._resume_pointer_file)
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to clear resume pointer: {e}")

    def _load_scenario_from_file(self, scenario_path: str) -> Optional[Scenario]:
        """Завантажує сценарій з JSON у повні dataclass-об'єкти."""
        if not scenario_path or not os.path.exists(scenario_path):
            return None
        try:
            with open(scenario_path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception as e:
            logger.error(f"[Orchestrator] Failed to read scenario file: {e}")
            return None
        
        try:
            scenes: List[Scene] = []
            for scene_data in data.get("scenes", []):
                dialogue_lines: List[DialogueLine] = []
                for dl in scene_data.get("dialogue_lines", []):
                    raw_text = dl.get("text", "")
                    raw_emotion = dl.get("emotion", "neutral")
                    dialogue_lines.append(DialogueLine(
                        character_name=dl.get("character_name", NARRATOR_NAME),
                        text=ensure_terminal_punctuation(
                            raw_text,
                            "narrator" if is_narrator_alias(dl.get("character_name", NARRATOR_NAME)) else "supporting"
                        ),
                        emotion=normalize_emotion_label(raw_emotion, raw_text),
                        scene_id=int(dl.get("scene_id", scene_data.get("id", 1) or 1)),
                        order_in_scene=int(dl.get("order_in_scene", 1) or 1),
                        audio_path=dl.get("audio_path"),
                        duration=float(dl.get("duration", 0.0) or 0.0)
                    ))
                sound_cues: List[SoundCue] = []
                for sc in scene_data.get("sound_cues", []):
                    sound_cues.append(SoundCue(
                        scene_id=int(sc.get("scene_id", scene_data.get("id", 1) or 1)),
                        description=sc.get("description", ""),
                        keywords=sc.get("keywords", []),
                        start_time=float(sc.get("start_time", 0.0) or 0.0),
                        duration=float(sc.get("duration", 30.0) or 30.0),
                        volume=float(sc.get("volume", 0.3) or 0.3),
                        local_path=sc.get("local_path"),
                        fade_in=float(sc.get("fade_in", 1.0) or 1.0),
                        fade_out=float(sc.get("fade_out", 1.0) or 1.0),
                        loop=bool(sc.get("loop", False))
                    ))
                scenes.append(Scene(
                    id=int(scene_data.get("id", 1) or 1),
                    title=scene_data.get("title", "Scene"),
                    setting=scene_data.get("setting", ""),
                    dialogue_lines=dialogue_lines,
                    sound_cues=sound_cues,
                    duration_estimate=float(scene_data.get("duration_estimate", 0.0) or 0.0)
                ))
            
            characters = [Character.from_dict(char_data) for char_data in data.get("characters", [])]
            scenario = Scenario(
                title=data.get("title", "Untitled"),
                author=data.get("author", "Unknown"),
                scenes=scenes,
                characters=characters,
                total_duration_estimate=float(data.get("total_duration_estimate", 0.0) or 0.0),
                created_at=data.get("created_at", datetime.now().isoformat())
            )
            return scenario
        except Exception as e:
            logger.error(f"[Orchestrator] Failed to parse scenario data: {e}")
            return None

    def _save_partial_audio_checkpoint(self, scenario: Scenario, tag: str = "auto_pause") -> Optional[str]:
        """Збирає partial-аудіо з уже готових реплік."""
        if not self._audio_assembler or not self._current_work_dir:
            return None
        try:
            checkpoint_dir = os.path.join(self._current_work_dir, "checkpoints")
            os.makedirs(checkpoint_dir, exist_ok=True)
            checkpoint_path = os.path.join(
                checkpoint_dir,
                f"{tag}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.wav"
            )
            self._audio_assembler.assemble_from_scenes(
                scenario.scenes,
                os.path.join(self._current_work_dir, "dialogue_chunks"),
                MixingConfig(),
                checkpoint_path
            )
            if os.path.exists(checkpoint_path) and os.path.getsize(checkpoint_path) > 0:
                return checkpoint_path
        except Exception as e:
            logger.warning(f"[Orchestrator] Failed to save partial checkpoint: {e}")
        return None

    async def resume_from_last_pause(self) -> Optional[str]:
        """Відновлює генерацію з останньої paused-сесії."""
        pointer = self._read_resume_pointer()
        work_dir = pointer.get("work_dir")
        pointer_status = pointer.get("status", "")
        pointer_reason = pointer.get("reason", "")
        if not work_dir or not os.path.exists(work_dir):
            await self._send_progress("ℹ️ Немає збереженої paused-сесії для продовження.")
            return None
        
        state = self._load_generation_state(work_dir)
        scenario_path = state.get("scenario_path") or os.path.join(work_dir, "scenario.json")
        scenario = self._load_scenario_from_file(scenario_path)
        if not scenario:
            await self._send_progress(f"❌ Не вдалося завантажити сценарій для resume: {scenario_path}")
            return None
        
        self._current_work_dir = work_dir
        self._current_scenario = scenario
        self._current_language = normalize_language_code(
            state.get("language"),
            normalize_language_code(self.config.language, "russian")
        )
        dirty_from_state = state.get("dirty_characters", [])
        self._dirty_characters.clear()
        if isinstance(dirty_from_state, list):
            for item in dirty_from_state:
                if not isinstance(item, str):
                    continue
                value = canonical_character_name(item) or item.strip()
                if value:
                    self._dirty_characters.add(value)
        self._stop_requested = False
        self._checkpoint_targets_sent.clear()
        previous_duration = float(state.get("generated_duration", 0.0) or 0.0)
        if previous_duration >= 60:
            self._checkpoint_targets_sent.add(60)
        if previous_duration >= 1800:
            self._checkpoint_targets_sent.add(1800)
        self._checkpoint_resume_event.set()
        self._waiting_for_checkpoint_resume = False
        self._pre_voice_reference_event.set()
        self._waiting_for_pre_voice_reference = False
        self._last_progress_percent = -1
        self._last_progress_min_bucket = -1
        if self._bot:
            self._bot.reset_progress_message()
        
        if self._voice_manager:
            for character in scenario.characters:
                self._voice_manager.register_character(character)
        
        await self._send_progress(
            f"▶️ Відновлюю генерацію з паузи.\n"
            f"📁 {work_dir}\n"
            f"🌐 Мова: {self.get_current_language()}"
        )

        if pointer_status == "paused_manual" and pointer_reason == "pre_voice_reference_upload_pending":
            require_manual_approval = bool(self._bot and self.config.auto_send_telegram)
            self.initialize_voice_approvals(scenario.characters, require_manual=require_manual_approval)
            self.sync_approvals_with_ready_references(scenario.characters)
            if require_manual_approval and scenario.characters and self.config.pre_voice_reference_stage_enabled:
                await self._send_progress(
                    "📥 Resume: етап завантаження референсів перед генерацією тестових голосів."
                )
                continue_from_pre_stage = await self._wait_for_pre_voice_reference_stage(scenario.characters)
                if not continue_from_pre_stage:
                    return None

            await self._send_progress("🎤 Resume: створення тестових голосів...")
            await self.generate_test_voices(scenario.characters)

            missing_references = self._collect_characters_missing_any_reference(scenario.characters)
            if missing_references:
                preview_missing = ", ".join(missing_references[:8])
                if len(missing_references) > 8:
                    preview_missing += f", ... (+{len(missing_references) - 8})"
                await self._send_progress(
                    "⏸ Resume: не вдалося створити прев'ю/референси для частини героїв.\n"
                    f"⌛ Без аудіо: {preview_missing}\n"
                    "Онови TTS runtime або завантаж свої голоси, потім /resume."
                )
                self._save_generation_state(
                    status="paused_manual",
                    reason="test_voice_generation_failed",
                    extra={"missing_characters": missing_references}
                )
                if self._current_work_dir:
                    self._write_resume_pointer(
                        self._current_work_dir,
                        status="paused_manual",
                        reason="test_voice_generation_failed"
                    )
                return None

            if self._bot and self._current_work_dir:
                test_voices_dir = os.path.join(self._current_work_dir, "test_voices")
                preview_dir = os.path.join(self._current_work_dir, "preview_voices")
                await self._bot.send_message(
                    (
                        f"📁 Папка з preview-голосами:\n{preview_dir}\n{self.format_drive_hint(preview_dir)}\n\n"
                        f"📁 Папка з повними референсами:\n{test_voices_dir}\n{self.format_drive_hint(test_voices_dir)}"
                    ).strip()
                )

            if require_manual_approval and scenario.characters:
                approved = await self.wait_for_all_voice_approvals()
                if not approved:
                    self._save_generation_state(status="paused_manual", reason="voice_approval_pending")
                    if self._current_work_dir:
                        self._write_resume_pointer(
                            self._current_work_dir,
                            status="paused_manual",
                            reason="voice_approval_pending"
                        )
                    return None
                refs_ready = await self.wait_for_reference_renders()
                if not refs_ready:
                    self._save_generation_state(status="paused_manual", reason="reference_render_pending")
                    if self._current_work_dir:
                        self._write_resume_pointer(
                            self._current_work_dir,
                            status="paused_manual",
                            reason="reference_render_pending"
                        )
                    return None
                await self._send_progress(
                    "✅ Resume: усі голоси затверджені і повні референси готові. "
                    "Продовжую генерацію аудіо."
                )
        
        if pointer_status == "paused_manual" and pointer_reason == "voice_approval_pending":
            self.initialize_voice_approvals(scenario.characters, require_manual=True)
            self.sync_approvals_with_ready_references(scenario.characters)
            await self._send_progress("⌛ Перед продовженням потрібно затвердити голоси.")
            approved = await self.wait_for_all_voice_approvals()
            if not approved:
                return None
            await self._send_progress("✅ Усі голоси затверджені, продовжую генерацію.")

        if pointer_status == "paused_manual" and pointer_reason == "test_voice_generation_failed":
            require_manual_approval = bool(self._bot and self.config.auto_send_telegram)
            self.initialize_voice_approvals(scenario.characters, require_manual=require_manual_approval)
            self.sync_approvals_with_ready_references(scenario.characters)
            await self._send_progress("🎤 Resume: повторна генерація тестових голосів...")
            await self.generate_test_voices(scenario.characters)

            missing_references = self._collect_characters_missing_any_reference(scenario.characters)
            if missing_references:
                preview_missing = ", ".join(missing_references[:8])
                if len(missing_references) > 8:
                    preview_missing += f", ... (+{len(missing_references) - 8})"
                await self._send_progress(
                    "⏸ Resume: частина голосів досі без аудіо.\n"
                    f"⌛ Без аудіо: {preview_missing}\n"
                    "Завантаж свої голоси або виправ TTS runtime, потім /resume."
                )
                self._save_generation_state(
                    status="paused_manual",
                    reason="test_voice_generation_failed",
                    extra={"missing_characters": missing_references}
                )
                if self._current_work_dir:
                    self._write_resume_pointer(
                        self._current_work_dir,
                        status="paused_manual",
                        reason="test_voice_generation_failed"
                    )
                return None

            if require_manual_approval and scenario.characters:
                approved = await self.wait_for_all_voice_approvals()
                if not approved:
                    self._save_generation_state(status="paused_manual", reason="voice_approval_pending")
                    if self._current_work_dir:
                        self._write_resume_pointer(
                            self._current_work_dir,
                            status="paused_manual",
                            reason="voice_approval_pending"
                        )
                    return None
                refs_ready = await self.wait_for_reference_renders()
                if not refs_ready:
                    self._save_generation_state(status="paused_manual", reason="reference_render_pending")
                    if self._current_work_dir:
                        self._write_resume_pointer(
                            self._current_work_dir,
                            status="paused_manual",
                            reason="reference_render_pending"
                        )
                    return None
                await self._send_progress("✅ Resume: голоси згенеровані та готові до подальшої генерації.")
        
        if pointer_status == "paused_manual" and pointer_reason == "reference_render_pending":
            full_dir = os.path.join(self._current_work_dir, "test_voices")
            os.makedirs(full_dir, exist_ok=True)
            for character in scenario.characters:
                self._schedule_full_reference_render(character.name, full_dir)
            refs_ready = await self.wait_for_reference_renders()
            if not refs_ready:
                return None
            await self._send_progress("✅ Повні референси голосів готові, продовжую генерацію.")

        if pointer_status == "paused_manual" and pointer_reason == "reference_voice_prompt_missing":
            extra = state.get("extra", {}) if isinstance(state.get("extra", {}), dict) else {}
            character_name = (extra.get("character_name") or "").strip()
            ref_path = (extra.get("reference_audio_path") or "").strip()
            ref_hint = f"\n📁 {ref_path}" if ref_path else ""
            if character_name:
                await self._send_progress(
                    f"⌛ Resume: перевіряю референс героя '{character_name}'.{ref_hint}\n"
                    "Якщо референс оновлено, продовжую генерацію."
                )
            else:
                await self._send_progress(
                    "⌛ Resume після strict-reference pause. Перевіряю оновлені референси і продовжую."
                )
        
        result = await self.generate_audio_drama(scenario)
        return result
    
    def _save_scenario(self, scenario: Scenario) -> None:
        """Зберігає сценарій у файл."""
        if not self._current_work_dir:
            return
        
        scenario_path = os.path.join(self._current_work_dir, "scenario.json")
        
        try:
            with open(scenario_path, 'w', encoding='utf-8') as f:
                json.dump(scenario.to_dict(), f, ensure_ascii=False, indent=2)
            
            logger.info(f"[Orchestrator] Scenario saved: {scenario_path}")
        except Exception as e:
            logger.error(f"[Orchestrator] Failed to save scenario: {e}")
    
    def _save_audio(self, audio: np.ndarray, sr: int, path: str) -> None:
        """Зберігає аудіо у файл."""
        try:
            from scipy.io import wavfile
            
            # Нормалізація
            if np.max(np.abs(audio)) > 1.0:
                audio = audio / np.max(np.abs(audio))
            
            audio_int16 = (audio * 32767).astype(np.int16)
            wavfile.write(path, sr, audio_int16)
        except Exception as e:
            logger.error(f"[Orchestrator] Failed to save audio: {e}")
            raise
    
    async def _send_progress(self, message: str) -> None:
        """Відправляє повідомлення про прогрес."""
        logger.info(f"[Orchestrator] {message}")
        
        if self._bot:
            await self._bot.send_message(message)
    
    async def handle_checkpoint_decision(self, continue_generation: bool) -> None:
        """Обробляє рішення користувача на контрольній точці."""
        if continue_generation:
            logger.info("[Orchestrator] User chose to continue generation")
            self._stop_requested = False
            self._checkpoint_resume_event.set()
            if self._waiting_for_pre_voice_reference:
                self._pre_voice_reference_event.set()
            self._save_generation_state(status="running", reason="checkpoint_continue")
        else:
            logger.info("[Orchestrator] User chose to stop generation")
            self._stop_requested = True
            self._checkpoint_resume_event.set()
            if self._waiting_for_pre_voice_reference:
                self._pre_voice_reference_event.set()
            self._save_generation_state(status="paused_manual", reason="checkpoint_stop")
            if self._current_work_dir:
                self._write_resume_pointer(self._current_work_dir, status="paused_manual", reason="checkpoint_stop")


# =============================================================================
# SECTION 9: COLAB INTEGRATION
# =============================================================================

_orchestrator_instance: Optional[AudioDramaOrchestrator] = None
_bot_instance: Optional[AudioDramaBot] = None
_bot_polling_thread: Optional[threading.Thread] = None
_bot_polling_lock = threading.Lock()


def is_colab_environment() -> bool:
    """Перевіряє чи запущено в Google Colab."""
    try:
        import google.colab
        return True
    except ImportError:
        return False


def mount_google_drive() -> bool:
    """Монтує Google Drive в Colab."""
    if not is_colab_environment():
        logger.warning("Not in Colab environment - Drive mount skipped")
        return False
    
    if os.path.exists('/content/drive/MyDrive'):
        logger.info("Google Drive already mounted")
        return True
    
    try:
        from google.colab import drive
        drive.mount('/content/drive')
        logger.info("Google Drive mounted")
        return True
    except Exception as e:
        logger.error(f"Drive mount error: {e}")
        return False


def get_colab_secrets(allow_userdata: Optional[bool] = None) -> Dict[str, Optional[str]]:
    """Отримує секрети з Colab."""
    secrets = {
        "telegram_token": None,
        "telegram_chat_id": None,
        "gemini_api_key": None,
        "pexels_api_key": None
    }
    
    # 1) Пріоритет: environment variables (ніколи не блокують)
    secrets["telegram_token"] = os.getenv("TELEGRAM_BOT_TOKEN")
    secrets["telegram_chat_id"] = os.getenv("TELEGRAM_CHAT_ID")
    secrets["gemini_api_key"] = os.getenv("GEMINI_API_KEY")
    secrets["pexels_api_key"] = os.getenv("PEXELS_API_KEY")
    
    for key, val in secrets.items():
        if val:
            logger.info(f"✓ {key.upper()} found in environment")
    
    if not is_colab_environment():
        logger.warning("Not in Colab environment - secrets not available")
        return secrets
    
    # 2) Опційно: Colab userdata (може блокуватися permission prompt)
    if allow_userdata is None:
        allow_userdata = os.getenv("AUDIO_DRAMA_USE_COLAB_USERDATA", "1").lower() in ("1", "true", "yes")
    if not allow_userdata:
        logger.info("[AudioDramaBot] Colab userdata secret lookup disabled (AUDIO_DRAMA_USE_COLAB_USERDATA=0)")
        return secrets
    
    try:
        from google.colab import userdata
        
        def _safe_userdata_get(secret_name: str, timeout_sec: float = 6.0) -> Optional[str]:
            """Безпечне отримання secret з timeout, щоб не зависати на permission prompt."""
            with ThreadPoolExecutor(max_workers=1) as executor:
                future = executor.submit(userdata.get, secret_name)
                try:
                    value = future.result(timeout=timeout_sec)
                    return value
                except FuturesTimeoutError:
                    logger.warning(f"⚠ {secret_name} read timed out after {timeout_sec:.0f}s")
                    return None
                except Exception:
                    return None
        
        try:
            if not secrets["telegram_token"]:
                secrets["telegram_token"] = _safe_userdata_get("TELEGRAM_BOT_TOKEN")
            if secrets["telegram_token"]:
                logger.info("✓ TELEGRAM_BOT_TOKEN found")
        except:
            logger.warning("⚠ TELEGRAM_BOT_TOKEN not found")
        
        try:
            if not secrets["telegram_chat_id"]:
                secrets["telegram_chat_id"] = _safe_userdata_get("TELEGRAM_CHAT_ID")
            if secrets["telegram_chat_id"]:
                logger.info("✓ TELEGRAM_CHAT_ID found")
        except:
            logger.warning("⚠ TELEGRAM_CHAT_ID not found")
        
        try:
            if not secrets["gemini_api_key"]:
                secrets["gemini_api_key"] = _safe_userdata_get("GEMINI_API_KEY")
            if secrets["gemini_api_key"]:
                logger.info("✓ GEMINI_API_KEY found")
        except:
            logger.warning("⚠ GEMINI_API_KEY not found")
        
        try:
            if not secrets["pexels_api_key"]:
                secrets["pexels_api_key"] = _safe_userdata_get("PEXELS_API_KEY")
            if secrets["pexels_api_key"]:
                logger.info("✓ PEXELS_API_KEY found")
        except:
            logger.warning("⚠ PEXELS_API_KEY not found")
        
    except Exception as e:
        logger.error(f"Failed to access Colab secrets: {e}")
    
    return secrets


def init_audio_drama_bot(
    drive_path: str = "/content/drive/MyDrive/audio_drama",
    cache_dir: str = "/content/cache/audio_drama",
    auto_mount_drive: bool = True,
    use_colab_userdata: Optional[bool] = None,
    init_timeout_sec: Optional[int] = None
) -> AudioDramaOrchestrator:
    """
    Ініціалізація Audio Drama Bot для Google Colab.
    
    Returns:
        AudioDramaOrchestrator: Ініціалізований оркестратор
    """
    global _orchestrator_instance
    
    if _orchestrator_instance is not None:
        logger.info("Orchestrator already initialized")
        print("[AudioDramaBot] Orchestrator already initialized", flush=True)
        return _orchestrator_instance
    
    logger.info("="*60)
    logger.info("AUDIO DRAMA BOT - INITIALIZATION")
    logger.info("="*60)
    print("[AudioDramaBot] Initialization started...", flush=True)
    
    # Монтуємо Drive
    if auto_mount_drive:
        mount_google_drive()
    
    # Отримуємо секрети
    secrets = get_colab_secrets(allow_userdata=use_colab_userdata)

    def _env_bool(name: str, default: bool) -> bool:
        raw = os.getenv(name)
        if raw is None:
            return default
        return str(raw).strip().lower() in {"1", "true", "yes", "on"}

    def _env_int(name: str, default: int, minimum: int = 1, maximum: int = 6) -> int:
        raw = os.getenv(name, str(default))
        try:
            value = int(raw)
        except Exception:
            value = default
        return max(minimum, min(maximum, value))

    def _env_float(name: str, default: float, minimum: float, maximum: float) -> float:
        raw = os.getenv(name, str(default))
        try:
            value = float(raw)
        except Exception:
            value = default
        return max(minimum, min(maximum, value))
    
    # Створюємо конфігурацію
    qa_enabled = _env_bool("AUDIO_DRAMA_QA_ENABLED", True)
    qa_max_attempts = _env_int("AUDIO_DRAMA_QA_MAX_ATTEMPTS", 2, minimum=1, maximum=5)
    qa_similarity = _env_float("AUDIO_DRAMA_QA_SIMILARITY", 0.78, minimum=0.30, maximum=0.99)
    qa_tail_ms = _env_float("AUDIO_DRAMA_QA_MIN_TAIL_SILENCE_MS", 120.0, minimum=20.0, maximum=800.0)
    qa_use_openai_asr = _env_bool("AUDIO_DRAMA_QA_USE_OPENAI_ASR", False)
    qa_asr_model = os.getenv("AUDIO_DRAMA_QA_ASR_MODEL", "gpt-4o-mini-transcribe").strip() or "gpt-4o-mini-transcribe"
    qa_use_local_asr = _env_bool("AUDIO_DRAMA_QA_USE_LOCAL_ASR", True)
    qa_local_asr_model = os.getenv("AUDIO_DRAMA_QA_LOCAL_ASR_MODEL", "small").strip() or "small"
    qa_local_asr_device = os.getenv("AUDIO_DRAMA_QA_LOCAL_ASR_DEVICE", "auto").strip().lower() or "auto"
    qa_local_asr_compute_type = os.getenv("AUDIO_DRAMA_QA_LOCAL_ASR_COMPUTE_TYPE", "auto").strip().lower() or "auto"
    strict_reference_mode = _env_bool("AUDIO_DRAMA_STRICT_REFERENCE_MODE", True)
    pre_voice_stage_enabled = _env_bool("AUDIO_DRAMA_PREVOICE_STAGE_ENABLED", True)
    pre_voice_stage_timeout = _env_int("AUDIO_DRAMA_PREVOICE_STAGE_TIMEOUT_SEC", 0, minimum=0, maximum=3600)
    one_script_mode = _env_bool("AUDIO_DRAMA_ONE_SCRIPT_MODE", True)
    allow_external_tts_fallback = _env_bool("AUDIO_DRAMA_ALLOW_EXTERNAL_TTS_FALLBACK", False)
    require_qwen_dual_models = _env_bool("AUDIO_DRAMA_REQUIRE_QWEN_DUAL_MODELS", True)
    qwen_precheck_on_init = _env_bool("AUDIO_DRAMA_QWEN_PRECHECK_ON_INIT", False)
    qwen_auto_install_on_load = _env_bool("AUDIO_DRAMA_QWEN_AUTO_INSTALL_ON_LOAD", True)

    config = OrchestratorConfig(
        one_script_mode=one_script_mode,
        drive_base_path=drive_path,
        cache_dir=cache_dir,
        prefer_internal_tts=True,
        allow_external_tts_fallback=False if one_script_mode else allow_external_tts_fallback,
        quality_guard_enabled=qa_enabled,
        quality_guard_max_attempts=qa_max_attempts,
        quality_guard_similarity_threshold=qa_similarity,
        quality_guard_min_tail_silence_ms=qa_tail_ms,
        quality_guard_use_openai_asr=qa_use_openai_asr,
        quality_guard_asr_model=qa_asr_model,
        quality_guard_use_local_asr=qa_use_local_asr,
        quality_guard_local_asr_model=qa_local_asr_model,
        quality_guard_local_asr_device=qa_local_asr_device,
        quality_guard_local_asr_compute_type=qa_local_asr_compute_type,
        strict_reference_mode_enabled=strict_reference_mode,
        pre_voice_reference_stage_enabled=pre_voice_stage_enabled,
        pre_voice_reference_timeout_sec=pre_voice_stage_timeout,
        require_qwen_dual_models=require_qwen_dual_models,
        qwen_precheck_on_init=qwen_precheck_on_init,
        qwen_auto_install_on_load=qwen_auto_install_on_load
    )
    
    # Створюємо оркестратор
    _orchestrator_instance = AudioDramaOrchestrator(config)
    
    init_kwargs = {
        "telegram_token": secrets["telegram_token"],
        "telegram_chat_id": secrets["telegram_chat_id"],
        "gemini_api_key": secrets["gemini_api_key"],
        "pexels_api_key": secrets["pexels_api_key"]
    }
    qwen_precheck_on_init = _env_bool("AUDIO_DRAMA_QWEN_PRECHECK_ON_INIT", False)
    if init_timeout_sec is None:
        if qwen_precheck_on_init and require_qwen_dual_models:
            default_timeout = 600
        else:
            default_timeout = 90
        timeout_sec = int(os.getenv("AUDIO_DRAMA_INIT_TIMEOUT_SEC", str(default_timeout)))
    else:
        timeout_sec = int(init_timeout_sec)
    timeout_sec = max(0, timeout_sec)

    def _run_initialize_sync(timeout_value: int) -> bool:
        """Запускає async initialize в sync-контексті з optional timeout."""
        try:
            loop = asyncio.get_event_loop()
            if loop.is_running():
                import nest_asyncio
                nest_asyncio.apply()
            init_coro = _orchestrator_instance.initialize(**init_kwargs)
            if timeout_value > 0:
                init_coro = asyncio.wait_for(init_coro, timeout=timeout_value)
            return loop.run_until_complete(init_coro)
        except RuntimeError:
            init_coro = _orchestrator_instance.initialize(**init_kwargs)
            if timeout_value > 0:
                init_coro = asyncio.wait_for(init_coro, timeout=timeout_value)
            return asyncio.run(init_coro)

    success = False
    try:
        success = _run_initialize_sync(timeout_sec)
        if success:
            logger.info("="*60)
            logger.info("✓ Audio Drama Bot initialized successfully!")
            logger.info("="*60)
            print("[AudioDramaBot] Initialization success", flush=True)
        else:
            logger.error("✗ Audio Drama Bot initialization failed")
            print("[AudioDramaBot] Initialization failed", flush=True)
    except asyncio.TimeoutError:
        logger.error(f"Audio Drama Bot initialization timed out after {timeout_sec}s")
        print(f"[AudioDramaBot] Initialization timed out after {timeout_sec}s", flush=True)
        if require_qwen_dual_models:
            raw_retry_timeout = os.getenv("AUDIO_DRAMA_INIT_RETRY_TIMEOUT_SEC", "300")
            try:
                retry_timeout_sec = int(raw_retry_timeout)
            except Exception:
                retry_timeout_sec = 300
            retry_timeout_sec = max(0, retry_timeout_sec)

            if retry_timeout_sec <= 0:
                logger.warning("[AudioDramaBot] Retry on init timeout disabled (AUDIO_DRAMA_INIT_RETRY_TIMEOUT_SEC<=0)")
                print("[AudioDramaBot] Retry disabled after timeout", flush=True)
            else:
                logger.warning(
                    "[AudioDramaBot] Retrying initialization after timeout "
                    f"(retry timeout={retry_timeout_sec}s)..."
                )
                print(
                    f"[AudioDramaBot] Retrying initialization (timeout={retry_timeout_sec}s)...",
                    flush=True
                )
                try:
                    success = _run_initialize_sync(retry_timeout_sec)
                    if success:
                        logger.info("✓ Audio Drama Bot initialized successfully (retry)")
                        print("[AudioDramaBot] Initialization success (retry)", flush=True)
                    else:
                        logger.error("✗ Audio Drama Bot initialization failed after retry")
                        print("[AudioDramaBot] Initialization failed after retry", flush=True)
                except asyncio.TimeoutError:
                    logger.error(
                        "Audio Drama Bot retry initialization timed out after "
                        f"{retry_timeout_sec}s"
                    )
                    print(
                        f"[AudioDramaBot] Retry initialization timed out after {retry_timeout_sec}s",
                        flush=True
                    )
                except Exception as retry_error:
                    logger.error(f"Audio Drama Bot retry initialization failed: {retry_error}")
                    print(f"[AudioDramaBot] Retry initialization error: {retry_error}", flush=True)
    except Exception as e:
        logger.error(f"Audio Drama Bot initialization failed: {e}")
        print(f"[AudioDramaBot] Initialization error: {e}", flush=True)
    
    return _orchestrator_instance


def get_orchestrator() -> Optional[AudioDramaOrchestrator]:
    """Повертає поточний оркестратор."""
    return _orchestrator_instance


async def run_bot() -> None:
    """Запускає Telegram бота."""
    global _orchestrator_instance, _bot_instance
    
    if _orchestrator_instance is None or not getattr(_orchestrator_instance, "_is_initialized", False):
        _orchestrator_instance = init_audio_drama_bot()

    if not _orchestrator_instance or not getattr(_orchestrator_instance, "_is_initialized", False):
        logger.error("Orchestrator not initialized. Bot will not start.")
        return

    existing_bot = _orchestrator_instance._bot
    should_recreate_bot = existing_bot is None
    if existing_bot is not None and hasattr(existing_bot, "is_bound_to_current_loop"):
        try:
            should_recreate_bot = not existing_bot.is_bound_to_current_loop()
        except Exception:
            should_recreate_bot = True

    if should_recreate_bot:
        token = (
            getattr(_orchestrator_instance, "_telegram_token", "") or
            (existing_bot.token if existing_bot else "") or
            (os.getenv("TELEGRAM_BOT_TOKEN") or "")
        ).strip()
        chat_id = (
            getattr(_orchestrator_instance, "_telegram_chat_id", "") or
            (existing_bot.chat_id if existing_bot else "") or
            (os.getenv("TELEGRAM_CHAT_ID") or "")
        ).strip()
        if not token or not chat_id:
            logger.error("Bot credentials are missing. Check TELEGRAM_BOT_TOKEN and TELEGRAM_CHAT_ID.")
            return
        if existing_bot is not None:
            logger.info("[AudioDramaBot] Recreating bot for current asyncio loop/thread...")
        fresh_bot = AudioDramaBot(token, chat_id)
        fresh_bot.set_orchestrator(_orchestrator_instance)
        initialized = await fresh_bot.initialize()
        if not initialized:
            logger.error("Failed to initialize Telegram bot in current loop.")
            return
        _orchestrator_instance._bot = fresh_bot

    if _orchestrator_instance._bot:
        logger.info("Starting Telegram bot...")
        await _orchestrator_instance._bot.run()
        return
    logger.error("Bot not initialized. Check TELEGRAM_BOT_TOKEN secret.")


def run_bot_sync() -> None:
    """Запускає бота синхронно (для Colab)."""
    try:
        loop = asyncio.get_event_loop()
        if loop.is_running():
            import nest_asyncio
            nest_asyncio.apply()
        loop.run_until_complete(run_bot())
    except RuntimeError:
        asyncio.run(run_bot())


def start_bot_polling_background() -> bool:
    """Запускає polling у daemon-thread, щоб не блокувати виконання клітинки Colab."""
    global _bot_polling_thread
    with _bot_polling_lock:
        if _bot_polling_thread is not None and _bot_polling_thread.is_alive():
            return False

        def _runner() -> None:
            loop = asyncio.new_event_loop()
            asyncio.set_event_loop(loop)
            try:
                loop.run_until_complete(run_bot())
            except Exception as exc:
                logger.error(f"[AudioDramaBot] Background polling thread crashed: {exc}")
            finally:
                try:
                    loop.run_until_complete(loop.shutdown_asyncgens())
                except Exception:
                    pass
                try:
                    loop.close()
                except Exception:
                    pass
                try:
                    asyncio.set_event_loop(None)
                except Exception:
                    pass

        _bot_polling_thread = threading.Thread(
            target=_runner,
            name="AudioDramaBotPolling",
            daemon=True
        )
        _bot_polling_thread.start()
        return True


# =============================================================================
# EXPORTS
# =============================================================================

__all__ = [
    # Data structures
    'Character',
    'DialogueLine',
    'SoundCue',
    'Scene',
    'Scenario',
    'SoundResult',
    'MixingConfig',
    'Checkpoint',
    'OrchestratorConfig',
    'DialogueQualityIssue',
    'DialogueQualityReport',
    
    # Enums
    'CharacterRole',
    'EmotionType',
    
    # API Clients
    'GeminiClient',
    'PexelsClient',
    
    # File parsing
    'FileParser',
    
    # Voice management
    'PronunciationManager',
    'AudioQualityGuard',
    'VoiceManager',
    
    # Audio processing
    'AudioMixer',
    'AudioAssembler',
    
    # Bot
    'AudioDramaBot',
    
    # Orchestrator
    'AudioDramaOrchestrator',
    
    # Colab integration
    'init_audio_drama_bot',
    'get_orchestrator',
    'run_bot',
    'run_bot_sync',
    'start_bot_polling_background',
    'is_colab_environment',
    'mount_google_drive',
    'get_colab_secrets',
]

__version__ = '1.1.20'
__author__ = 'VIBEMODLY Team'


# =============================================================================
# AUTO-INITIALIZATION FOR COLAB
# =============================================================================

def _auto_init_for_colab() -> None:
    """Автоматична ініціалізація при імпорті в Google Colab."""
    if not is_colab_environment():
        logger.info("[AudioDramaBot] Not in Colab - auto-init skipped")
        return
    
    auto_init_enabled = os.getenv("AUDIO_DRAMA_AUTO_INIT", "1").lower() in ("1", "true", "yes")
    if not auto_init_enabled:
        logger.info("[AudioDramaBot] AUTO-INIT disabled by AUDIO_DRAMA_AUTO_INIT=0")
        print("[AudioDramaBot] Auto-init disabled (AUDIO_DRAMA_AUTO_INIT=0)", flush=True)
        return
    
    print(f"[AudioDramaBot] Version: {__version__}", flush=True)
    print("[AudioDramaBot] Colab auto-init started...", flush=True)
    
    logger.info("\n" + "="*60)
    logger.info("[AudioDramaBot] AUTO-INITIALIZATION FOR COLAB")
    logger.info("="*60)
    
    # Перевіряємо секрети
    secrets = get_colab_secrets()
    
    missing_secrets = [k for k, v in secrets.items() if v is None]
    missing_required = [
        key for key in ("telegram_token", "telegram_chat_id", "gemini_api_key")
        if not secrets.get(key)
    ]
    
    if missing_secrets:
        logger.warning("\n⚠️  MISSING SECRETS:")
        for secret in missing_secrets:
            logger.warning(f"   - {secret.upper()}")
        
        logger.info("\nAdd these secrets in Colab (🔑 icon in left sidebar):")
        logger.info("   1. TELEGRAM_BOT_TOKEN - from @BotFather")
        logger.info("   2. TELEGRAM_CHAT_ID - from @userinfobot")
        logger.info("   3. GEMINI_API_KEY - from Google AI Studio")
        logger.info("   4. PEXELS_API_KEY - from Pexels.com")
        print("[AudioDramaBot] Missing some secrets. Check Colab Secrets panel (🔑).", flush=True)
    
    allow_partial = os.getenv("AUDIO_DRAMA_ALLOW_PARTIAL_INIT", "0").lower() in ("1", "true", "yes")
    if missing_required and not allow_partial:
        logger.warning(
            "[AudioDramaBot] Auto-init skipped: missing required secrets "
            "(TELEGRAM_BOT_TOKEN, TELEGRAM_CHAT_ID, GEMINI_API_KEY)."
        )
        logger.info(
            "[AudioDramaBot] Set secrets or export env vars, then run init_audio_drama_bot(use_colab_userdata=True)."
        )
        print(
            "[AudioDramaBot] Auto-init skipped: add TELEGRAM_BOT_TOKEN, TELEGRAM_CHAT_ID, GEMINI_API_KEY.",
            flush=True
        )
        return
    
    # Ініціалізуємо (safe defaults to avoid blocking on secret permission dialogs)
    use_userdata = os.getenv("AUDIO_DRAMA_USE_COLAB_USERDATA", "1").lower() in ("1", "true", "yes")
    auto_mount = os.getenv("AUDIO_DRAMA_AUTO_MOUNT_DRIVE", "1").lower() in ("1", "true", "yes")
    require_qwen_dual_models = os.getenv("AUDIO_DRAMA_REQUIRE_QWEN_DUAL_MODELS", "1").lower() in ("1", "true", "yes", "on")
    qwen_precheck_on_init = os.getenv("AUDIO_DRAMA_QWEN_PRECHECK_ON_INIT", "0").lower() in ("1", "true", "yes", "on")
    default_timeout = "600" if (require_qwen_dual_models and qwen_precheck_on_init) else "90"
    init_timeout = int(os.getenv("AUDIO_DRAMA_INIT_TIMEOUT_SEC", default_timeout))
    init_timeout = max(0, init_timeout)
    
    orchestrator = init_audio_drama_bot(
        auto_mount_drive=auto_mount,
        use_colab_userdata=use_userdata,
        init_timeout_sec=init_timeout
    )
    
    # One-script mode: автоматичний запуск Telegram polling, якщо бот ініціалізовано
    auto_start_bot = os.getenv("AUDIO_DRAMA_AUTO_START_BOT", "1").lower() in ("1", "true", "yes")
    auto_start_mode = os.getenv("AUDIO_DRAMA_AUTO_START_MODE", "background").strip().lower()
    start_in_background = auto_start_mode in ("background", "bg", "thread", "daemon")
    bot_ready = bool(
        orchestrator
        and getattr(orchestrator, "_is_initialized", False)
        and getattr(orchestrator, "_bot", None)
    )
    
    if auto_start_bot and bot_ready:
        mode_label = "background" if start_in_background else "foreground"
        print(f"[AudioDramaBot] Starting Telegram bot polling (one-script mode, {mode_label})...", flush=True)
        try:
            if start_in_background:
                started = start_bot_polling_background()
                if started:
                    print(
                        "[AudioDramaBot] Polling started in background. Cell execution can continue.",
                        flush=True
                    )
                else:
                    print("[AudioDramaBot] Polling is already running in background.", flush=True)
            else:
                run_bot_sync()
        except Exception as e:
            logger.error(f"[AudioDramaBot] Bot auto-start failed: {e}")
            print(f"[AudioDramaBot] Bot auto-start failed: {e}", flush=True)
    else:
        if not bot_ready:
            print("[AudioDramaBot] Bot not started: missing TELEGRAM_BOT_TOKEN/TELEGRAM_CHAT_ID.", flush=True)
        else:
            print("[AudioDramaBot] Auto-start bot disabled (AUDIO_DRAMA_AUTO_START_BOT=0).", flush=True)
    
    logger.info("\n" + "="*60)
    logger.info("✓ Audio Drama Bot is ready!")
    logger.info("="*60 + "\n")


# Додаємо alias модуля для сценарію "запуск як один великий cell у Colab"
if __name__ == "__main__":
    sys.modules.setdefault("audio_drama_bot", sys.modules[__name__])

# Автоматична ініціалізація
_auto_init_for_colab()
